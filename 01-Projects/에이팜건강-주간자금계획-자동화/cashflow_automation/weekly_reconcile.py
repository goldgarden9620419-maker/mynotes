# -*- coding: utf-8 -*-
"""금요일 주간 대조: 이번 주 계획 vs 실제 계좌별 입출금.

월요일 자금계획의 지급예정과 은행 3사 실제 입출금을 주말 전에 대조해,
차이(미집행·계획없는출금·금액차이)를 사용자에게 먼저 확인받는 보고서
(06_확인필요/주간대조_일시.xlsx)를 만든다 (2026-09-21 사용자 요청:
"원래 계획과 실 지출을 비교해 한번 더 확인").

기본 일정은 금요일 17:00 자동 생성 + 트레이 메뉴 수동 실행.
"""
from __future__ import annotations

from datetime import date, datetime, time as dtime, timedelta
from pathlib import Path
from typing import Optional

from common import (
    BANK_REFLECT_OK, MATCH_AMOUNT_DIFF, MATCH_DATE_DIFF, MATCH_MANUAL,
    MATCH_NOT_FOUND, MATCH_PAID, MATCH_PARTIAL, MATCH_UNPLANNED,
    REFLECT_PAID, now_local, week_monday,
)
from excel_report import account_label

_NAVY = "1F4E79"
_DIFF_COLORS = {
    MATCH_NOT_FOUND: "FFE699",     # 미집행(계획했는데 출금 없음) — 주황
    MATCH_UNPLANNED: "FFC7CE",     # 계획에 없는 출금 — 붉은
    MATCH_AMOUNT_DIFF: "FFF2CC",   # 금액 차이 — 노랑
    MATCH_DATE_DIFF: "FFF2CC",
    MATCH_PARTIAL: "FFF2CC",
    MATCH_MANUAL: "FFF2CC",
}
_DIFF_ORDER = [MATCH_NOT_FOUND, MATCH_UNPLANNED, MATCH_AMOUNT_DIFF,
               MATCH_DATE_DIFF, MATCH_PARTIAL, MATCH_MANUAL]


def _d(value):
    return value.date() if isinstance(value, datetime) else value


def build_reconcile_data(plans: list[dict], bank_rows: list[dict],
                         balances: dict, run_date: date,
                         date_window_days: int = 3,
                         name_threshold: float = 70) -> dict:
    """이번 주(월요일~run_date) 계획 vs 실제 대조 자료를 만든다.

    plans: 통합 지출계획 행(정상반영·지급완료), bank_rows: 표준 은행 행.
    balances: (은행, 계좌)별 현재 잔액. 반환 dict는 write_reconcile_workbook
    입력으로 쓴다.
    """
    import payment_matcher

    monday = week_monday(run_date)
    week_end = min(run_date, monday + timedelta(days=4))   # 이번 주 금요일까지

    week_plans = []
    for p in plans:
        pd = _d(p.get("자금계획 반영일") or p.get("지급예정일"))
        if pd is not None and monday <= pd <= week_end:
            week_plans.append(p)

    matched = payment_matcher.match_payments(
        week_plans, bank_rows, run_date,
        window_days=date_window_days, name_threshold=name_threshold,
        unplanned_days=max(0, (run_date - monday).days))

    diffs = [r for r in matched["results"]
             if r.get("대조결과") != MATCH_PAID]
    diffs += matched["unplanned"]
    diffs.sort(key=lambda r: (_DIFF_ORDER.index(r.get("대조결과"))
                              if r.get("대조결과") in _DIFF_ORDER else 9,
                              _d(r.get("지급예정일")
                                 or r.get("실제출금일")) or date.max))
    ok_rows = [r for r in matched["results"]
               if r.get("대조결과") == MATCH_PAID]

    # 계좌별: 주간 실제 흐름과 주초(월요일 시작) 잔액 (내부이체 포함 되돌림)
    accounts: dict = {}
    last_bank_date: Optional[date] = None
    for tx in bank_rows:
        if tx.get("반영상태") != BANK_REFLECT_OK:
            continue
        td = _d(tx.get("거래일"))
        if td is not None and (last_bank_date is None or td > last_bank_date):
            last_bank_date = td
        key = (tx.get("은행"), tx.get("계좌") or "")
        acct = accounts.setdefault(key, {
            "입금": 0.0, "출금": 0.0, "내부이체": 0.0, "마지막거래일": None})
        if td is None or td < monday or td > run_date:
            continue
        amt_in = tx.get("입금액") or 0.0
        amt_out = tx.get("출금액") or 0.0
        if tx.get("내부이체"):
            acct["내부이체"] += amt_in - amt_out
        else:
            acct["입금"] += amt_in
            acct["출금"] += amt_out
        if acct["마지막거래일"] is None or td > acct["마지막거래일"]:
            acct["마지막거래일"] = td

    rows = []
    for key in sorted(set(balances) | set(accounts)):
        acct = accounts.get(key) or {"입금": 0.0, "출금": 0.0,
                                     "내부이체": 0.0, "마지막거래일": None}
        current = float(balances.get(key) or 0.0)
        opening = current - acct["입금"] + acct["출금"] - acct["내부이체"]
        rows.append({
            "계좌": account_label(*key),
            "주초잔액": opening, "입금": acct["입금"], "출금": acct["출금"],
            "내부이체": acct["내부이체"], "현재잔액": current,
            "마지막거래일": acct["마지막거래일"],
        })

    # 은행 파일이 아직 없는 날짜의 미집행은 판정 보류 안내를 덧붙인다
    for rw in diffs:
        if rw.get("대조결과") == MATCH_NOT_FOUND:
            pdd = _d(rw.get("지급예정일"))
            if pdd is not None and (last_bank_date is None
                                    or pdd > last_bank_date):
                rw["비고"] = ((rw.get("비고") or "") +
                            " 은행 내역이 아직 없는 날짜 — 파일 갱신 후 "
                            "재확인").strip()

    plan_out = sum(p.get("예상금액") or 0 for p in week_plans)
    actual_out = sum(a["출금"] for a in rows)
    counts = {k: sum(1 for r in diffs if r.get("대조결과") == k)
              for k in _DIFF_ORDER}
    return {
        "기간": (monday, week_end), "실행일": run_date,
        "은행기준일": last_bank_date,
        "계좌": rows,
        "계획지출": plan_out, "실제출금": actual_out,
        "일치": len(ok_rows), "차이내역": diffs, "차이건수": counts,
    }


def write_reconcile_workbook(data: dict, out_path: Path) -> Path:
    """대조 자료를 주간대조 워크북(안내·계좌별대조·차이내역)으로 저장한다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def _head(ws, row, headers, widths):
        for c, (h, w) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.fill = PatternFill("solid", start_color=_NAVY)
            cell.font = Font(name="맑은 고딕", color="FFFFFF", bold=True,
                             size=10)
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(c)].width = w

    def _put(ws, row, col, value, *, bold=False, size=10, color="000000",
             fmt=None, fill=None):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(name="맑은 고딕", bold=bold, size=size, color=color)
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = PatternFill("solid", start_color=fill)
        return cell

    monday, week_end = data["기간"]
    n_diff = len(data["차이내역"])

    # ── 안내 ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "안내"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 90
    _put(ws, 1, 1, f"주간 대조 — 계획 vs 실제 입출금 ({monday} ~ {week_end})",
         bold=True, size=13, color=_NAVY)
    _put(ws, 2, 1, "확인 완료", bold=True)
    _put(ws, 2, 2, "아니오", fill="FFF2CC")
    msg = (f"차이 {n_diff}건을 확인해 주세요 — ① '계좌별대조' 시트에서 "
           "계좌별 입출금·잔액을 보고 ② '차이내역' 시트에서 "
           "미집행(주황)·계획없는출금(붉은)·금액차이(노랑)를 검토한 뒤 "
           "위 B2를 '예'로 바꿔 저장하면 검토 완료로 기록됩니다. "
           "미집행 건은 팀 지출계획 수정(다음 주 재제출)으로 관리하세요."
           if n_diff else
           "이번 주 계획과 실제 입출금이 모두 일치합니다. "
           "B2를 '예'로 바꿔 저장하면 검토 완료로 기록됩니다.")
    _put(ws, 3, 1, msg, color="C00000" if n_diff else "2E7D32")
    last = data.get("은행기준일")
    if last is None or last < data["실행일"]:
        _put(ws, 4, 1,
             f"⚠ 은행 파일이 {last or '없음'}까지만 있어 오늘 거래가 빠져 "
             "있을 수 있습니다 — 은행 3사 파일을 새로 받아 넣고 "
             "'주간 대조 확인'을 다시 실행하면 정확해집니다.",
             color="C00000")

    # ── 계좌별대조 ──────────────────────────────────────────────────
    ws = wb.create_sheet("계좌별대조")
    _put(ws, 1, 1, f"계좌별 입출금 대조 ({monday} ~ {week_end})",
         bold=True, size=12, color=_NAVY)
    _head(ws, 3, ["계좌", "주초 잔액", "이번 주 입금", "이번 주 출금",
                  "내부이체(순)", "현재 잔액", "마지막 거래일"],
          [26, 15, 15, 15, 14, 15, 13])
    r = 4
    for a in data["계좌"]:
        _put(ws, r, 1, a["계좌"])
        for c, k in ((2, "주초잔액"), (3, "입금"), (4, "출금"),
                     (5, "내부이체"), (6, "현재잔액")):
            _put(ws, r, c, round(a[k]), fmt="#,##0;[Red]-#,##0")
        d = a["마지막거래일"]
        if d is not None:
            _put(ws, r, 7, datetime.combine(d, dtime()),
                 fmt="yyyy-mm-dd")
        r += 1
    _put(ws, r, 1, "합계", bold=True)
    for c, k in ((2, "주초잔액"), (3, "입금"), (4, "출금"),
                 (5, "내부이체"), (6, "현재잔액")):
        _put(ws, r, c, round(sum(a[k] for a in data["계좌"])),
             bold=True, fmt="#,##0;[Red]-#,##0")
    r += 2
    _put(ws, r, 1, "계획 대비 요약", bold=True, color=_NAVY)
    _put(ws, r + 1, 1, "이번 주 계획 지출")
    _put(ws, r + 1, 2, round(data["계획지출"]), fmt="#,##0")
    _put(ws, r + 2, 1, "실제 외부 출금")
    _put(ws, r + 2, 2, round(data["실제출금"]), fmt="#,##0")
    _put(ws, r + 3, 1, "차이 (계획-실제)", bold=True)
    _put(ws, r + 3, 2, round(data["계획지출"] - data["실제출금"]),
         bold=True, fmt="#,##0;[Red]-#,##0")
    _put(ws, r + 4, 1, "대조 결과")
    cnt = data["차이건수"]
    _put(ws, r + 4, 2,
         f"일치 {data['일치']}건 · 미집행 {cnt.get(MATCH_NOT_FOUND, 0)}건 · "
         f"계획없는출금 {cnt.get(MATCH_UNPLANNED, 0)}건 · 금액·기타 차이 "
         f"{n_diff - cnt.get(MATCH_NOT_FOUND, 0) - cnt.get(MATCH_UNPLANNED, 0)}건")

    # ── 차이내역 ────────────────────────────────────────────────────
    ws = wb.create_sheet("차이내역")
    _put(ws, 1, 1, "차이 내역 — 계획했는데 안 나갔거나, 계획에 없거나, "
                   "금액이 다른 건", bold=True, size=12, color=_NAVY)
    _head(ws, 3, ["구분", "팀명", "거래처·내용", "지급예정일", "실제출금일",
                  "계획금액", "실제금액", "차이", "은행", "비고"],
          [14, 11, 34, 12, 12, 13, 13, 12, 10, 26])
    r = 4
    for row in data["차이내역"]:
        fill = _DIFF_COLORS.get(row.get("대조결과"))
        _put(ws, r, 1, row.get("대조결과"), fill=fill)
        _put(ws, r, 2, row.get("팀명") or "")
        _put(ws, r, 3, row.get("거래처") or "")
        for c, k in ((4, "지급예정일"), (5, "실제출금일")):
            d = _d(row.get(k))
            if d is not None:
                _put(ws, r, c, datetime.combine(d, dtime()),
                     fmt="yyyy-mm-dd")
        for c, k in ((6, "예정금액"), (7, "실제금액"), (8, "차이금액")):
            v = row.get(k)
            if v is not None:
                _put(ws, r, c, round(v), fmt="#,##0;[Red]-#,##0")
        _put(ws, r, 9, row.get("은행") or "")
        _put(ws, r, 10, row.get("비고") or "")
        r += 1
    if r == 4:
        _put(ws, 4, 1, "차이 없음 — 이번 주 계획과 실제가 모두 일치합니다.",
             color="2E7D32")
    else:
        ws.auto_filter.ref = f"A3:J{r - 1}"
    ws.freeze_panes = "A4"

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    wb.close()
    return out_path


def run_weekly_reconcile(cfg, log, now=None, open_file: bool = True) -> Path:
    """입력을 읽어 주간대조 파일을 만들고 (기본) 자동으로 연다."""
    import bank_classifier
    import bank_loader
    import duplicate_checker
    import forecast_engine
    import team_loader
    from card_payment import CardPaymentCalculator

    now = now or now_local(cfg.timezone_name)
    run_date = now.date()

    card_calc = CardPaymentCalculator.from_workbook(cfg.base_workbook_path())
    team_data = team_loader.load_all_teams(cfg)
    plan = team_loader.build_integrated_plan(
        team_data["rows"], card_calc.settlement_date,
        include_unconfirmed=cfg.get("forecast", "include_unconfirmed",
                                    default=True))
    forecast_engine.split_apalm_marked(plan)   # 에이팜 별도관리 건 제외

    bank_data = bank_loader.load_all_banks(cfg)
    kept, _dups = duplicate_checker.remove_duplicates(bank_data["rows"])
    rules = bank_classifier.load_rules(
        cfg.state_dir / cfg.get("bank", "classify_rules_file",
                                default="classify_rules.json"))
    bank_classifier.classify_rows(kept, rules)
    balances, _total = bank_loader.summarize_balances(kept)

    match_targets = plan["countable"] + [
        r for r in plan["integrated"] if r.get("반영상태") == REFLECT_PAID]
    data = build_reconcile_data(
        match_targets, kept, balances, run_date,
        date_window_days=cfg.get("matching", "date_window_days", default=3),
        name_threshold=cfg.get("matching", "name_similarity_threshold",
                               default=70))

    out = cfg.folder("review") / f"주간대조_{now:%Y%m%d_%H%M}.xlsx"
    write_reconcile_workbook(data, out)
    log.info("주간 대조(%s~%s): 일치 %d건, 차이 %d건 — %s",
             data["기간"][0], data["기간"][1], data["일치"],
             len(data["차이내역"]), out.name)
    if open_file:
        from app import _open_file
        _open_file(out)
    return out
