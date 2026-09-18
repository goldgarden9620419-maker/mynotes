# -*- coding: utf-8 -*-
"""팀별 제출양식·경영지원 대외비 양식·자금계획 기준파일 생성 스크립트.

실행:
    python tools/make_templates.py [출력폴더]

- 시트 '지출계획' + 표 'tbl_지출계획'
- 요청ID 자동수식 + 셀 잠금(시트 보호 암호: apharm)
- 지급방법·카드구분·확정여부·진행상태 드롭다운
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from common import (
    CONFIDENTIAL_CATEGORIES, CONFIRM_VALUES, PAY_METHODS, PROGRESS_VALUES,
    REQUIRED_TEAM_COLUMNS, TEAMS,
)
from card_payment import CARD_SHEET_COLUMNS, CARD_SHEET_NAME, DEFAULT_CARD_RULES

SHEET_PASSWORD = "apharm"
DATA_ROWS = 200

_HEADER_FILL = PatternFill("solid", start_color="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_THIN = Side(style="thin", color="C0C0C0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_COLUMN_WIDTHS = {
    "요청ID": 20, "팀코드": 8, "팀명": 14, "최초등록일": 12, "일련번호": 9,
    "지급예정일": 12, "거래처": 18, "지출내용": 28, "예상금액": 14,
    "지급방법": 11, "카드구분": 11, "확정여부": 9, "진행상태": 9,
    "최종수정일": 12, "비고": 20, "대외비구분": 22,
}


def _dv(values: list[str], ranges: str) -> DataValidation:
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"',
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="입력 오류",
                        error="목록에서 선택해 주세요.")
    dv.add(ranges)
    return dv


def _guide_sheet(wb: Workbook, team_name: str, confidential: bool) -> None:
    ws = wb.create_sheet("사용안내")
    ws.column_dimensions["A"].width = 95
    lines = [
        f"[{team_name}] 주간 지출계획 작성 안내",
        "",
        "1. '지출계획' 시트에만 입력합니다. 열 이름·시트명은 바꾸지 마세요.",
        "2. 요청ID는 자동으로 만들어집니다(잠금). 직접 입력하지 않습니다.",
        "   형식: 팀코드-최초등록일-일련번호 (예: 물류-20260918-002)",
        "3. 일련번호는 같은 날 등록한 순서대로 1, 2, 3… 을 직접 입력합니다.",
        "   (행 번호가 아닙니다. 15번째 요청이면 15를 입력 → 015로 표시)",
        "4. 금액·지급일·상태가 바뀌어도 요청ID는 그대로 두고,",
        "   진행상태를 '변경'으로 바꾸고 최종수정일을 수정합니다.",
        "5. 취소된 건은 행을 지우지 말고 진행상태를 '취소'로 바꿉니다.",
        "6. 지급방법이 '법인카드'이면 카드구분(우리카드/국민카드)을 선택하고,",
        "   지급예정일에는 카드 '사용일'을 입력합니다. 결제일은 자동 계산됩니다.",
        "7. 확정여부가 '확정'이고 진행상태가 '신규' 또는 '변경'인 건만",
        "   자금계획 합계에 반영됩니다.",
        "8. 매주 금요일에 다음 주 계획을 입력하고, 월요일 오전 9시 전까지",
        "   지정 폴더에 저장합니다.",
    ]
    if confidential:
        lines += [
            "",
            "[대외비 안내]",
            "9. 이 파일은 경영지원팀 대외비입니다. 02_경영지원_대외비 폴더에만",
            "   저장하고 다른 팀과 공유하지 않습니다.",
            "10. '대외비구분' 열에서 항목 분류를 선택하면 보고서에는",
            "    분류별 총액만 표시됩니다 (상세내용 미노출).",
        ]
    for r, line in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=line)
        if r == 1:
            cell.font = Font(bold=True, size=13, color="1F4E79")
    ws.sheet_view.showGridLines = False


def create_team_template(path: Path, team: dict) -> Path:
    confidential = bool(team.get("confidential"))
    columns = list(REQUIRED_TEAM_COLUMNS)
    if confidential:
        columns.append("대외비구분")

    wb = Workbook()
    ws = wb.active
    ws.title = "지출계획"

    for c, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = \
            _COLUMN_WIDTHS.get(name, 12)

    last_row = 1 + DATA_ROWS
    for r in range(2, last_row + 1):
        # 요청ID 자동 수식 (11번 항목의 일반 셀 수식)
        ws.cell(row=r, column=1, value=(
            f'=IF(OR(B{r}="",D{r}="",E{r}=""),"",'
            f'B{r}&"-"&TEXT(D{r},"yyyymmdd")&"-"&TEXT(E{r},"000"))'))
        for c, name in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c)
            cell.border = _BORDER
            # 요청ID만 잠금, 나머지는 입력 가능
            cell.protection = Protection(locked=(c == 1))
            if name in ("최초등록일", "지급예정일", "최종수정일"):
                cell.number_format = "yyyy-mm-dd"
            elif name == "예상금액":
                cell.number_format = "#,##0"
            elif name == "일련번호":
                cell.number_format = "0"
        ws.cell(row=r, column=2).value = None  # 팀코드는 직접 입력(아래 예시 참고)

    # 드롭다운
    def col_of(name: str) -> str:
        return get_column_letter(columns.index(name) + 1)

    rng = lambda name: f"{col_of(name)}2:{col_of(name)}{last_row}"
    ws.add_data_validation(_dv(PAY_METHODS, rng("지급방법")))
    ws.add_data_validation(_dv(["우리카드", "국민카드"], rng("카드구분")))
    ws.add_data_validation(_dv(CONFIRM_VALUES, rng("확정여부")))
    ws.add_data_validation(_dv(PROGRESS_VALUES, rng("진행상태")))
    ws.add_data_validation(_dv([team["code"]], rng("팀코드")))
    if confidential:
        ws.add_data_validation(_dv(CONFIDENTIAL_CATEGORIES, rng("대외비구분")))

    # 표 (구조화 참조 지원)
    table_ref = f"A1:{get_column_letter(len(columns))}{last_row}"
    table = Table(displayName="tbl_지출계획", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2",
                                          showRowStripes=True)
    ws.add_table(table)

    # 시트 보호: 잠금 해제된 셀만 입력 가능
    ws.protection.sheet = True
    ws.protection.password = SHEET_PASSWORD
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.insertRows = True
    ws.protection.sort = True
    ws.protection.autoFilter = True

    ws.freeze_panes = "A2"
    _guide_sheet(wb, team["name"], confidential)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def create_base_workbook(path: Path) -> Path:
    wb = Workbook()

    ws = wb.active
    ws.title = "안내"
    ws.column_dimensions["A"].width = 95
    for r, line in enumerate([
        "(주)에이팜건강 자금계획 기준파일",
        "",
        "· 카드결제기준: 법인카드 결제일 계산 기준. 실제 카드 약정에 맞게",
        "  결제일과 이용기간을 반드시 수정해 주세요.",
        "  이용기간은 '전전월/전월/당월 N일' 또는 '… 말일' 형식으로 입력합니다.",
        "· 주간조정: 자금계획에 수동으로 반영할 입금/지출 조정액을 입력합니다.",
        "  (예: 일회성 대출 상환, 투자금 입금 등)",
        "· 계좌목록: 회사 계좌 목록(참고용). 내부이체 판정 키워드는",
        "  00_프로그램/classify_rules.json에서 관리합니다.",
        "· 이 파일은 04_기준파일 폴더에 두고 파일명을 바꾸지 않습니다.",
    ], start=1):
        cell = ws.cell(row=r, column=1, value=line)
        if r == 1:
            cell.font = Font(bold=True, size=13, color="1F4E79")
    ws.sheet_view.showGridLines = False

    ws = wb.create_sheet(CARD_SHEET_NAME)
    for c, name in enumerate(CARD_SHEET_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = 16
    for r, rule in enumerate(DEFAULT_CARD_RULES, start=2):
        for c, name in enumerate(CARD_SHEET_COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=rule.get(name))
            cell.border = _BORDER
    ws.add_data_validation(_dv(["사용", "미사용"], "G2:G20"))

    ws = wb.create_sheet("주간조정")
    for c, (name, width) in enumerate([("일자", 12), ("조정입금", 14),
                                       ("조정지출", 14), ("내용", 40)],
                                      start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
    for r in range(2, 102):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=2).number_format = "#,##0"
        ws.cell(row=r, column=3).number_format = "#,##0"

    ws = wb.create_sheet("계좌목록")
    for c, (name, width) in enumerate([("은행", 12), ("계좌번호", 22),
                                       ("용도", 30)], start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, (bank, note) in enumerate([("농협", "주거래"),
                                      ("우리은행", "세금·보험 출금"),
                                      ("국민은행", "급여 이체")], start=2):
        ws.cell(row=r, column=1, value=bank)
        ws.cell(row=r, column=3, value=note)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def main() -> None:
    out_dir = Path(sys.argv[1]) if len(sys.argv) > 1 \
        else Path(__file__).resolve().parent.parent.parent / "양식"
    created = []
    for team in TEAMS:
        created.append(create_team_template(out_dir / team["file"], team))
    created.append(create_base_workbook(
        out_dir / "에이팜건강_자금계획_기준파일.xlsx"))
    for p in created:
        print("생성:", p)


if __name__ == "__main__":
    main()
