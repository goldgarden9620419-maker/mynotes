# -*- coding: utf-8 -*-
"""법인카드 결제일 계산 (13~14번 항목).

기준파일의 '카드결제기준' 시트를 읽어, 카드 사용일과 카드구분으로
예상 카드 결제일을 계산한다. 이용기간은 "전전월/전월/당월 N일|말일"
형식의 문자열로 정의한다.
"""
from __future__ import annotations

import calendar
import re
from datetime import date
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from common import normalize_text

CARD_SHEET_NAME = "카드결제기준"
CARD_SHEET_COLUMNS = ["카드구분", "카드사", "결제일", "이용기간 시작일",
                      "이용기간 종료일", "출금계좌", "사용여부"]

# 기본 기준표. 실제 카드 약정에 맞게 기준파일에서 반드시 수정한다.
DEFAULT_CARD_RULES = [
    {"카드구분": "우리카드", "카드사": "우리카드", "결제일": 15,
     "이용기간 시작일": "전월 1일", "이용기간 종료일": "전월 말일",
     "출금계좌": "우리은행 법인계좌", "사용여부": "사용"},
    {"카드구분": "국민카드", "카드사": "KB국민카드", "결제일": 25,
     "이용기간 시작일": "전월 13일", "이용기간 종료일": "당월 12일",
     "출금계좌": "국민은행 법인계좌", "사용여부": "사용"},
]

_OFFSETS = {"당월": 0, "전월": -1, "전전월": -2}
_PERIOD_RE = re.compile(r"(당월|전월|전전월)\s*(말일|(\d{1,2})\s*일?)")


def _last_day(year: int, month: int) -> int:
    return calendar.monthrange(year, month)[1]


def _shift_month(year: int, month: int, offset: int) -> tuple[int, int]:
    idx = year * 12 + (month - 1) + offset
    return idx // 12, idx % 12 + 1


def parse_period_spec(text: str) -> Optional[tuple[int, Optional[int]]]:
    """'전월 13일' → (-1, 13), '전월 말일' → (-1, None). 실패 시 None."""
    m = _PERIOD_RE.search(normalize_text(text))
    if not m:
        return None
    offset = _OFFSETS[m.group(1)]
    if m.group(2).startswith("말일"):
        return offset, None
    return offset, int(m.group(3))


def _resolve_period_date(settle_year: int, settle_month: int,
                         spec: tuple[int, Optional[int]]) -> date:
    offset, day = spec
    y, m = _shift_month(settle_year, settle_month, offset)
    if day is None:
        day = _last_day(y, m)
    return date(y, m, min(day, _last_day(y, m)))


class CardRule:
    def __init__(self, raw: dict):
        self.card_type = normalize_text(raw.get("카드구분"))
        self.issuer = normalize_text(raw.get("카드사"))
        try:
            self.settle_day = int(raw.get("결제일") or 0)
        except (TypeError, ValueError):
            self.settle_day = 0
        self.start_spec = parse_period_spec(raw.get("이용기간 시작일") or "")
        self.end_spec = parse_period_spec(raw.get("이용기간 종료일") or "")
        self.account = normalize_text(raw.get("출금계좌"))
        self.active = normalize_text(raw.get("사용여부")) in ("사용", "Y", "O", "예")

    @property
    def valid(self) -> bool:
        return bool(self.card_type) and 1 <= self.settle_day <= 31 \
            and self.start_spec is not None and self.end_spec is not None


class CardPaymentCalculator:
    def __init__(self, rules: Optional[list[dict]] = None):
        raw_rules = rules if rules is not None else DEFAULT_CARD_RULES
        self.rules = [CardRule(r) for r in raw_rules]

    # ------------------------------------------------------------------
    @classmethod
    def from_workbook(cls, path: Path) -> "CardPaymentCalculator":
        """기준파일에서 카드결제기준 시트를 읽는다. 없으면 기본값 사용."""
        path = Path(path)
        if not path.exists():
            return cls()
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception:
            return cls()
        try:
            sheet = None
            for name in wb.sheetnames:
                if normalize_text(name) == CARD_SHEET_NAME:
                    sheet = wb[name]
                    break
            if sheet is None:
                return cls()
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                return cls()
            header_idx, header = None, None
            for i, row in enumerate(rows[:10]):
                values = [normalize_text(v) for v in row]
                if "카드구분" in values and "결제일" in values:
                    header_idx, header = i, values
                    break
            if header is None:
                return cls()
            parsed = []
            for row in rows[header_idx + 1:]:
                raw = {header[i]: row[i] for i in range(len(header))
                       if i < len(row) and header[i]}
                if normalize_text(raw.get("카드구분")):
                    parsed.append(raw)
            return cls(parsed) if parsed else cls()
        finally:
            wb.close()

    # ------------------------------------------------------------------
    def find_rule(self, card_type: str) -> Optional[CardRule]:
        key = normalize_text(card_type)
        if not key:
            return None
        for rule in self.rules:
            if rule.valid and rule.active and rule.card_type == key:
                return rule
        # 부분 일치 (예: '우리' ↔ '우리카드')
        for rule in self.rules:
            if rule.valid and rule.active and (
                    key in rule.card_type or rule.card_type in key):
                return rule
        return None

    def settlement_date(self, usage_date: Optional[date],
                        card_type: str) -> Optional[date]:
        """카드 사용일 → 예상 결제일. 계산 불가면 None (결제일확인필요)."""
        if usage_date is None:
            return None
        rule = self.find_rule(card_type)
        if rule is None:
            return None
        # 사용일이 속할 수 있는 결제월 후보를 순서대로 확인
        for offset in range(0, 4):
            y, m = _shift_month(usage_date.year, usage_date.month, offset)
            start = _resolve_period_date(y, m, rule.start_spec)
            end = _resolve_period_date(y, m, rule.end_spec)
            if start <= usage_date <= end:
                return date(y, m, min(rule.settle_day, _last_day(y, m)))
        return None
