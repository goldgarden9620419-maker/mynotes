# -*- coding: utf-8 -*-
"""팀별 주간 지출계획 파일 읽기와 통합 (9~13, 19번 항목).

- 시트명 '지출계획', 표 'tbl_지출계획'을 우선 사용하되
  표가 없으면 열 이름을 찾아 읽는다.
- 요청ID 기준 중복/변경/취소 처리는 request_id 모듈 규칙을 따른다.
- 경영지원 대외비 행은 통합자료에 상세를 남기지 않고
  분류·총액 단위로만 노출한다.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Callable, Optional

from openpyxl import load_workbook

from common import (
    CONFIDENTIAL_MASK, PAY_METHOD_AUTO, PAY_METHOD_CARD,
    PAY_METHOD_TRANSFER, PAY_METHOD_UNDECIDED, PAY_METHODS,
    REFLECT_CARD_DATE_NEEDED, REFLECT_MISSING_INFO, REFLECT_OK,
    REFLECT_UNCONFIRMED, REQUIRED_TEAM_COLUMNS, TEAM_SHEET_NAME,
    TEAM_TABLE_NAME, TEAMS, classify_confidential, normalize_text,
    parse_amount, parse_date,
)
from request_id import (
    decide_reflect_status, is_valid_request_id, make_request_id,
    select_latest,
)

# 열 이름 유연 매핑: 공백 제거 후 비교
_COLUMN_ALIASES = {c.replace(" ", ""): c for c in REQUIRED_TEAM_COLUMNS}
_COLUMN_ALIASES["대외비구분"] = "대외비구분"
_COLUMN_ALIASES["신청자"] = "신청자"
# 취합본 양식(2026-09-18~): 확정여부 대신 품의승인(승인대기/승인/반려) 사용
_COLUMN_ALIASES["품의승인"] = "품의승인"

# 품의승인 값 → 확정여부 해석
_APPROVAL_CONFIRMED = {"승인", "승인완료", "완료", "확정"}
_APPROVAL_PENDING = {"승인대기", "대기", "상신", "미승인", "반려", "보류"}


class TeamFileError(Exception):
    pass


def _match_column(header: str) -> Optional[str]:
    key = normalize_text(header).replace(" ", "")
    return _COLUMN_ALIASES.get(key)


def _find_sheet(wb):
    for name in wb.sheetnames:
        if normalize_text(name) == TEAM_SHEET_NAME:
            return wb[name]
    return wb[wb.sheetnames[0]]


def _header_map_from_table(ws) -> Optional[tuple[int, dict[int, str]]]:
    """tbl_지출계획 표에서 (헤더행 번호, {열번호: 표준열이름}) 추출."""
    tables = getattr(ws, "tables", {}) or {}
    # openpyxl의 TableList.items()는 (이름, 범위문자열)을 돌려주므로
    # 이름으로 다시 조회해 Table 객체를 얻는다.
    for tname in list(tables):
        if normalize_text(tname) != TEAM_TABLE_NAME:
            continue
        table = tables[tname]
        ref = table.ref if hasattr(table, "ref") else str(table)
        from openpyxl.utils import range_boundaries
        min_col, min_row, max_col, _ = range_boundaries(ref)
        mapping = {}
        for col in range(min_col, max_col + 1):
            std = _match_column(ws.cell(row=min_row, column=col).value or "")
            if std:
                mapping[col] = std
        if mapping:
            return min_row, mapping
    return None


def _header_map_by_scan(ws) -> Optional[tuple[int, dict[int, str]]]:
    """상위 30행에서 필수 열 이름이 가장 많이 맞는 행을 헤더로 삼는다."""
    best = None
    for row in range(1, min(ws.max_row, 30) + 1):
        mapping = {}
        for col in range(1, min(ws.max_column, 40) + 1):
            std = _match_column(ws.cell(row=row, column=col).value or "")
            if std and std not in mapping.values():
                mapping[col] = std
        if len(mapping) >= 5 and (best is None or len(mapping) > len(best[1])):
            best = (row, mapping)
    return best


def load_team_file(path: Path, team: dict) -> tuple[list[dict], list[dict]]:
    """팀 파일 하나를 읽어 (행 목록, 확인필요 목록)을 돌려준다."""
    path = Path(path)
    issues: list[dict] = []
    if not path.exists():
        return [], [{"구분": "팀 자료 미제출", "팀명": team["name"],
                     "내용": f"파일 없음: {path.name}", "원본파일": path.name}]
    try:
        wb = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:
        return [], [{"구분": "손상된 파일", "팀명": team["name"],
                     "내용": f"파일 열기 실패: {exc}", "원본파일": path.name}]
    try:
        ws = _find_sheet(wb)
        header = _header_map_from_table(ws) or _header_map_by_scan(ws)
        if header is None:
            return [], [{"구분": "필수 열 누락", "팀명": team["name"],
                         "내용": "지출계획 열 이름을 찾지 못했습니다",
                         "원본파일": path.name}]
        header_row, mapping = header
        found_cols = set(mapping.values())
        missing_cols = [c for c in REQUIRED_TEAM_COLUMNS
                        if c not in found_cols and c != "비고"]
        # 품의승인 열이 있으면 확정여부를 대신한다
        if "확정여부" in missing_cols and "품의승인" in found_cols:
            missing_cols.remove("확정여부")
        if missing_cols:
            issues.append({"구분": "필수 열 누락", "팀명": team["name"],
                           "내용": "누락 열: " + ", ".join(missing_cols),
                           "원본파일": path.name})

        file_mtime = path.stat().st_mtime
        rows: list[dict] = []
        for r in range(header_row + 1, ws.max_row + 1):
            raw = {std: ws.cell(row=r, column=col).value
                   for col, std in mapping.items()}
            if all(v is None or str(v).strip() == "" for v in raw.values()):
                continue
            row = _normalize_row(raw, team, path.name, file_mtime, r)
            rows.append(row)
        return rows, issues
    finally:
        wb.close()


def _normalize_row(raw: dict, team: dict, filename: str,
                   file_mtime: float, row_order: int) -> dict:
    row = {
        "요청ID": normalize_text(raw.get("요청ID")),
        "팀코드": normalize_text(raw.get("팀코드")) or team["code"],
        "팀명": normalize_text(raw.get("팀명")) or team["name"],
        "최초등록일": parse_date(raw.get("최초등록일")),
        "일련번호": raw.get("일련번호"),
        "지급예정일": parse_date(raw.get("지급예정일")),
        "거래처": normalize_text(raw.get("거래처")),
        "지출내용": normalize_text(raw.get("지출내용")),
        "예상금액": parse_amount(raw.get("예상금액")),
        "지급방법": normalize_text(raw.get("지급방법")),
        "카드구분": normalize_text(raw.get("카드구분")),
        "확정여부": normalize_text(raw.get("확정여부")),
        "진행상태": normalize_text(raw.get("진행상태")),
        "최종수정일": parse_date(raw.get("최종수정일")),
        "비고": normalize_text(raw.get("비고")),
        "대외비구분": normalize_text(raw.get("대외비구분")),
        "신청자": normalize_text(raw.get("신청자")),
        "품의승인": normalize_text(raw.get("품의승인")),
        "원본파일": filename,
        "file_mtime": file_mtime,
        "row_order": row_order,
        "confidential": bool(team.get("confidential")),
        "반영상태": "",
        "확인사항": "",
        "자금계획 반영일": None,
    }
    # 품의승인 값으로 확정여부를 해석 (승인=확정, 대기·반려=미확정)
    if not row["확정여부"] and row["품의승인"]:
        approval = row["품의승인"]
        if approval in _APPROVAL_CONFIRMED:
            row["확정여부"] = "확정"
        elif approval in _APPROVAL_PENDING:
            row["확정여부"] = "미확정"
        else:
            row["확정여부"] = "미확정"
            row["확인사항"] = f"품의승인 값 해석 불가({approval})"
    # 요청ID가 비어있으면 구성요소로 재생성 시도
    if not row["요청ID"]:
        row["요청ID"] = make_request_id(row["팀코드"], row["최초등록일"],
                                        row["일련번호"])
    return row


def validate_row(row: dict) -> list[str]:
    """행 단위 필수값 검사 (18, 27번 항목). 문제 목록을 돌려준다."""
    problems = []
    if not row.get("요청ID"):
        problems.append("요청ID 누락")
    elif not is_valid_request_id(row["요청ID"]):
        problems.append("요청ID 형식 오류")
    if not row.get("팀명"):
        problems.append("팀명 누락")
    if row.get("지급예정일") is None:
        problems.append("지급예정일 누락")
    amount = row.get("예상금액")
    if amount is None or amount == 0:
        problems.append("금액 누락 또는 0원")
    if not row.get("지급방법"):
        problems.append("지급방법 누락")
    elif row["지급방법"] not in PAY_METHODS:
        problems.append(f"지급방법 값 오류({row['지급방법']})")
    return problems


# ---------------------------------------------------------------------------
# 전체 팀 취합 + 통합 데이터 생성
# ---------------------------------------------------------------------------

def load_all_teams(cfg) -> dict:
    """모든 팀 파일을 읽는다.

    반환: {"rows": [...], "issues": [...], "missing_teams": [...]}
    """
    all_rows: list[dict] = []
    issues: list[dict] = []
    missing_teams: list[str] = []
    for team, path in cfg.team_files():
        rows, file_issues = load_team_file(path, team)
        issues.extend(file_issues)
        if not path.exists():
            missing_teams.append(team["name"])
        all_rows.extend(rows)
    return {"rows": all_rows, "issues": issues, "missing_teams": missing_teams}


def build_integrated_plan(rows: list[dict],
                          card_date_fn: Optional[Callable] = None) -> dict:
    """요청ID 중복 정리 → 반영상태 결정 → 자금계획 반영일 계산.

    card_date_fn(사용일: date, 카드구분: str) -> Optional[date]
    반환: {"integrated": 전체 행(중복 포함), "countable": 반영 행,
           "issues": 확인필요 목록}
    """
    issues: list[dict] = []
    selected, duplicates = select_latest(rows)

    for row in selected:
        problems = validate_row(row)
        status = decide_reflect_status(row)
        if problems:
            # 취소 건은 어차피 합계 제외이므로 정보누락으로 바꾸지 않는다
            if status == REFLECT_OK:
                status = REFLECT_MISSING_INFO
            row["확인사항"] = "; ".join(problems)
            for p in problems:
                issues.append({"구분": p.split("(")[0], "팀명": row.get("팀명"),
                               "요청ID": row.get("요청ID"),
                               "내용": p, "원본파일": row.get("원본파일")})
        row["반영상태"] = status
        if status == REFLECT_UNCONFIRMED:
            # 미확정(승인대기 포함)은 합계 제외 + 확인필요 표시 (12번 항목)
            reason = row.get("품의승인") or row.get("확정여부") or "미확정"
            if row.get("confidential"):
                subject = classify_confidential(row.get("지출내용", ""),
                                                row.get("대외비구분", ""))
            else:
                subject = row.get("거래처", "")
            issues.append({"구분": "미확정(승인대기)", "팀명": row.get("팀명"),
                           "요청ID": row.get("요청ID"),
                           "내용": f"{reason} 상태라 합계 미반영: {subject} "
                                  f"{(row.get('예상금액') or 0):,.0f}원 "
                                  f"(예정일 {row.get('지급예정일')})",
                           "원본파일": row.get("원본파일")})
        _apply_plan_date(row, card_date_fn, issues)

    for dup in duplicates:
        issues.append({"구분": "요청ID 중복", "팀명": dup.get("팀명"),
                       "요청ID": dup.get("요청ID"),
                       "내용": dup.get("확인사항", "요청ID 중복"),
                       "원본파일": dup.get("원본파일")})

    integrated = selected + duplicates
    integrated.sort(key=lambda r: (r.get("자금계획 반영일") or r.get("지급예정일")
                                   or date.max, r.get("요청ID") or ""))
    countable = [r for r in selected if r["반영상태"] == REFLECT_OK]
    return {"integrated": integrated, "countable": countable, "issues": issues}


def _apply_plan_date(row: dict, card_date_fn, issues: list[dict]) -> None:
    """지급방법별 자금계획 반영일 결정 (13번 항목)."""
    method = row.get("지급방법")
    due = row.get("지급예정일")
    if row["반영상태"] != REFLECT_OK:
        row["자금계획 반영일"] = due
        return
    if method in (PAY_METHOD_TRANSFER, PAY_METHOD_AUTO):
        row["자금계획 반영일"] = due
    elif method == PAY_METHOD_CARD:
        settle = None
        if card_date_fn is not None and due is not None:
            settle = card_date_fn(due, row.get("카드구분") or "")
        if settle is None:
            row["반영상태"] = REFLECT_CARD_DATE_NEEDED
            row["자금계획 반영일"] = due
            note = "카드 결제일 계산 불가"
            row["확인사항"] = (row["확인사항"] + "; " + note).strip("; ")
            issues.append({"구분": "카드 결제일 누락", "팀명": row.get("팀명"),
                           "요청ID": row.get("요청ID"), "내용": note,
                           "원본파일": row.get("원본파일")})
        else:
            row["자금계획 반영일"] = settle
    elif method == PAY_METHOD_UNDECIDED:
        row["반영상태"] = REFLECT_MISSING_INFO
        row["자금계획 반영일"] = due
        note = "지급방법 미정: 확정 지출에서 제외"
        row["확인사항"] = (row["확인사항"] + "; " + note).strip("; ")
        issues.append({"구분": "지급방법 누락", "팀명": row.get("팀명"),
                       "요청ID": row.get("요청ID"), "내용": note,
                       "원본파일": row.get("원본파일")})
    else:
        row["자금계획 반영일"] = due


# ---------------------------------------------------------------------------
# 대외비 마스킹
# ---------------------------------------------------------------------------

def mask_confidential_rows(integrated: list[dict]) -> list[dict]:
    """통합 시트 표시용: 대외비 행을 분류·총액 집계로 치환한다.

    일반 행은 그대로 두고, 대외비 행은 (분류, 반영일, 지급방법, 반영상태)
    단위 합계 행으로 바꾼다. 거래처·지출내용·비고는 노출하지 않는다.
    """
    normal = [r for r in integrated if not r.get("confidential")]
    conf = [r for r in integrated if r.get("confidential")]
    grouped: dict[tuple, dict] = {}
    for row in conf:
        category = classify_confidential(row.get("지출내용", ""),
                                         row.get("대외비구분", ""))
        key = (category, row.get("자금계획 반영일"), row.get("지급방법"),
               row.get("반영상태"))
        g = grouped.setdefault(key, {
            "요청ID": f"{CONFIDENTIAL_MASK}-집계",
            "팀코드": "경영",
            "팀명": "경영지원팀",
            "최초등록일": None,
            "일련번호": "",
            "지급예정일": row.get("지급예정일"),
            "거래처": CONFIDENTIAL_MASK,
            "지출내용": category,
            "예상금액": 0.0,
            "지급방법": row.get("지급방법"),
            "카드구분": row.get("카드구분"),
            "확정여부": row.get("확정여부"),
            "진행상태": "",
            "최종수정일": None,
            "비고": f"{CONFIDENTIAL_MASK} 항목 합계",
            "원본파일": row.get("원본파일"),
            "반영상태": row.get("반영상태"),
            "확인사항": "",
            "자금계획 반영일": row.get("자금계획 반영일"),
            "confidential": True,
            "건수": 0,
        })
        g["예상금액"] = (g["예상금액"] or 0) + (row.get("예상금액") or 0)
        g["건수"] += 1
        g["비고"] = f"{CONFIDENTIAL_MASK} {g['건수']}건 합계"
    masked = normal + list(grouped.values())
    masked.sort(key=lambda r: (r.get("자금계획 반영일") or r.get("지급예정일")
                               or date.max, str(r.get("요청ID") or "")))
    return masked
