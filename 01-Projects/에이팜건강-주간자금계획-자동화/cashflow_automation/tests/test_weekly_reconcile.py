# -*- coding: utf-8 -*-
"""금요일 주간 대조(계획 vs 실제 입출금) 테스트."""
from datetime import date

from openpyxl import load_workbook

import weekly_reconcile as wr
from common import (BANK_REFLECT_OK, MATCH_NOT_FOUND, MATCH_PAID,
                    MATCH_UNPLANNED, REFLECT_OK)

RUN = date(2026, 9, 25)          # 금요일
MONDAY = date(2026, 9, 21)


def _plan(day, amount, vendor, req="건강-20260921-001"):
    return {"요청ID": req, "팀명": "건강사업팀", "거래처": vendor,
            "지출내용": f"{vendor} 결제", "예상금액": amount,
            "자금계획 반영일": day, "지급예정일": day,
            "지급방법": "계좌송금", "반영상태": REFLECT_OK}


def _tx(day, out=0.0, inc=0.0, name="", bank="우리은행", acct="220351",
        internal=False):
    return {"거래일": day, "거래일시": None, "은행": bank, "계좌": acct,
            "출금액": out, "입금액": inc, "거래후잔액": 0.0,
            "적요": "인터넷", "기재내용·상대방": name, "취급점": "",
            "자동분류": "", "내부이체": internal, "정기지출후보": False,
            "반영상태": BANK_REFLECT_OK, "원본파일": "은행.xlsx"}


def _data():
    plans = [
        _plan(date(2026, 9, 22), 1_000_000, "네이버"),          # 집행됨
        _plan(date(2026, 9, 24), 4_900_000, "메트라이프",
              req="경영-20260921-002"),                          # 미집행
        _plan(date(2026, 9, 15), 700_000, "지난주건",
              req="건강-20260914-003"),                          # 창 밖 — 제외
    ]
    rows = [
        _tx(date(2026, 9, 22), out=1_000_000, name="네이버"),
        _tx(date(2026, 9, 23), out=500_000, name="미지정업체"),   # 계획없는출금
        _tx(date(2026, 9, 24), out=300_000, name="농협이체", internal=True),
        _tx(date(2026, 9, 24), inc=300_000, name="우리이체", bank="농협",
            acct="301-569003", internal=True),
        _tx(date(2026, 9, 18), inc=700_000, name="지난주입금"),   # 주 밖
    ]
    balances = {("우리은행", "220351"): 5_000_000,
                ("농협", "301-569003"): 2_000_000}
    return wr.build_reconcile_data(plans, rows, balances, RUN)


def test_주간대조_집계():
    data = _data()
    assert data["기간"] == (MONDAY, RUN)
    assert data["일치"] == 1
    kinds = [r["대조결과"] for r in data["차이내역"]]
    assert kinds == [MATCH_NOT_FOUND, MATCH_UNPLANNED]
    miss = data["차이내역"][0]
    assert miss["거래처"] == "메트라이프" and miss["예정금액"] == 4_900_000
    # 은행 파일이 9/24까지라 미집행 판정에 보류 안내는 없다(예정일 9/24)
    assert "파일 갱신" not in (miss["비고"] or "")
    # 지난주 계획(9/15)은 이번 주 대조 대상이 아니다
    assert all(r.get("거래처") != "지난주건" for r in data["차이내역"])
    # 계획지출(1,000,000+4,900,000) vs 실제 외부 출금(1,500,000)
    assert round(data["계획지출"]) == 5_900_000
    assert round(data["실제출금"]) == 1_500_000
    # 계좌별: 주초잔액 = 현재 − 입금 + 출금 − 내부이체(순)
    by = {a["계좌"]: a for a in data["계좌"]}
    woori = next(v for k, v in by.items() if "우리" in k)
    nh = next(v for k, v in by.items() if "농협" in k)
    assert round(woori["주초잔액"]) == 6_800_000   # 5,000,000+1,500,000+300,000
    assert round(woori["출금"]) == 1_500_000
    assert round(woori["내부이체"]) == -300_000
    assert round(nh["주초잔액"]) == 1_700_000
    assert round(nh["내부이체"]) == 300_000
    # 주 밖(9/18) 입금은 주간 흐름에 포함되지 않는다
    assert round(woori["입금"]) == 0


def test_주간대조_워크북(tmp_path):
    data = _data()
    out = tmp_path / "주간대조_test.xlsx"
    wr.write_reconcile_workbook(data, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["안내", "계좌별대조", "차이내역"]
    info = wb["안내"]
    assert info["B2"].value == "아니오"          # 확인 완료 칸
    assert "차이 2건" in str(info["A3"].value)
    # 은행 파일(9/24)이 실행일(금 9/25)보다 이전 → 갱신 경고
    assert "은행 3사 파일을 새로 받아" in str(info["A4"].value)
    acc = wb["계좌별대조"]
    assert acc["A3"].value == "계좌"
    labels = [acc.cell(row=r, column=1).value for r in (4, 5, 6)]
    assert "합계" in labels
    diff = wb["차이내역"]
    assert diff["A4"].value == MATCH_NOT_FOUND
    assert str(diff["A4"].fill.start_color.rgb).endswith("FFE699")  # 주황
    assert diff["F4"].value == 4_900_000
    assert diff["A5"].value == MATCH_UNPLANNED
    assert str(diff["A5"].fill.start_color.rgb).endswith("FFC7CE")  # 붉은
    assert diff["G5"].value == 500_000
    assert diff.auto_filter.ref == "A3:J5"
    wb.close()
