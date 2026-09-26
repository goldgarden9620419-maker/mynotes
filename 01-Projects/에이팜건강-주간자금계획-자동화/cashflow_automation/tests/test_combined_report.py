# -*- coding: utf-8 -*-
"""통합 결과 파일(경영보고+대표보고+라이브) 테스트 (2026-09-23 사용자 요청).

경영보고 시트의 ③ 입금예정·④ 지출예정 표와 반영률(F12)이 입력 기준이고,
라이브 시트들은 SUMIFS·참조 수식으로 연동되어 경영보고에서 일자·금액을
고치면 전체가 재계산된다. 대표보고는 A4 한 장 인쇄용 수식 시트다.
"""
from datetime import date

from openpyxl import load_workbook

import management_report as mr
from tests.test_live_report import (
    NEW_MONDAY, _fake_report, _make_stub_template,
)


def _combined_report():
    rep = _fake_report(NEW_MONDAY)
    rep["meta"].update({"company": "(주)에이팜건강",
                        "run_at": "2026-09-23 17:10:00",
                        "run_date": date(2026, 9, 23)})
    rep["forecast"].update({
        "rate": 0.8, "actual_until": None, "intraday": None,
        "rate_scenarios": {
            0.8: {"4주 기말잔액": 1_000_000, "4주 최저잔액": 500_000,
                  "자금부족 예상일": None},
            0.9: {"4주 기말잔액": 2_000_000, "4주 최저잔액": 900_000,
                  "자금부족 예상일": None},
            1.0: {"4주 기말잔액": 3_000_000, "4주 최저잔액": 1_500_000,
                  "자금부족 예상일": None},
        },
    })
    rep["holidays"] = {}
    rep["receipt_rates"] = [0.8, 0.9, 1.0]
    rep["stability_target"] = 0
    rep["issues"] = []
    rep["missing_teams"] = []
    rep["week_expenses"] = [
        {"일자": date(2026, 9, 24), "구분": "물류팀",
         "내용": "한진택배 택배비", "금액": 500_000,
         "지급방법": "계좌송금", "요청ID": "물류-20260918-001"},
        {"일자": date(2026, 9, 25), "구분": "디자인팀",
         "내용": "리플렛 인쇄", "금액": 25_000, "지급방법": "법인카드",
         "요청ID": "디자인-20260918-001"},
        {"일자": date(2026, 9, 25), "구분": "확인반영",
         "내용": "확인필요 지시 반영: METLIFE", "금액": 300_000,
         "지급방법": ""},
    ]
    rep["adjustments"] = [{"일자": date(2026, 9, 24),
                           "조정입금": 1_000_000.0, "조정지출": 0.0,
                           "내용": "거래처 정산 입금"},
                          {"일자": date(2026, 10, 15),
                           "조정입금": 700_000.0, "조정지출": 0.0,
                           "내용": "창 밖 확정입금"}]
    return rep


def test_통합파일_시트구성과_경영보고_연동(tmp_path):
    template = tmp_path / "템플릿.xlsx"
    _make_stub_template(template)
    out = tmp_path / "주간자금계획_test.xlsx"
    rep = _combined_report()
    mr.create_combined_workbook(rep, template, out)
    assert mr.verify_combined_workbook(out, NEW_MONDAY)

    wb = load_workbook(out)
    # 시트 순서(2026-09-23 사용자 요청): 요약·업데이트운영이 맨 앞,
    # 그다음 경영보고(입력 기준) → 대표보고 → 라이브 시트들
    assert wb.sheetnames[:4] == ["요약", "업데이트운영",
                                 mr.SHEET_NAME, mr.CEO_SHEET]
    assert "4주일별계획" in wb.sheetnames

    # 반영률 입력은 경영보고 F12 하나 — 요약 B13이 참조한다
    assert wb["요약"]["B13"].value == f"='{mr.SHEET_NAME}'!$F$12"

    # 4주일별계획 계획 행의 지출 칸 = 경영보고 ④ 표 SUMIFS (계좌 0개
    # 배치: 온라인3·확정4·송금5·카드6·조정7). 9/24 = 8+3행
    daily = wb["4주일별계획"]
    tr = str(daily.cell(row=11, column=5).value)
    assert tr.startswith("=SUMIFS") and "'경영보고'!$E$16:$E$109" in tr
    assert "계좌송금" in tr
    card = str(daily.cell(row=11, column=6).value)
    assert "법인카드" in card
    conf = str(daily.cell(row=11, column=4).value)
    assert "'경영보고'!$E$114:$E$135" in conf and "확정입금" in conf
    # 조건 범위(C열)도 같은 행 구간이어야 한다 (v9 회귀 방지)
    assert f"$C${mr._INC_FIRST}:$C${mr._INC_LAST}" in conf.replace(
        "'경영보고'!", "")
    etc = str(daily.cell(row=11, column=7).value)
    assert etc.startswith("=SUMIFS") and etc.endswith("-E11-F11")

    # 경영보고 ④의 확정입금 도우미(K)도 ③ 표 SUMIFS — ③에서 일자를
    # 고치면 ④·라이브가 함께 움직인다 (9/24 = 143행)
    ws = wb[mr.SHEET_NAME]
    k17 = str(ws["K143"].value)
    assert k17.startswith("=SUMIFS($E$114") and "확정입금" in k17
    # ③ 확정입금 행은 수정 가능(노란 칸)
    inc_rows = [r for r in range(114, 136) if ws.cell(row=r, column=3).value
                == "확정입금"]
    assert inc_rows
    first_inc = inc_rows[0]
    assert str(ws.cell(row=first_inc, column=5).fill.start_color.rgb) \
        .endswith("FFF2CC")
    # 창 밖(10/15) 확정입금도 ③에 담긴다(숨김 행) — SUMIFS 반영용
    hidden_inc = [r for r in inc_rows
                  if ws.cell(row=r, column=4).value == "창 밖 확정입금"]
    assert hidden_inc and ws.row_dimensions[hidden_inc[0]].hidden

    wb.close()


def test_대표보고_시트는_A4_한장_수식_연동(tmp_path):
    template = tmp_path / "템플릿.xlsx"
    _make_stub_template(template)
    out = tmp_path / "주간자금계획_ceo.xlsx"
    mr.create_combined_workbook(_combined_report(), template, out)

    wb = load_workbook(out)
    ceo = wb[mr.CEO_SHEET]
    # 핵심 수치가 전부 경영보고 수식 참조 (F12·④ 수정 시 즉시 갱신)
    assert str(ceo["C5"].value) == f"='{mr.SHEET_NAME}'!C12"  # 총잔액(고정)
    assert str(ceo["C6"].value) == f"='{mr.SHEET_NAME}'!E167"  # 4주 기말
    assert "SUMIFS" in str(ceo["C9"].value)                   # 카드 4주
    assert str(ceo["C10"].value).startswith("=IFERROR(")      # 부족 예상일
    # 일별 전망·시나리오 표도 경영보고 참조
    joined = " ".join(str(c.value) for row in ceo.iter_rows()
                      for c in row if c.value is not None)
    assert f"'{mr.SHEET_NAME}'!" in joined
    # A4 세로 한 장 인쇄 설정 — 그대로 프린트하면 바로 보고 가능
    assert ceo.page_setup.paperSize == 9
    assert ceo.page_setup.orientation == "portrait"
    assert int(ceo.page_setup.fitToWidth or 0) == 1
    assert int(ceo.page_setup.fitToHeight or 0) == 1
    assert str(ceo.print_area)
    wb.close()


def test_연동_수식은_정적_값과_일치한다(tmp_path):
    """SUMIFS 연동 수식을 표에서 시뮬레이션하면 엔진의 일별 계획 값과
    같아야 한다 (이월일 리터럴 포함)."""
    from datetime import datetime, timedelta
    import re

    template = tmp_path / "템플릿.xlsx"
    _make_stub_template(template)
    rep = _combined_report()
    base = NEW_MONDAY
    # 엔진 일별 값과 ③·④ 표를 서로 일치하게 구성한다 (운영에선 같은
    # 원천에서 나오므로 항상 일치): 9/24 송금 50만+확정 100만,
    # 9/25 카드 2.5만+기타(확인반영) 30만. 9/23 실행일은 실적 마감,
    # 미집행 송금 20만은 9/24로 이월
    for row in rep["forecast"]["daily"]:
        row.update({"확정·기타입금": 0, "팀별 송금예정": 0, "카드결제": 0,
                    "자동이체": 0, "기타지출": 0, "실적": False})
        if row["일자"] == date(2026, 9, 23):
            row.update({"당일실적": True, "기타지출": 150_000})
        if row["일자"] == date(2026, 9, 24):
            row.update({"팀별 송금예정": 200_000 + 500_000,
                        "확정·기타입금": 1_000_000})
        if row["일자"] == date(2026, 9, 25):
            row.update({"카드결제": 25_000, "기타지출": 300_000})
        if row["일자"] == date(2026, 10, 15):
            # 창 밖 확정입금 — ③ 숨김 행으로 담겨 SUMIFS에 잡힌다
            row.update({"확정·기타입금": 700_000})
    rep["forecast"]["intraday"] = {
        "일자": date(2026, 9, 23), "이월일": date(2026, 9, 24),
        "남은계획": {"송금": 200_000, "카드": 0, "자동이체": 0},
        "_이월확정": 0, "_이월조정지출": 0,
        "온라인실제": 0, "기타입금실제": 0, "입금실제": 0,
        "출금실제": 150_000, "계획외지출": 150_000,
        "집행": {"송금": 0, "카드": 0, "자동이체": 0}, "보류내역": [],
        "확정도착": 0, "대조내역": []}
    out = tmp_path / "주간자금계획_sim.xlsx"
    mr.create_combined_workbook(rep, template, out)

    wb = load_workbook(out)
    ws = wb[mr.SHEET_NAME]

    def _day(v):
        return v.date() if isinstance(v, datetime) else v

    def _sum_exp(d, method=None):
        total = 0.0
        for r in range(16, 110):
            if _day(ws.cell(row=r, column=1).value) != d:
                continue
            m = str(ws.cell(row=r, column=6).value or "")
            if method is not None and m != method:
                continue
            total += ws.cell(row=r, column=5).value or 0
        return total

    def _sum_conf(d):
        return sum(ws.cell(row=r, column=5).value or 0
                   for r in range(114, 136)
                   if _day(ws.cell(row=r, column=1).value) == d
                   and ws.cell(row=r, column=3).value == "확정입금")

    daily = wb["4주일별계획"]
    lit = re.compile(r"\+(\d+)(?=[-]|$)")
    for i, src in enumerate(rep["forecast"]["daily"]):
        row = 8 + i
        d = src["일자"]
        if src.get("실적") or src.get("당일실적"):
            continue
        tr_f = str(daily.cell(row=row, column=5).value)
        card_f = str(daily.cell(row=row, column=6).value)
        conf_f = str(daily.cell(row=row, column=4).value)
        etc_f = str(daily.cell(row=row, column=7).value)

        def _carr(f):
            m = lit.search(f)
            return int(m.group(1)) if m else 0
        tr = _sum_exp(d, "계좌송금") + _carr(tr_f)
        card = _sum_exp(d, "법인카드") + _carr(card_f)
        etc = _sum_exp(d) + _carr(etc_f) - tr - card
        conf = _sum_conf(d) + _carr(conf_f)
        assert tr == src["팀별 송금예정"], (d, "송금", tr)
        assert card == src["카드결제"], (d, "카드", card)
        assert etc == src["자동이체"] + src["기타지출"], (d, "기타", etc)
        assert conf == src["확정·기타입금"], (d, "확정", conf)
    wb.close()


def test_요약_수정안내표와_업데이트운영_안내(tmp_path):
    """2026-09-23 사용자 요청: 요약 시트에 경영보고 수정 방법과 수정 시
    각 시트 어디가 바뀌는지 표로, 업데이트운영 시트는 이 파일 운영법
    안내로 재작성, 두 시트를 맨 앞에 배치."""
    template = tmp_path / "템플릿.xlsx"
    _make_stub_template(template)
    out = tmp_path / "주간자금계획_guide.xlsx"
    rep = _combined_report()
    rep["bank_rows"] = [
        {"은행": "농협", "거래일": date(2026, 9, 21), "입금액": 1000},
        {"은행": "농협", "거래일": date(2026, 9, 18), "출금액": 500},
        {"은행": "우리은행", "거래일": date(2026, 9, 22), "입금액": 200,
         "반영상태": "중복제외"},
    ]
    mr.create_combined_workbook(rep, template, out)

    wb = load_workbook(out)
    ws = wb["요약"]
    # 사용 방법 문구가 통합 파일 기준으로 갱신됨
    assert mr.SHEET_NAME in str(ws["D13"].value)
    # 수정 안내 표: 띠 + 반영률·보류·확정입금 행이 있고 반영 위치가 적힘
    assert "수정 방법" in str(ws["A19"].value)
    body = {str(ws.cell(row=r, column=1).value or ""):
            str(ws.cell(row=r, column=4).value or "")
            for r in range(21, 27)}
    assert any("반영률" in k for k in body)
    assert any("I열" in k and "대표보고" in v for k, v in body.items())
    assert any("실지출" in k and "H열" in v for k, v in body.items())

    guide = wb["업데이트운영"]
    text = " ".join(str(c.value) for row in guide.iter_rows(max_row=40)
                    for c in row if c.value is not None)
    # 옛 채팅 업로드 안내가 없어지고 폴더 운영 안내로 바뀜
    assert "대화창" not in text
    assert "03_은행거래내역" in text and "매주 운영 순서" in text
    assert "시트 구성" in text and mr.CEO_SHEET in text
    # 은행 자료 자동 집계: 중복제외 행은 빠진다
    rows = {str(guide.cell(row=r, column=1).value):
            guide.cell(row=r, column=2).value
            for r in range(1, 45)}
    assert rows.get("농협") == 2
    assert "우리은행" not in rows
    wb.close()
