# -*- coding: utf-8 -*-
"""전 팀 통일 취합 양식(2026-09-18~, 확정여부·품의승인 열 없음) 테스트."""
from datetime import date

from openpyxl import Workbook

import team_loader

COLUMNS = ["요청ID", "팀코드", "팀명", "신청자", "최초등록일", "일련번호",
           "지급예정일", "거래처", "지출내용", "예상금액", "지급방법",
           "카드구분", "진행상태", "최종수정일", "비고", "대외비구분"]


def _write(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "지출계획"
    for c, name in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=c, value=name)
    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(COLUMNS, start=1):
            ws.cell(row=r, column=c, value=row.get(name))
    wb.save(path)
    wb.close()


def _row(code, name, serial, vendor, amount, method="계좌송금", **extra):
    row = {"팀코드": code, "팀명": name, "신청자": "홍길동",
           "최초등록일": date(2026, 9, 18), "일련번호": serial,
           "지급예정일": date(2026, 9, 25), "거래처": vendor,
           "지출내용": "테스트 지출", "예상금액": amount,
           "지급방법": method, "최종수정일": date(2026, 9, 18)}
    row.update(extra)
    return row


def _sample_rows():
    return [
        _row("디자인", "디자인팀", 1, "와우프레스", 25000,
             method="법인카드", 카드구분="국민카드"),
        _row("경영", "경영지원팀", 1, "서울보증보험", 632250,
             대외비구분="기타 대외비 지출"),
        _row("물류", "물류팀", 1, "택배사", 300000, 진행상태="취소"),
        _row("외계", "외계팀", 1, "미지의거래처", 10000),  # 인식 불가 팀
    ]


class _Cfg:
    """load_all_teams용 최소 설정 스텁."""

    def __init__(self, base):
        self.base = base

    def folder(self, name):
        return {"confidential_admin": self.base / "02",
                "general_teams": self.base / "01"}[name]

    def team_files(self):
        from common import TEAMS
        out = []
        for t in TEAMS:
            folder = ("confidential_admin" if t.get("confidential")
                      else "general_teams")
            out.append((t, self.folder(folder) / t["file"]))
        return out


def test_취합파일_전팀_인식(tmp_path):
    (tmp_path / "02").mkdir()
    (tmp_path / "01").mkdir()
    _write(tmp_path / "02" / "지출계획취합_2026-09-18.xlsx", _sample_rows())
    # 개별 팀 파일은 더 이상 읽지 않는다
    _write(tmp_path / "01" / "디자인팀_주간지출계획.xlsx",
           [_row("디자인", "디자인팀", 9, "무시될거래처", 99999)])

    data = team_loader.load_all_teams(_Cfg(tmp_path))
    assert data["consolidated_file"] == "지출계획취합_2026-09-18.xlsx"
    assert data["missing_teams"] == []
    vendors = {r["거래처"] for r in data["rows"]}
    assert "와우프레스" in vendors and "무시될거래처" not in vendors
    # 경영지원 행만 대외비
    conf = {r["팀명"]: r["confidential"] for r in data["rows"]}
    assert conf["경영지원팀"] is True and conf["디자인팀"] is False
    # 인식 불가 팀은 일반팀으로 처리 + 확인필요 표시
    assert any(i["구분"] == "팀 인식 불가" for i in data["issues"])
    # 취합에 없는 팀은 정보성 표시 (건강·CS·데이터·연구)
    absent = [i["팀명"] for i in data["issues"] if i["구분"] == "취합 미포함 팀"]
    assert set(absent) == {"건강사업팀", "CS팀", "데이터사이언스팀", "연구팀"}


def test_취합파일_정책_기본반영(tmp_path):
    """확정여부 열이 없으면 전 건 확정 취급, 취소만 제외. 잡음 없음."""
    (tmp_path / "02").mkdir()
    path = tmp_path / "02" / "취합.xlsx"
    _write(path, _sample_rows())
    rows, issues, present = team_loader.load_consolidated_file(path)
    assert present >= {"디자인팀", "경영지원팀", "물류팀"}

    plan = team_loader.build_integrated_plan(
        rows, card_date_fn=lambda due, card: due, include_unconfirmed=True)
    countable_ids = {r["팀명"] for r in plan["countable"]}
    assert "물류팀" not in countable_ids          # 취소 건 제외
    assert {"디자인팀", "경영지원팀"} <= countable_ids
    # 열 자체가 없으므로 승인대기 잡음이 생기지 않는다
    assert not any(i["구분"].startswith("승인대기") for i in plan["issues"])
    assert not any(i["구분"].startswith("미확정") for i in plan["issues"])

    masked = team_loader.mask_confidential_rows(plan["integrated"])
    assert not any("서울보증보험" in str(r.get("거래처", "")) for r in masked)
