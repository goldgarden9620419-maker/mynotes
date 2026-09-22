# -*- coding: utf-8 -*-
"""실적 반영(지난 날짜)·추정 중복 제외·정기지출분류 오버라이드 테스트."""
from datetime import date

import forecast_engine as fe


def _tx(d, in_amt=0.0, out_amt=0.0, cls="", internal=False):
    return {"거래일": d, "입금액": in_amt, "출금액": out_amt,
            "자동분류": cls, "내부이체": internal}


def test_지난날짜는_실적으로_채운다():
    base = date(2026, 9, 14)
    history = [
        _tx(date(2026, 9, 1), in_amt=1_000_000, cls="온라인매출입금"),
        _tx(date(2026, 9, 15), in_amt=300_000, cls="온라인매출입금"),
        _tx(date(2026, 9, 15), out_amt=100_000),
        _tx(date(2026, 9, 16), in_amt=50_000),
        _tx(date(2026, 9, 16), in_amt=999_999, internal=True),  # 제외
    ]
    fc = fe.build_forecast([], base, 1_000_000, history, [], [],
                           [0.8], 0.8)
    assert fc["actual_until"] == date(2026, 9, 16)
    # 시작잔액 = 현재잔액 - 실적 순증감(30만-10만+5만)
    assert round(fc["start_balance"]) == 1_000_000 - 250_000
    daily = {r["일자"]: r for r in fc["daily"]}
    d15 = daily[date(2026, 9, 15)]
    assert d15["실적"] is True
    assert d15["온라인 예상입금"] == 300_000
    assert d15["기타지출"] == 100_000
    d16 = daily[date(2026, 9, 16)]
    assert d16["확정·기타입금"] == 50_000
    # 실적 마지막 날 기말잔액이 현재잔액과 일치 (이중계산 없음)
    assert round(d16["기말잔액"]) == 1_000_000
    assert daily[date(2026, 9, 17)]["실적"] is False


def test_실행일은_실적으로_마감하고_미집행은_익일_이월():
    """실행일 행은 은행 파일 그대로 실적으로 마감하고(순현금흐름 =
    실제 입출금만), 안 나간 지급예정은 사라지지 않고 익일로 이월된다
    (2026-09-21 사용자 확정 — 익일부터 실잔고 기준 자금 예산).
    """
    base = date(2026, 9, 21)                       # 월요일 = 실행일
    history = [
        _tx(date(2026, 9, 1), in_amt=1_000_000, cls="온라인매출입금"),
        # 실행일 새벽 실적: 입금 500,000 / 출금 100,000 (순 +400,000)
        _tx(base, in_amt=500_000, cls="온라인매출입금"),
        _tx(base, out_amt=100_000),
    ]
    plans = [{"자금계획 반영일": base, "예상금액": 3_498_000.0,
              "지급방법": "계좌송금", "확정여부": "확정"}]
    fc = fe.build_forecast(plans, base, 10_000_000, history, [], [],
                           [0.8], 0.8, run_date=base)
    # 당일은 실적 구간에서 제외 (기준일 전날 → None)
    assert fc["actual_until"] is None
    # 시작잔액 = 현재잔액 − 당일 순증감 (전일 마감)
    assert round(fc["start_balance"]) == 10_000_000 - 400_000
    d21 = fc["daily"][0]
    assert d21["당일실적"] is True
    # 실행일 = 실적 마감: 순현금흐름은 실제 입출금만, 기말 = 현재 실잔고
    assert round(d21["온라인 예상입금"]) == 500_000
    assert round(d21["기타지출"]) == 100_000       # 계획 밖 실제 출금
    assert round(d21["팀별 송금예정"]) == 0        # 예정 금액 미포함
    assert round(d21["기초잔액"]) == 9_600_000
    assert round(d21["기말잔액"]) == 10_000_000
    # 안 나간 지급예정 3,498,000은 익일(9/22 화)로 이월된다
    d22 = fc["daily"][1]
    assert round(d22["팀별 송금예정"]) == 3_498_000
    assert "이월" in d22["비고"]

    # 실행일이 수요일이면 월·화까지만 실적으로 확정한다
    run_wed = date(2026, 9, 23)
    history2 = history + [
        _tx(date(2026, 9, 22), in_amt=200_000, cls="온라인매출입금"),
        _tx(run_wed, out_amt=50_000),
    ]
    fc2 = fe.build_forecast(plans, base, 10_000_000, history2, [], [],
                            [0.8], 0.8, run_date=run_wed)
    assert fc2["actual_until"] == date(2026, 9, 22)
    daily2 = {r["일자"]: r for r in fc2["daily"]}
    assert daily2[date(2026, 9, 22)]["실적"] is True
    assert daily2[run_wed]["실적"] is False
    # run_date를 안 주면(과거 호환) 기존처럼 마지막 은행일까지 실적
    fc3 = fe.build_forecast(plans, base, 10_000_000, history2, [], [],
                            [0.8], 0.8)
    assert fc3["actual_until"] == run_wed


def test_당일거래가_은행파일에_없으면_실적마감하지_않는다():
    """아침 실행(전일자 은행 파일)이면 당일 거래가 없으므로 실행일 행은
    예전처럼 송금예정·조정 계획 그대로 표시한다 (2026-09-22 사용자 확정).
    당일 파일을 올린 뒤 실행하면 위 테스트처럼 실적으로 마감된다."""
    base = date(2026, 9, 21)                       # 월요일 = 실행일
    history = [                                     # 전일까지의 거래뿐
        _tx(date(2026, 9, 1), in_amt=1_000_000, cls="온라인매출입금"),
        _tx(date(2026, 9, 18), in_amt=300_000, cls="온라인매출입금"),
    ]
    plans = [{"자금계획 반영일": base, "예상금액": 3_498_000.0,
              "지급방법": "계좌송금", "확정여부": "확정"}]
    fc = fe.build_forecast(plans, base, 10_000_000, history, [], [],
                           [0.8], 0.8, run_date=base)
    assert fc["intraday"] is None
    d21 = fc["daily"][0]
    assert not d21.get("당일실적")
    # 오늘 행에 지출예정이 계획대로 남아 있고, 익일로 이월되지 않는다
    assert round(d21["팀별 송금예정"]) == 3_498_000
    d22 = fc["daily"][1]
    assert round(d22["팀별 송금예정"]) == 0
    assert "이월" not in (d22["비고"] or "")


def test_계좌별시나리오_당일거래_되돌리기():
    """backout_from(실행일) 이후 거래는 시작 잔액에서 되돌린다 (내부이체 포함)."""
    run = date(2026, 9, 21)
    acc_w = ("우리은행", "220351")
    acc_n = ("농협", "301-569003")
    balances = {acc_w: 5_000_000, acc_n: 2_000_000}
    hist = [
        {"거래일": date(2026, 9, 18), "은행": "우리은행", "계좌": "220351",
         "입금액": 700_000, "출금액": 0, "내부이체": False},   # 과거 — 무관
        {"거래일": run, "은행": "우리은행", "계좌": "220351",
         "입금액": 500_000, "출금액": 100_000, "내부이체": False},
        {"거래일": run, "은행": "농협", "계좌": "301-569003",
         "입금액": 0, "출금액": 300_000, "내부이체": True},    # 내부이체도 되돌림
    ]
    daily = [{"일자": run, "요일": "월", "실적": False,
              "온라인 예상입금": 0, "확정·기타입금": 0, "팀별 송금예정": 0,
              "카드결제": 0, "자동이체": 0, "기타지출": 0}]
    sc = fe.build_account_scenario(daily, balances, hist, None,
                                   backout_from=run)
    # 우리: 5,000,000 − (500,000−100,000) = 4,600,000 / 농협: 2,000,000 + 300,000
    assert round(sc["opening"][acc_w]) == 4_600_000
    assert round(sc["opening"][acc_n]) == 2_300_000
    # backout_from 없으면 기존 그대로
    sc0 = fe.build_account_scenario(daily, balances, hist, None)
    assert round(sc0["opening"][acc_w]) == 5_000_000


def test_팀계획과_겹치는_자동초안_조정_제외():
    adjustments = [
        # 이름·금액·날짜가 모두 비슷 → 중복으로 제외
        {"일자": date(2026, 9, 21), "조정입금": 0.0, "조정지출": 223_019.0,
         "내용": "정기지출 추정(자동 초안): SKB 신뢰도 상"},
        # 금액은 거의 같지만 이름이 다름 → 우연 일치, 유지
        {"일자": date(2026, 9, 24), "조정입금": 0.0, "조정지출": 2_000_000.0,
         "내용": "정기지출 추정(자동 초안): 세이브더칠드런 신뢰도 중"},
        # 사용자가 직접 넣은 조정은 건드리지 않는다
        {"일자": date(2026, 9, 21), "조정입금": 0.0, "조정지출": 220_000.0,
         "내용": "수동 입력 지출"},
    ]
    plans = [
        {"자금계획 반영일": date(2026, 9, 21), "예상금액": 220_000.0,
         "거래처": "SK브로드밴드", "지출내용": "인터넷 요금"},
        {"자금계획 반영일": date(2026, 9, 21), "예상금액": 1_996_000.0,
         "거래처": "트라이앵글하모니", "지출내용": "용역비"},
    ]
    kept, skipped = fe.filter_duplicate_adjustments(adjustments, plans)
    assert len(skipped) == 1 and "SKB" in skipped[0]["내용"]
    assert {k["내용"] for k in kept} == {
        "정기지출 추정(자동 초안): 세이브더칠드런 신뢰도 중", "수동 입력 지출"}


def test_정기지출분류_시트_왕복(tmp_path):
    from openpyxl import Workbook, load_workbook
    p = tmp_path / "기준.xlsx"
    wb = Workbook()
    wb.active.title = "카드결제기준"
    wb.save(p)
    wb.close()

    items = [{"정기지출명": "콜마비앤에이치", "분류": "외상대"},
             {"정기지출명": "SKB", "분류": "통신비"}]
    fe.apply_recurring_overrides(items, fe.load_recurring_overrides(p))
    assert items[0]["성격"] == "변동"   # 분류에 '외상' 포함 → 기본 변동
    assert items[1]["성격"] == "정기"

    assert fe.ensure_recurring_override_sheet(p, items) == 2

    # 사용자가 SKB의 성격을 '제외'로 수정하면 다음 실행에 반영된다
    wb = load_workbook(p)
    ws = wb[fe.OVERRIDE_SHEET_NAME]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "SKB":
            row[2].value = "제외"
    wb.save(p)
    wb.close()

    fe.apply_recurring_overrides(items, fe.load_recurring_overrides(p))
    assert items[1]["성격"] == "제외"
    # 이미 있는 항목은 다시 추가하지 않는다 (사용자 수정 보존)
    assert fe.ensure_recurring_override_sheet(p, items) == 0


def test_정기지출분석_수정_수확_반영(tmp_path):
    """결과 파일의 분류·성격 수정이 기준파일 정기지출분류로 흘러간다."""
    from openpyxl import Workbook

    # 사용자가 편집한 라이브 결과물 흉내
    live = tmp_path / "주간자금계획_라이브_20260921_0910.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "정기지출분석"
    for c, h in enumerate(["은행", "정기지출명", "분류"], start=1):
        ws.cell(row=5, column=c, value=h)
    ws.cell(row=5, column=11, value="성격")
    ws.cell(row=6, column=2, value="기업콜마비앤에이치")
    ws.cell(row=6, column=3, value="외상대 지급")      # 사용자가 입력
    ws.cell(row=6, column=11, value="변동")            # 사용자가 입력
    ws.cell(row=7, column=2, value="세이브더칠드런")
    ws.cell(row=7, column=3, value="성격확인필요")      # 미입력 → 무시
    wb.save(live)
    wb.close()

    edits = fe.harvest_recurring_edits(live)
    assert edits["기업콜마비앤에이치"] == {"분류": "외상대 지급", "성격": "변동"}
    assert "세이브더칠드런" not in edits

    # 기준파일에 반영 → 다음 실행의 오버라이드로 적용
    base = tmp_path / "기준.xlsx"
    wb = Workbook()
    wb.active.title = "카드결제기준"
    wb.save(base)
    wb.close()
    assert fe.update_override_sheet(base, edits) >= 1
    ov = fe.load_recurring_overrides(base)
    items = [{"정기지출명": "기업콜마비앤에이치", "분류": ""}]
    fe.apply_recurring_overrides(items, ov)
    assert items[0]["분류"] == "외상대 지급"
    assert items[0]["성격"] == "변동"
    # 같은 값 재수확은 변경 0건 (사용자 수정 보존)
    assert fe.update_override_sheet(base, edits) == 0


def test_정기지출분석_검토파일과_일괄변경_왕복(tmp_path):
    """확인 파일의 정기지출분석 시트: 필터·드롭다운·K4 전체 일괄·N열
    분류별 일괄이 붙고, 수정이 수확돼 기준파일 정기지출분류로 흘러간다.
    우선순위: 전체 일괄 > 분류별 일괄 > 개별 행. '비정기'='변동'."""
    from openpyxl import load_workbook
    import excel_report as er

    assert er.recurring_review_name("확인필요_20260921_0910.xlsx") \
        == "정기지출분석_20260921_0910.xlsx"
    out = tmp_path / "확인필요_20260921_0910.xlsx"
    recurring = [
        {"은행": "우리은행", "정기지출명": "SKB", "분류": "통신비",
         "발생개월수": 6, "거래건수": 6, "평균 월지출": 220000.0,
         "대표 지급일": 25, "신뢰도": "상", "성격": "정기"},
        {"은행": "우리은행", "정기지출명": "KT", "분류": "통신비",
         "발생개월수": 6, "거래건수": 6, "평균 월지출": 90000.0,
         "대표 지급일": 26, "신뢰도": "상", "성격": "정기"},
        {"은행": "농협", "정기지출명": "코웨이", "분류": "렌탈료",
         "발생개월수": 6, "거래건수": 6, "평균 월지출": 113398.0,
         "대표 지급일": 28, "신뢰도": "상", "성격": "변동"},
    ]
    er.create_issue_workbook([], out, week_key="2026-W39", signature="sig",
                             recurring=recurring)

    wb = load_workbook(out)
    # 시트 순서: 안내(컨트롤) → 확인필요 → 정기지출분석
    assert wb.sheetnames[0] == "안내"
    assert wb.sheetnames[-1] == "정기지출분석"
    ws = wb["정기지출분석"]
    # 머리글(5행) 필터 + K4 전체 일괄 + 분류별 일괄 블록(M·N열)
    assert ws.auto_filter.ref == "A5:K8"
    assert ws["K4"].value == "변경 안 함"
    # 성격 칸은 일괄 컨트롤과 수식으로 연결 (엑셀에서 즉시 반영되어 보임)
    assert str(ws["K6"].value).startswith("=IF($K$4")
    assert "MATCH($C6" in str(ws["K6"].value)
    assert '"비정기"' in str(ws["K8"].value)          # 변동은 비정기로 표시
    assert ws["M5"].value == "분류별 일괄"
    assert [ws[f"M{r}"].value for r in (6, 7)] == ["통신비", "렌탈료"]
    # N열은 각 분류의 '현재 성격'을 보여준다 (지난 적용 상태 유지)
    assert ws["N6"].value == "정기"                   # 통신비: 둘 다 정기
    assert ws["N7"].value == "비정기"                 # 렌탈료: 변동 1건
    assert ws.column_dimensions["O"].hidden           # 기준값 열은 숨김
    # '적용할 성격' 선택지 안내 상자 (P열)
    assert ws["P5"].value == "적용할 성격 안내"
    guide = " ".join(str(ws[f"P{r}"].value) for r in range(6, 11))
    for word in ("혼합", "정기 —", "비정기 —", "제외 —", "우선순위"):
        assert word in guide
    formulas = [str(dv.formula1) for dv in ws.data_validations.dataValidation]
    assert any("전체 비정기" in f for f in formulas)
    assert any("정기,비정기,제외" in f for f in formulas)
    # 전체 일괄이 켜지면 나타나는 실시간 경고 (L4)
    assert str(ws["L4"].value).startswith('=IF($K$4=')
    assert "무시됩니다" in str(ws["L4"].value)
    # 개별 수정: SKB만 '비정기'로 직접 입력 (수식 대신 값)
    ws["K6"] = "비정기"
    wb.save(out)
    wb.close()
    edits = fe.harvest_recurring_edits(out)
    assert edits["SKB"]["성격"] == "변동"             # 비정기 → 내부 변동
    # 수식 그대로인 행은 '변경 없음'으로 읽힌다 (기준파일 값 유지)
    assert edits["KT"]["성격"] == ""

    # 분류별 일괄: 통신비 전체를 '제외' → 개별 값보다 우선
    wb = load_workbook(out)
    wb["정기지출분석"]["N6"] = "제외"
    wb.save(out)
    wb.close()
    edits = fe.harvest_recurring_edits(out)
    assert edits["SKB"]["성격"] == "제외"
    assert edits["KT"]["성격"] == "제외"
    assert edits["코웨이"]["성격"] == ""              # 다른 분류는 그대로

    # 전체 일괄(K4)은 분류별보다도 우선한다
    wb = load_workbook(out)
    wb["정기지출분석"]["K4"] = "전체 비정기"
    wb.save(out)
    wb.close()
    edits = fe.harvest_recurring_edits(out)
    assert all(e["성격"] == "변동" for e in edits.values())

    # 기준파일 반영 → 다음 오버라이드로 적용
    from openpyxl import Workbook
    base = tmp_path / "기준.xlsx"
    nb = Workbook()
    nb.active.title = "카드결제기준"
    nb.save(base)
    nb.close()
    assert fe.update_override_sheet(base, edits) >= 2
    fe.apply_recurring_overrides(recurring, fe.load_recurring_overrides(base))
    assert all(i["성격"] == "변동" for i in recurring)


def test_자동추정_시트_수정_원본목록_역반영(tmp_path):
    """확인 파일의 자동추정_지출목록 시트에서 고친 반영/제외·일괄 설정이
    원본 목록 파일(자동추정_지출목록.xlsx)에 되쓰인다."""
    from datetime import date as _date
    from openpyxl import load_workbook
    import excel_report as er

    draft = tmp_path / "자동추정_지출목록.xlsx"
    recur = [{"정기지출명": "SKB", "대표 지급일": 25,
              "평균 월지출": 220_000.0, "신뢰도": "상"},
             {"정기지출명": "코웨이", "대표 지급일": 28,
              "평균 월지출": 113_398.0, "신뢰도": "상"}]
    fe.refresh_auto_draft_file(draft, recur, _date(2026, 9, 21))

    table = fe.read_draft_table(draft)
    assert table["mode"] == "개별 관리" and len(table["rows"]) == 2

    out = tmp_path / "확인필요_20260921_0910.xlsx"
    er.create_issue_workbook([], out, week_key="2026-W39", signature="sig",
                             draft_table=table)
    wb = load_workbook(out)
    ws = wb["자동추정_지출목록"]
    assert ws["B2"].value == "개별 관리"          # 일괄 설정 드롭다운
    assert ws["B5"].value in ("SKB", "코웨이")    # 자료 5행~
    # 사용자가 SKB를 '제외'로, 일괄 설정을 '전체 제외'로 변경
    for r in range(5, 7):
        if ws.cell(row=r, column=2).value == "SKB":
            ws.cell(row=r, column=5, value="제외")
    ws["B2"] = "전체 제외"
    wb.save(out)
    wb.close()

    assert fe.apply_review_draft_edits(out, draft) >= 2
    # 같은 내용 재반영은 0건 (변경 없음)
    assert fe.apply_review_draft_edits(out, draft) == 0
    assert fe.read_draft_table(draft)["mode"] == "전체 제외"
    # 다음 갱신 때 일괄 모드가 전 행에 적용된다
    fe.refresh_auto_draft_file(draft, recur, _date(2026, 9, 21))
    drafts, excluded = fe.load_auto_drafts(draft, _date(2026, 9, 21))
    assert not drafts and excluded == 2


def test_새_항목은_주황색_강조와_안내멘트(tmp_path):
    """지난 확인 파일에 없던 항목은 주황색 행으로 강조되고,
    안내 시트에 '확인해 달라'는 멘트가 나타난다."""
    from datetime import date as _date
    from openpyxl import load_workbook
    import excel_report as er

    def _rec(name, cls):
        return {"은행": "우리", "정기지출명": name, "분류": cls,
                "발생개월수": 6, "거래건수": 6, "평균 월지출": 100000.0,
                "대표 지급일": 25, "신뢰도": "상", "성격": "정기"}

    issues1 = [{"구분": "계획 없는 실제출금", "일자": _date(2026, 9, 22),
                "내용": "사무용품", "금액": 10000.0}]
    recurring1 = [_rec("SKB", "통신비")]
    draft1 = {"mode": "개별 관리",
              "rows": [(_date(2026, 9, 25), "SKB", 220000.0, "상",
                        "제외", "")]}
    prev = tmp_path / "확인필요_20260921_0900.xlsx"
    er.create_issue_workbook(issues1, prev, week_key="2026-W39",
                             signature="s", recurring=recurring1,
                             draft_table=draft1)
    snap = er.review_snapshot(prev)
    assert snap and snap["recurring"] == {"SKB"}

    issues2 = issues1 + [{"구분": "정기지출 누락 의심",
                          "일자": _date(2026, 9, 26), "내용": "KT 통신",
                          "금액": 90000.0, "지시항목": "KT"}]
    recurring2 = recurring1 + [_rec("KT", "통신비")]
    draft2 = {"mode": "개별 관리",
              "rows": draft1["rows"]
              + [(_date(2026, 9, 28), "코웨이", 113398.0, "상",
                  "반영", "")]}
    marks = er.diff_new_items(issues2, recurring2, draft2, snap)
    assert marks["count"] == 3

    out = tmp_path / "확인필요_20260921_0910.xlsx"
    er.create_issue_workbook(issues2, out, week_key="2026-W39",
                             signature="s", recurring=recurring2,
                             draft_table=draft2, new_marks=marks)
    wb = load_workbook(out)
    assert "새로 생긴 항목이 3건" in str(wb["안내"]["A3"].value)
    def _fill(ws, r):
        return str(ws.cell(row=r, column=1).fill.start_color.rgb or "")
    assert _fill(wb["확인필요"], 6).endswith("FFE699")      # 새 항목 행
    assert not _fill(wb["확인필요"], 5).endswith("FFE699")  # 기존 항목
    assert _fill(wb["정기지출분석"], 7).endswith("FFE699")  # KT
    assert _fill(wb["자동추정_지출목록"], 6).endswith("FFE699")  # 코웨이
    wb.close()

    # 변화가 없으면 '그대로'라는 안내가 뜬다
    marks0 = er.diff_new_items(issues1, recurring1, draft1, snap)
    assert marks0["count"] == 0
    same = tmp_path / "확인필요_20260921_0920.xlsx"
    er.create_issue_workbook(issues1, same, week_key="2026-W39",
                             signature="s", recurring=recurring1,
                             draft_table=draft1, new_marks=marks0)
    wb = load_workbook(same)
    assert "그대로" in str(wb["안내"]["A3"].value)
    wb.close()


def test_계좌별_시나리오_인출_우선순위():
    """우리은행 부족분은 농협→국민 순으로 이체해 채운다."""
    daily = [{"일자": date(2026, 9, 21), "요일": "월", "실적": False,
              "온라인 예상입금": 0.0, "확정·기타입금": 0.0,
              "팀별 송금예정": 0.0, "카드결제": 0.0, "자동이체": 0.0,
              "기타지출": 300.0}]
    balances = {("우리은행", "W"): 100.0, ("농협", "N"): 150.0,
                ("국민은행", "K"): 1000.0}
    sc = fe.build_account_scenario(daily, balances, [])
    assert sc["accounts"][0] == ("우리은행", "W")
    row = sc["rows"][0]
    # 부족 200 = 농협 150 + 국민 50
    assert row["이체"] == {("농협", "N"): 150.0, ("국민은행", "K"): 50.0}
    assert round(row["잔액"][("우리은행", "W")]) == 0
    assert round(row["잔액"][("농협", "N")]) == 0
    assert round(row["잔액"][("국민은행", "K")]) == 950
    assert row["지출"] == 300.0          # 당일 지출도 행에 기록
    assert row["비고"] == ""

    # 전 계좌 소진 시 부족 경고
    daily2 = [dict(daily[0], 기타지출=2000.0)]
    sc2 = fe.build_account_scenario(daily2, balances, [])
    assert "부족" in sc2["rows"][0]["비고"]

    # 예상입금은 최근 외부입금 비중대로 나뉘고, 배분액이 행에 기록된다
    hist = [{"은행": "우리은행", "계좌": "W", "입금액": 300.0},
            {"은행": "국민은행", "계좌": "K", "입금액": 700.0},
            {"은행": "국민은행", "계좌": "K", "입금액": 500.0,
             "내부이체": True}]                     # 내부이체는 비중 제외
    daily3 = [dict(daily[0], **{"온라인 예상입금": 1000.0, "기타지출": 0.0})]
    sc3 = fe.build_account_scenario(daily3, balances, hist)
    assert sc3["shares"][("우리은행", "W")] == 0.3
    assert sc3["shares"][("농협", "N")] == 0.0
    assert sc3["opening"][("우리은행", "W")] == 100.0   # 수식 출발 잔액
    row3 = sc3["rows"][0]
    assert round(row3["입금"][("우리은행", "W")]) == 300
    assert round(row3["입금"][("국민은행", "K")]) == 700
    assert round(row3["잔액"][("국민은행", "K")]) == 1700


def test_에이팜_판별과_지출예정_수집():
    """㈜에이팜 관련만 골라내고 자사명(에이팜건강)은 제외한다."""
    assert fe.mentions_apalm("㈜에이팜")
    assert fe.mentions_apalm("", "에이팜 외상대 지급")
    assert not fe.mentions_apalm("(주)에이팜건강")
    assert not fe.mentions_apalm("에이팜 건강 급여")   # 공백 변형도 자사
    assert not fe.mentions_apalm("콜마비앤에이치")

    masked = [
        {"반영상태": "정상반영", "자금계획 반영일": date(2026, 9, 23),
         "팀명": "경영지원팀", "거래처": "㈜에이팜", "지출내용": "외상대 지급",
         "예상금액": 5_000_000.0, "지급방법": "계좌송금"},
        {"반영상태": "정상반영", "자금계획 반영일": date(2026, 9, 22),
         "팀명": "디자인팀", "거래처": "OO인쇄", "지출내용": "리플렛",
         "예상금액": 25_000.0, "지급방법": "법인카드"},
        # 대외비 집계행은 수집하지 않는다
        {"반영상태": "정상반영", "confidential": True,
         "자금계획 반영일": date(2026, 9, 23), "거래처": "㈜에이팜",
         "지출내용": "(대외비 집계)", "예상금액": 1.0},
    ]
    adjustments = [
        {"일자": date(2026, 9, 25), "조정입금": 0.0, "조정지출": 3_000_000.0,
         "내용": "정기지출 추정(자동 초안): ㈜에이팜 신뢰도 상"},
        {"일자": date(2026, 9, 25), "조정입금": 0.0, "조정지출": 100_000.0,
         "내용": "정기지출 추정(자동 초안): SKB 신뢰도 상"},
    ]
    rows = fe.collect_apalm_expenses(masked, adjustments)
    assert [(r["출처"], r["거래처"], r["반영"]) for r in rows] == [
        ("팀 지출계획", "㈜에이팜", "반영"),
        ("정기지출 추정", "㈜에이팜", "반영")]
    assert rows[0]["예상금액"] == 5_000_000.0
    assert rows[1]["예상금액"] == 3_000_000.0


def test_비고_에이팜_표시는_자금계획에서_제외():
    """비고에 '에이팜'이 적힌 지출계획은 집계에서 빠지고 별도관리된다."""
    rows = [
        {"반영상태": "정상반영", "자금계획 반영일": date(2026, 9, 30),
         "팀명": "경영지원팀", "거래처": "서원회계법인", "지출내용": "기장료",
         "예상금액": 330_000.0, "지급방법": "계좌송금", "비고": "에이팜"},
        {"반영상태": "정상반영", "자금계획 반영일": date(2026, 9, 21),
         "팀명": "건강사업팀", "거래처": "네이버SA&GFA", "지출내용": "광고비",
         "예상금액": 3_498_000.0, "지급방법": "계좌송금",
         "비고": "화/수 2일치 금액"},
    ]
    plan = {"countable": list(rows), "integrated": rows}
    marked = fe.split_apalm_marked(plan)
    assert [r["거래처"] for r in marked] == ["서원회계법인"]
    assert marked[0]["반영상태"] == fe.APALM_EXCLUDED_STATUS
    assert [r["거래처"] for r in plan["countable"]] == ["네이버SA&GFA"]
    # 취합(integrated)에는 별도관리 상태로 남는다
    assert rows[0]["반영상태"] == fe.APALM_EXCLUDED_STATUS

    # 별도관리 건은 원본 상세 그대로 '미반영'으로 수집된다
    out = fe.collect_apalm_expenses(plan["countable"], [], marked)
    assert len(out) == 1
    assert out[0]["반영"] == "미반영(별도 관리)"
    assert out[0]["거래처"] == "서원회계법인"
    assert out[0]["비고"] == "에이팜"

    # 경영지원팀(대외비) 행도 상세를 싣되, 급여·세금보험류 분류만 가린다
    conf = [{"반영상태": fe.APALM_EXCLUDED_STATUS, "confidential": True,
             "자금계획 반영일": date(2026, 9, 30), "팀명": "경영지원팀",
             "거래처": "가가사무기", "지출내용": "복합기 임대",
             "예상금액": 280_000.0, "대외비구분": "", "비고": "에이팜"},
            {"반영상태": fe.APALM_EXCLUDED_STATUS, "confidential": True,
             "자금계획 반영일": date(2026, 9, 25), "팀명": "경영지원팀",
             "거래처": "메트라이프", "지출내용": "정기보험",
             "예상금액": 2_497_000.0, "대외비구분": "세금·보험(대외비)",
             "비고": "에이팜"}]
    out = fe.collect_apalm_expenses([], [], conf)
    by_amt = {r["예상금액"]: r for r in out}
    assert by_amt[280_000.0]["거래처"] == "가가사무기"      # 일반 내용은 노출
    assert by_amt[2_497_000.0]["거래처"] == "(대외비)"      # 민감 분류는 가림
    assert by_amt[2_497_000.0]["지출내용"] == "세금·보험(대외비)"


def test_비고에서_에이팜_제외():
    """4주일별계획 비고에는 에이팜 상세를 넣지 않는다."""
    import app as app_module

    daily = [{"일자": date(2026, 9, 23), "실적": False,
              "비고": "정기지출 추정(자동 초안): ㈜에이팜 신뢰도 상; "
                    "정기지출 추정(자동 초안): SKB 신뢰도 상"}]
    masked = [
        {"반영상태": "정상반영", "자금계획 반영일": date(2026, 9, 23),
         "팀명": "경영지원팀", "거래처": "㈜에이팜", "지출내용": "외상대 지급",
         "예상금액": 5_000_000.0},
        {"반영상태": "정상반영", "자금계획 반영일": date(2026, 9, 23),
         "팀명": "디자인팀", "거래처": "OO인쇄", "지출내용": "리플렛",
         "예상금액": 25_000.0},
    ]
    app_module._annotate_daily_notes(daily, masked)
    note = daily[0]["비고"]
    assert "에이팜" not in note
    assert "OO인쇄" in note and "SKB" in note


def test_요일평균은_최근주에_가중치():
    """새로 올린 주의 실적이 예상 입금액에 더 크게 반영된다."""
    base = date(2026, 9, 21)  # 월요일
    history = [
        # 1주 전 월요일: 2백만 / 4주 전 월요일: 1백만
        _tx(date(2026, 9, 14), in_amt=2_000_000, cls="온라인매출입금"),
        _tx(date(2026, 8, 24), in_amt=1_000_000, cls="온라인매출입금"),
    ]
    plain = fe.weekday_online_averages(history, base, 12, recency_halflife=0)
    weighted = fe.weekday_online_averages(history, base, 12)
    # 단순 평균: 3백만 / 월요일 4회 = 75만
    assert round(plain[0]) == 750_000
    # 가중 평균은 최근 주(2백만) 쪽으로 끌려 올라간다
    assert weighted[0] > plain[0]
    # 다른 요일은 실적이 없으므로 0
    assert weighted[1] == 0.0


def test_성격_변동이면_자동초안_추정_제외():
    """외상매입금처럼 금액이 변하는 항목은 자동 초안 추정을 쓰지 않는다."""
    adjustments = [
        {"일자": date(2026, 9, 28), "조정입금": 0.0, "조정지출": 1_369_025.0,
         "내용": "정기지출 추정(자동 초안): 국민주식회사 에이팜 신뢰도 하"},
        {"일자": date(2026, 9, 30), "조정입금": 0.0, "조정지출": 100_000.0,
         "내용": "정기지출 추정(자동 초안): SKB 신뢰도 상"},
        # 입금 추정과 수동 조정은 건드리지 않는다
        {"일자": date(2026, 9, 30), "조정입금": 10_021_867.0, "조정지출": 0.0,
         "내용": "관계사 외상대 입금 추정(주식회사에이팜, 말일)"},
    ]
    overrides = {"국민주식회사 에이팜": {"분류": "외상매입금", "성격": "변동"},
                 "SKB": {"분류": "통신비", "성격": "정기"}}
    kept, dropped = fe.filter_adjustments_by_overrides(adjustments, overrides)
    assert len(dropped) == 1 and "에이팜" in dropped[0]["내용"]
    assert {a["내용"][:2] for a in kept} == {"정기", "관계"}

    # 성격을 '정기'로 되돌리면 다시 반영된다
    overrides["국민주식회사 에이팜"]["성격"] = "정기"
    kept2, dropped2 = fe.filter_adjustments_by_overrides(
        adjustments, overrides)
    assert not dropped2 and len(kept2) == 3


def test_금주일별_시나리오_생성():
    base = date(2026, 9, 21)  # 월요일
    fc = fe.build_forecast([], base, 1_000_000, [], [], [], [0.8, 0.9], 0.8)
    days = fc["rate_scenarios"][0.9]["금주일별"]
    assert len(days) == 5
    assert days[0][0] == base and days[-1][0] == date(2026, 9, 25)


def test_당일_행은_실제_입출금만으로_마감된다():
    """실행일 행 = 은행 파일의 실제 입출금(실적 마감), 기말 = 오늘
    실잔고. 실제는 today_actual로도 담겨 출발 행(7행)에 표시된다."""
    base = date(2026, 9, 21)
    history = [
        {"거래일": base, "입금액": 10_373_905.0, "출금액": 0.0,
         "자동분류": "온라인매출입금", "내부이체": False},
        {"거래일": base, "입금액": 20_110_000.0, "출금액": 0.0,
         "자동분류": "", "내부이체": False},
        {"거래일": base, "입금액": 0.0, "출금액": 3_468_293.0,
         "자동분류": "", "내부이체": False},
    ]
    plans = [{"자금계획 반영일": base, "예상금액": 3_498_000.0,
              "지급방법": "계좌송금", "확정여부": "확정"}]
    fc = fe.build_forecast(plans, base, 161_670_332, history, [], [],
                           [0.8], 0.8, run_date=base)
    ta = fc["today_actual"]
    assert round(ta["온라인"]) == 10_373_905
    assert round(ta["기타입금"]) == 20_110_000
    assert round(ta["출금"]) == 3_468_293
    assert round(ta["순증감"]) == 27_015_612
    # 실행일 행: 실제 값으로 마감 (출금 3,468,293은 계획 3,498,000과
    # 대조되어 집행으로 표시), 기말 = 현재 실잔고
    d21 = fc["daily"][0]
    assert round(d21["온라인 예상입금"]) == 10_373_905
    assert round(d21["확정·기타입금"]) == 20_110_000
    assert round(d21["팀별 송금예정"]) == 3_468_293   # 실제 집행액
    assert round(d21["기초잔액"]) == 161_670_332 - 27_015_612
    assert round(d21["기말잔액"]) == 161_670_332
    # 잔여 29,707은 익일로 이월
    assert round(fc["daily"][1]["팀별 송금예정"]) == 3_498_000 - 3_468_293


def test_13주에_정기지출_추정을_반영하지_않는다():
    """정기지출 추정은 확인용일 뿐 계획에 넣지 않는다 (2026-09-21)."""
    base = date(2026, 9, 21)
    recurring = [{"분류": "카드", "평균 월지출": 5_000_000.0,
                  "대표 지급일": 25}]
    fc = fe.build_forecast([], base, 10_000_000, [], [], recurring,
                           [0.8], 0.8, run_date=base)
    for w in fc["weekly"][4:]:
        assert round(w.get("카드결제") or 0) == 0
        assert round(w.get("기타지출") or 0) == 0


def test_보류_선택한_당일_예정은_계획에서_뺀다():
    """확인 단계에서 '보류(제외)'로 고른 항목의 잔여는 오늘 계획에서
    빠진다 (2026-09-21 사용자 요청 — 차이는 사용자가 결정)."""
    base = date(2026, 9, 21)
    history = [{"거래일": base, "입금액": 0.0, "출금액": 2_508_000.0,
                "자동분류": "", "내부이체": False,
                "기재내용·상대방": "국민네이버 apha"}]
    plans = [
        {"요청ID": "마케팅-001", "자금계획 반영일": base,
         "예상금액": 3_498_000.0, "거래처": "네이버SA&GFA",
         "지급방법": "계좌송금", "확정여부": "확정"},
        {"요청ID": "물류-002", "자금계획 반영일": base,
         "예상금액": 1_996_000.0, "거래처": "트라이앵글하모니",
         "지급방법": "계좌송금", "확정여부": "확정"},
    ]
    # 보류 없음: 실행일 행은 집행 2,508,000만, 잔여는 익일(9/22) 이월
    fc = fe.build_forecast(plans, base, 10_000_000, history, [], [],
                           [0.8], 0.8, run_date=base)
    assert round(fc["daily"][0]["팀별 송금예정"]) == 2_508_000
    assert round(fc["daily"][1]["팀별 송금예정"]) == 990_000 + 1_996_000
    # 트라이앵글 보류: 잔여 1,996,000이 이월에서 빠진다
    fc2 = fe.build_forecast(plans, base, 10_000_000, history, [], [],
                            [0.8], 0.8, run_date=base,
                            intraday_holds={"물류-002"})
    assert round(fc2["daily"][1]["팀별 송금예정"]) == 990_000
    assert "보류 제외" in fc2["daily"][0]["비고"]
    # 대조내역: 네이버 일부지급 / 트라이앵글 미집행(보류)
    det = {d["요청ID"]: d for d in fc2["intraday"]["대조내역"]}
    assert det["마케팅-001"]["상태"] == "일부지급"
    assert det["물류-002"]["상태"] == "미집행"
    assert det["물류-002"]["보류"] is True


def test_확인파일_보류_선택_왕복(tmp_path):
    """확인필요 파일의 '당일 지출·실제 차이' 행에서 '보류(제외)'를
    고르면 load_intraday_holds가 요청ID를 돌려준다."""
    import excel_report

    issues = [
        {"구분": excel_report.ISSUE_INTRADAY, "팀명": "물류팀",
         "요청ID": "물류-002", "일자": date(2026, 9, 21),
         "내용": "트라이앵글하모니 — 미집행", "금액": 1_996_000,
         "원본파일": "", "당일지시": True},
        {"구분": excel_report.ISSUE_INTRADAY, "팀명": "마케팅팀",
         "요청ID": "마케팅-001", "일자": date(2026, 9, 21),
         "내용": "네이버SA&GFA — 일부지급", "금액": 990_000,
         "원본파일": "", "당일지시": True},
        {"구분": "정기지출 누락 의심", "일자": date(2026, 9, 22),
         "내용": "메트라이프", "금액": 4_900_000, "원본파일": "",
         "지시항목": "메트라이프"},
    ]
    path = tmp_path / "확인필요_test.xlsx"
    excel_report.create_issue_workbook(issues, path, week_key="2026-W39",
                                       signature="sig")
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[excel_report.REVIEW_SHEET]
    # 당일 행 기본값은 '지출 예정(유지)'
    assert ws.cell(row=5, column=9).value == excel_report.ACTION_KEEP
    ws.cell(row=5, column=9).value = excel_report.ACTION_HOLD   # 물류-002 보류
    wb.save(path)
    wb.close()
    assert excel_report.load_intraday_holds(path) == {"물류-002"}
