# -*- coding: utf-8 -*-
"""'처리' 열 드롭다운·안내 말풍선 회귀 테스트 (2026-10-04 사용자 요청).

계획 없는 실제출금의 처리 칸에 뭘 입력할지 알 수 없던 문제 →
기본 분류 + 규칙 파일의 기존 분류를 드롭다운으로 제공하고(직접 입력
허용), 셀 선택 시 안내 말풍선, 머리글에 사용법 메모를 붙인다.
"""
from openpyxl import load_workbook

import excel_report
from excel_report import (ACTION_HOLD, ACTION_KEEP, DEFAULT_CATEGORY_CHOICES,
                         ISSUE_INTRADAY, ISSUE_UNPLANNED)


def _make(tmp_path, **kw):
    issues = [
        {"구분": ISSUE_UNPLANNED, "팀명": "", "은행": "우리은행",
         "내용": "SKB6486358411 통신요금", "금액": 85621},
        {"구분": ISSUE_INTRADAY, "팀명": "경영", "내용": "당일 차이",
         "금액": 10000, "당일지시": True, "요청ID": "경영-20261004-001"},
    ]
    out = excel_report.create_issue_workbook(
        issues, tmp_path / "확인필요_test.xlsx",
        week_key="2026-W41", signature="sig", **kw)
    return load_workbook(out)["확인필요"]


def _dv_for(ws, coord):
    for dv in ws.data_validations.dataValidation:
        if any(coord in rng for rng in dv.sqref.ranges):
            return dv
    return None


def test_계획없는출금_처리칸은_분류_드롭다운(tmp_path):
    ws = _make(tmp_path, category_choices=["네이버SA&GFA", "4대보험"])
    dv = _dv_for(ws, "I5")                       # 첫 자료 행(계획없는출금)
    assert dv is not None and dv.type == "list"
    assert dv.showErrorMessage is False          # 직접 입력 허용
    assert dv.showInputMessage and dv.prompt     # 안내 말풍선
    # 숨김 L열에 기본 후보 + 규칙 분류(중복 제거) 순서로 들어간다
    col = [ws.cell(row=r, column=12).value
           for r in range(2, 2 + len(DEFAULT_CATEGORY_CHOICES) + 1)]
    assert col[:len(DEFAULT_CATEGORY_CHOICES)] == DEFAULT_CATEGORY_CHOICES
    assert "네이버SA&GFA" in col                  # 4대보험은 기본과 중복
    assert ws.column_dimensions["L"].hidden


def test_당일차이_처리칸은_유지보류_드롭다운과_말풍선(tmp_path):
    ws = _make(tmp_path)
    dv = _dv_for(ws, "I6")                       # 둘째 자료 행(당일 차이)
    assert dv is not None
    assert ACTION_KEEP in dv.formula1 and ACTION_HOLD in dv.formula1
    assert dv.showInputMessage and "보류" in dv.prompt


def test_처리_머리글에_사용법_메모(tmp_path):
    ws = _make(tmp_path)
    head = ws.cell(row=4, column=9)              # I4: '처리' 머리글
    assert head.comment is not None
    assert "분류" in head.comment.text
