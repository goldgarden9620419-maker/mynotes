# -*- coding: utf-8 -*-
"""pytest 공용 fixture. 실제 급여·개인정보 없는 가상자료만 사용한다."""
from __future__ import annotations

import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook

from common import REQUIRED_TEAM_COLUMNS, TEAMS, week_monday
from config import Config


def write_team_file(path: Path, rows: list[dict],
                    extra_columns: list[str] | None = None) -> Path:
    """테스트용 팀 지출계획 파일(수식 없이 값으로) 생성."""
    columns = list(REQUIRED_TEAM_COLUMNS) + (extra_columns or [])
    wb = Workbook()
    ws = wb.active
    ws.title = "지출계획"
    for c, name in enumerate(columns, start=1):
        ws.cell(row=1, column=c, value=name)
    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(columns, start=1):
            ws.cell(row=r, column=c, value=row.get(name))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def team_row(code: str, name: str, serial: int, reg: date, due: date,
             vendor: str, desc: str, amount: float,
             method: str = "계좌송금", card: str = "",
             confirmed: str = "확정", progress: str = "신규",
             modified: date | None = None, **extra) -> dict:
    row = {
        "요청ID": f"{code}-{reg.strftime('%Y%m%d')}-{serial:03d}",
        "팀코드": code, "팀명": name, "최초등록일": reg, "일련번호": serial,
        "지급예정일": due, "거래처": vendor, "지출내용": desc,
        "예상금액": amount, "지급방법": method, "카드구분": card,
        "확정여부": confirmed, "진행상태": progress,
        "최종수정일": modified or reg, "비고": "",
    }
    row.update(extra)
    return row


def write_bank_csv(path: Path, bank: str, rows: list[dict]) -> Path:
    """은행별 실제 다운로드 형식과 유사한 CSV 생성."""
    headers = {
        "농협": ["거래일자", "거래시간", "출금금액", "입금금액", "거래후잔액",
               "거래내용", "거래기록사항", "거래점"],
        "우리은행": ["거래일시", "적요", "기재내용", "지급(원)", "입금(원)",
                 "거래후잔액", "취급점"],
        "국민은행": ["거래일시", "적요", "보낸분/받는분", "출금액(원)",
                 "입금액(원)", "잔액(원)", "처리점"],
    }[bank]
    path.parent.mkdir(parents=True, exist_ok=True)
    import csv
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([row.get(h, "") for h in headers])
    return path


def nh_tx(d: date, t: str, out_amt=0, in_amt=0, balance=0,
          desc="", memo="", branch="본점") -> dict:
    return {"거래일자": d.strftime("%Y-%m-%d"), "거래시간": t,
            "출금금액": out_amt or "", "입금금액": in_amt or "",
            "거래후잔액": balance, "거래내용": desc, "거래기록사항": memo,
            "거래점": branch}


def woori_tx(dt: datetime, summary="", memo="", out_amt=0, in_amt=0,
             balance=0) -> dict:
    return {"거래일시": dt.strftime("%Y-%m-%d %H:%M:%S"), "적요": summary,
            "기재내용": memo, "지급(원)": out_amt or "",
            "입금(원)": in_amt or "", "거래후잔액": balance, "취급점": "본점"}


def kb_tx(dt: datetime, summary="", memo="", out_amt=0, in_amt=0,
          balance=0) -> dict:
    return {"거래일시": dt.strftime("%Y-%m-%d %H:%M:%S"), "적요": summary,
            "보낸분/받는분": memo, "출금액(원)": out_amt or "",
            "입금액(원)": in_amt or "", "잔액(원)": balance, "처리점": "본점"}


@pytest.fixture
def env(tmp_path):
    """폴더 구조 + config + 빈 팀파일이 준비된 실행환경."""
    cfg = Config.load(config_path=tmp_path / "00_프로그램" / "config.yaml",
                      base_dir=tmp_path)
    cfg.ensure_folders()
    return cfg


def fill_all_team_files(cfg, rows_by_team: dict[str, list[dict]]) -> None:
    """모든 팀 파일을 생성한다. rows_by_team에 없는 팀은 빈 파일."""
    for team, path in cfg.team_files():
        rows = rows_by_team.get(team["name"], [])
        extra = ["대외비구분"] if team["confidential"] else None
        write_team_file(path, rows, extra)


def make_full_inputs(cfg, now: datetime) -> dict:
    """정상 실행이 가능한 최소 입력자료 세트를 만든다."""
    base = week_monday(now.date())
    reg = base - timedelta(days=3)  # 지난주 금요일 등록
    rows_by_team = {
        "물류팀": [
            team_row("물류", "물류팀", 1, reg, base + timedelta(days=2),
                     "한진택배", "택배비", 500000),
            team_row("물류", "물류팀", 2, reg, base + timedelta(days=3),
                     "대한통운", "상자 구입", 200000, method="법인카드",
                     card="국민카드"),
        ],
        "건강사업팀": [
            team_row("건강", "건강사업팀", 1, reg, base + timedelta(days=1),
                     "원료상사", "원료 구매", 1200000),
        ],
        "경영지원팀": [
            team_row("경영", "경영지원팀", 1, reg, base + timedelta(days=4),
                     "가상급여처리", "급여 이체", 3000000,
                     대외비구분="급여·인건비(대외비)"),
        ],
    }
    fill_all_team_files(cfg, rows_by_team)

    # 은행자료: 지난주 거래 + 온라인 입금 이력
    prev = base - timedelta(days=7)
    write_bank_csv(cfg.bank_dir("농협") / "농협_거래내역.csv", "농협", [
        nh_tx(prev, "09:10:11", out_amt=300000, balance=52000000,
              desc="이체", memo="사무용품(주)"),
        nh_tx(prev + timedelta(days=1), "14:22:33", in_amt=2500000,
              balance=54500000, desc="입금", memo="네이버페이"),
        nh_tx(prev + timedelta(days=2), "10:00:00", in_amt=1800000,
              balance=56300000, desc="입금", memo="KG이니시스"),
    ])
    write_bank_csv(cfg.bank_dir("우리은행") / "우리_거래내역.csv", "우리은행", [
        woori_tx(datetime.combine(prev, datetime.min.time())
                 + timedelta(days=1, hours=9), "타행이체", "사회보험료",
                 out_amt=450000, balance=23000000),
        woori_tx(datetime.combine(prev, datetime.min.time())
                 + timedelta(days=2, hours=11), "카드", "우리카드결제대금",
                 out_amt=1100000, balance=21900000),
    ])
    write_bank_csv(cfg.bank_dir("국민은행") / "국민_거래내역.csv", "국민은행", [
        kb_tx(datetime.combine(prev, datetime.min.time())
              + timedelta(days=3, hours=10), "이체", "(주)에이팜건강",
              out_amt=5000000, balance=8000000),
        kb_tx(datetime.combine(prev, datetime.min.time())
              + timedelta(days=4, hours=15), "입금", "스마트스토어",
              in_amt=900000, balance=8900000),
    ])
    return rows_by_team
