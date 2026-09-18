# -*- coding: utf-8 -*-
"""취합본 양식(신청자·품의승인 열) 인식 테스트."""
from datetime import date

from openpyxl import Workbook

from common import TEAMS
import team_loader

COLUMNS = ["요청ID", "팀코드", "팀명", "신청자", "최초등록일", "일련번호",
           "지급예정일", "거래처", "지출내용", "예상금액", "지급방법",
           "카드구분", "품의승인", "진행상태", "최종수정일", "비고",
           "대외비구분"]


def _write(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "지출계획"
    for c, name in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=c, value=name)
    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(COLUMNS, start=1):
            ws.cell(row=r, column=c, value=row.get(name))
    wb.save(path)
    wb.close()


def _row(serial, approval, vendor="서울보증보험", amount=632250,
         method="계좌송금"):
    return {"팀코드": "경영", "팀명": "경영지원팀", "신청자": "홍길동",
            "최초등록일": date(2026, 9, 18), "일련번호": serial,
            "지급예정일": date(2026, 9, 23), "거래처": vendor,
            "지출내용": "보증보험 갱신", "예상금액": amount,
            "지급방법": method, "품의승인": approval, "진행상태": "신규",
            "최종수정일": date(2026, 9, 18),
            "대외비구분": "기타 대외비 지출"}


def test_품의승인_열_인식(tmp_path):
    path = tmp_path / "취합본.xlsx"
    _write(path, [_row(1, "승인대기"), _row(2, "승인"), _row(3, "반려")])
    team = next(t for t in TEAMS if t["name"] == "경영지원팀")
    rows, issues = team_loader.load_team_file(path, team)
    assert len(rows) == 3
    # 확정여부 열이 없어도 '필수 열 누락'으로 처리하지 않는다
    assert not any("확정여부" in i.get("내용", "") for i in issues)
    by_serial = {r["일련번호"]: r for r in rows}
    assert by_serial[1]["확정여부"] == "미확정"   # 승인대기
    assert by_serial[2]["확정여부"] == "확정"     # 승인
    assert by_serial[3]["확정여부"] == "미확정"   # 반려
    assert by_serial[1]["신청자"] == "홍길동"
    assert by_serial[1]["요청ID"] == "경영-20260918-001"

    plan = team_loader.build_integrated_plan(rows)
    # 기본(명세) 모드: 승인된 1건만 합계 반영
    assert len(plan["countable"]) == 1
    assert plan["countable"][0]["일련번호"] == 2
    # 승인대기·반려는 확인필요에 표시되고 대외비 거래처는 노출하지 않는다
    pending = [i for i in plan["issues"] if i["구분"] == "미확정(승인대기)"]
    assert len(pending) == 2
    assert all("서울보증보험" not in i["내용"] for i in pending)
    assert all("기타 대외비 지출" in i["내용"] for i in pending)


def test_승인대기_포함_정책(tmp_path):
    """2026-09-18 사용자 결정: 확정여부 무관, 지급일자 기준 반영."""
    path = tmp_path / "취합본.xlsx"
    _write(path, [_row(1, "승인대기"), _row(2, "승인"), _row(3, "반려")])
    team = next(t for t in TEAMS if t["name"] == "경영지원팀")
    rows, _ = team_loader.load_team_file(path, team)
    plan = team_loader.build_integrated_plan(rows, include_unconfirmed=True)
    # 취소 건이 없으므로 3건 전부 반영
    assert len(plan["countable"]) == 3
    assert all(r["반영상태"] == "정상반영" for r in plan["countable"])
    # 반영됐다는 사실은 확인필요에 정보성으로 표시된다
    included = [i for i in plan["issues"] if i["구분"] == "승인대기(반영됨)"]
    assert len(included) == 2
    assert all("서울보증보험" not in i["내용"] for i in included)
