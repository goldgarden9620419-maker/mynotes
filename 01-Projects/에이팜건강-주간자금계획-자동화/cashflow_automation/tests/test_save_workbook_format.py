# -*- coding: utf-8 -*-
"""save_workbook 자동 서식 회귀 테스트 (2026-10-04 사용자 요청).

각 시트에서 긴 텍스트가 셀 안에 안 보이던 문제 → 줄바꿈 + 행 높이
자동 계산, 날짜·숫자가 셀 아래에 붙던 문제 → 전 셀 세로 가운데 정렬.
"""
from datetime import date

from openpyxl import Workbook, load_workbook

from common import save_workbook


def _make_wb():
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.column_dimensions["A"].width = 12
    ws["A1"] = date(2026, 10, 5)
    ws["B1"] = 1234567
    ws["A2"] = "짧은글"
    ws["A3"] = ("이 텍스트는 열 폭 12를 훨씬 넘는 아주 긴 한글 설명 문장이라서 "
                "줄바꿈과 행 높이 조정이 없으면 셀 안에서 보이지 않는다")
    ws["B3"] = "오른쪽 값"                      # 넘침을 막는 셀
    ws["A4"] = "오른쪽이 빈 아주 긴 제목 문장 — 옆으로 흘러넘쳐 보이므로 줄바꿈하면 안 된다"
    return wb


def test_값있는_셀은_전부_세로_가운데_정렬(tmp_path):
    out = tmp_path / "f.xlsx"
    save_workbook(_make_wb(), out)
    ws = load_workbook(out)["요약"]
    for coord in ("A1", "B1", "A2", "A3"):
        assert ws[coord].alignment.vertical == "center", coord


def test_긴_텍스트는_줄바꿈과_행높이_자동_조정(tmp_path):
    out = tmp_path / "f.xlsx"
    save_workbook(_make_wb(), out)
    ws = load_workbook(out)["요약"]
    assert ws["A3"].alignment.wrap_text is True
    assert ws.row_dimensions[3].height is not None
    assert ws.row_dimensions[3].height > 20          # 여러 줄 분량
    # 짧은 행은 높이를 건드리지 않는다
    assert ws.row_dimensions[2].height is None


def test_오른쪽이_비면_넘침_유지_줄바꿈_안함(tmp_path):
    out = tmp_path / "f.xlsx"
    save_workbook(_make_wb(), out)
    ws = load_workbook(out)["요약"]
    assert not ws["A4"].alignment.wrap_text     # 제목·안내문 넘침 유지
    assert ws.row_dimensions[4].height is None
    assert ws["A4"].alignment.vertical == "center"


def test_수식_셀은_줄바꿈과_높이_계산에서_제외(tmp_path):
    wb = _make_wb()
    ws = wb["요약"]
    ws["A5"] = ("=INDEX('설정및분류'!$C$6:$C$999,MATCH(B5,"
                "'설정및분류'!$B$6:$B$999,0))")      # 긴 수식
    out = tmp_path / "f.xlsx"
    save_workbook(wb, out)
    ws2 = load_workbook(out)["요약"]
    assert ws2["A5"].alignment.vertical == "center"  # 정렬은 적용
    assert not ws2["A5"].alignment.wrap_text         # 줄바꿈은 제외
    assert ws2.row_dimensions[5].height is None      # 높이도 그대로


def test_대표보고_시트는_높이_유지_정렬만(tmp_path):
    wb = _make_wb()
    ws = wb.create_sheet("대표보고")
    ws.column_dimensions["A"].width = 10
    ws["A1"] = "대표보고 A4 인쇄 레이아웃 보호를 확인하는 아주 긴 문장 " * 3
    out = tmp_path / "f.xlsx"
    save_workbook(wb, out)
    ws2 = load_workbook(out)["대표보고"]
    assert ws2["A1"].alignment.vertical == "center"
    assert ws2.row_dimensions[1].height is None      # 높이는 그대로
