# -*- coding: utf-8 -*-
"""전체 파이프라인 통합 테스트 (필수파일 누락·재시도·중복실행·재열기 검증)."""
import logging
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

import app as app_module
from common import (
    STATUS_REVIEW_WAIT, STATUS_SUCCESS, STATUS_WAITING_FILES, now_local,
    week_monday,
)
from excel_report import verify_workbook
from management_report import verify_management_workbook
from state_manager import StateManager
from tests.conftest import make_full_inputs, nh_tx, write_bank_csv

NOW = datetime(2026, 9, 21, 9, 15)  # 월요일 09:15
LOG = logging.getLogger("test")


def _run(cfg, state, **kwargs):
    return app_module.run_weekly_job(cfg, state, LOG, now=NOW, **kwargs)


def test_전체_파이프라인_성공(env):
    cfg = env
    make_full_inputs(cfg, NOW)
    state = StateManager(cfg.state_dir)

    result = _run(cfg, state, mode="auto")
    assert result.status == STATUS_SUCCESS, result.message

    out_dir = cfg.folder("output")
    excels = list(out_dir.glob("주간자금계획_*.xlsx"))
    pdfs = list(out_dir.glob("주간자금계획_대표보고_*.pdf"))
    assert len(excels) == 1 and len(pdfs) == 1
    # 확인필요는 05_결과가 아니라 06_확인필요 폴더에만 둔다
    assert not list(out_dir.glob("확인필요_*.xlsx"))
    assert len(list(cfg.folder("review").glob("확인필요_*.xlsx"))) == 1

    # 결과파일 재열기 검증 (31번 항목) — 경영보고(대화형)가 기본 결과물
    assert verify_management_workbook(excels[0])

    # 상태 저장 확인 (5번 항목)
    assert state.state["last_run_status"] == STATUS_SUCCESS
    assert state.state["last_successful_week"] == "2026-W39"
    assert state.state["last_output_file"].startswith("주간자금계획_")
    assert state.state["input_signature"]

    # 대외비 상세는 경영보고에 노출되지 않는다 (15번 항목)
    wb = load_workbook(excels[0], read_only=True)
    try:
        texts = []
        for row in wb["주간보고"].iter_rows(values_only=True):
            texts.extend(str(v) for v in row if v is not None)
        joined = " ".join(texts)
        assert "가상급여처리" not in joined
        assert "급여·인건비(대외비)" in joined
        # 일반 팀 자료는 정상 노출
        assert "한진택배" in joined
    finally:
        wb.close()

    # 계획 없는 출금은 확인필요에 기록된다
    assert result.issue_count > 0


def test_같은_주차_중복_실행_방지(env):
    cfg = env
    make_full_inputs(cfg, NOW)
    state = StateManager(cfg.state_dir)
    assert _run(cfg, state).status == STATUS_SUCCESS
    # 입력 변경이 없으면 재실행하지 않는다 (6번 항목)
    again = _run(cfg, state)
    assert again.status == "SKIPPED"
    # 강제 재실행 시에도 기존 결과물은 덮어쓰지 않는다 (28번 항목)
    # — 최신본만 05_결과에 남고, 이전 파일은 지난자료로 보존된다
    forced = _run(cfg, state, force=True)
    assert forced.status == STATUS_SUCCESS
    excels = sorted(cfg.folder("output")
                    .glob("주간자금계획_경영보고_*.xlsx"))
    assert len(excels) == 1
    assert excels[0].stem.endswith("_2")  # 새 실행분 (이름 충돌 회피)
    archived = sorted((cfg.folder("archive") / "지난결과")
                      .glob("주간자금계획_경영보고_*.xlsx"))
    assert len(archived) == 1 and not archived[0].stem.endswith("_2")


def test_필수파일_누락과_재시도(env):
    cfg = env
    make_full_inputs(cfg, NOW)
    # 국민은행 자료 제거 → 필수파일 누락
    for f in cfg.bank_dir("국민은행").glob("*"):
        f.unlink()
    state = StateManager(cfg.state_dir)
    result = _run(cfg, state)
    assert result.status == STATUS_WAITING_FILES
    assert "국민은행" in result.message
    # 결과물이 생성되지 않아야 한다
    assert not list(cfg.folder("output").glob("주간자금계획_*.xlsx"))

    # 파일이 준비된 뒤 재시도하면 성공한다 (18번 항목)
    prev = week_monday(NOW.date()) - timedelta(days=7)
    write_bank_csv(cfg.bank_dir("국민은행") / "국민_거래내역.csv", "국민은행", [
        nh_tx(prev, "10:00:00", in_amt=100000, balance=1000000,
              desc="입금", memo="스마트스토어")])
    retry = _run(cfg, state, mode="retry")
    assert retry.status == STATUS_SUCCESS


def test_스케줄러_파일대기_재시도(env):
    """WAITING_FILES 상태에서 주기 검사가 자동으로 재실행하는지."""
    cfg = env
    now = now_local(cfg.timezone_name)
    make_full_inputs(cfg, now)
    for f in cfg.bank_dir("농협").glob("*"):
        f.unlink()
    state = StateManager(cfg.state_dir)
    from scheduler import AutomationService
    service = AutomationService(cfg, state, LOG)

    result = service.run_job(mode="auto")
    assert result.status == STATUS_WAITING_FILES

    # 파일 없는 동안의 주기 검사: 실행하지 않음
    service._periodic_check()
    state.reload()
    assert state.state["last_run_status"] == STATUS_WAITING_FILES

    # 파일이 채워지면 주기 검사가 재실행한다
    prev = week_monday(now.date()) - timedelta(days=7)
    write_bank_csv(cfg.bank_dir("농협") / "농협_거래내역.csv", "농협", [
        nh_tx(prev, "09:00:00", in_amt=500000, balance=5000000,
              desc="입금", memo="네이버페이")])
    service._periodic_check()
    state.reload()
    assert state.state["last_run_status"] == STATUS_SUCCESS


def test_입력변경_감지(env):
    cfg = env
    make_full_inputs(cfg, NOW)
    state = StateManager(cfg.state_dir)
    assert _run(cfg, state).status == STATUS_SUCCESS

    import file_validator
    sig = file_validator.current_input_signature(cfg)
    assert not state.input_changed_after_success(sig)

    # 팀 파일 변경 → 서명 변화 감지 (29번 항목)
    team_path = cfg.folder("general_teams") / "물류팀_주간지출계획.xlsx"
    import os
    st = team_path.stat()
    os.utime(team_path, (st.st_atime, st.st_mtime + 120))
    new_sig = file_validator.current_input_signature(cfg)
    assert state.input_changed_after_success(new_sig)


def test_손상파일_재열기_검증_실패(tmp_path):
    bad = tmp_path / "깨진파일.xlsx"
    bad.write_bytes(b"this is not an xlsx file")
    assert not verify_workbook(bad, ["확인필요"])


def test_재실행시_이전_결과는_지난자료로_이동(env):
    """같은 주에 여러 번 실행해도 05_결과에는 최신 실행분만 남는다."""
    cfg = env
    make_full_inputs(cfg, NOW)
    state = StateManager(cfg.state_dir)
    assert _run(cfg, state, mode="manual").status == STATUS_SUCCESS
    assert _run(cfg, state, mode="manual",
                force=True).status == STATUS_SUCCESS

    out_dir = cfg.folder("output")
    assert len(list(out_dir.glob("주간자금계획_*.xlsx"))) == 1
    assert len(list(out_dir.glob("주간자금계획_대표보고_*.pdf"))) == 1
    assert not list(out_dir.glob("확인필요_*.xlsx"))   # 06 폴더에만 둔다
    # 이전 실행분은 삭제되지 않고 지난자료로 이동
    archive = cfg.folder("archive") / "지난결과"
    assert len(list(archive.glob("주간자금계획_*"))) >= 2
    # 06_확인필요 폴더도 최신 파일 하나만 유지
    assert len(list(cfg.folder("review").glob("확인필요_*.xlsx"))) == 1


def test_확인후_결과생성_2단계(env):
    """confirm_before_results: 확인필요 검토('확인 완료'=예) 후에만 결과 생성."""
    cfg = env
    cfg.data.setdefault("options", {})["confirm_before_results"] = True
    make_full_inputs(cfg, NOW)
    state = StateManager(cfg.state_dir)

    # 1단계: 확인필요만 만들어지고 결과물은 없다
    first = _run(cfg, state, mode="manual")
    assert first.status == STATUS_REVIEW_WAIT
    assert "확인 완료" in first.message
    assert not list(cfg.folder("output").glob("주간자금계획_*"))
    reviews = list(cfg.folder("review").glob("확인필요_*.xlsx"))
    assert len(reviews) == 1
    state.reload()
    assert state.state["last_run_status"] == STATUS_REVIEW_WAIT

    # 컨트롤이 들어 있다 (B2 드롭다운, 기본 '아니오')
    from excel_report import review_confirmed
    ok, week, sig = review_confirmed(reviews[0])
    assert not ok and week == "2026-W39" and sig

    # 사용자가 '확인 완료'를 '예'로 저장
    wb = load_workbook(reviews[0])
    wb["확인필요"]["B2"] = "예"
    wb.save(reviews[0])
    wb.close()

    # 2단계: 다시 실행하면 결과 3종이 만들어진다
    done = _run(cfg, state, mode="manual")
    assert done.status == STATUS_SUCCESS, done.message
    out = cfg.folder("output")
    assert len(list(out.glob("주간자금계획_경영보고_*.xlsx"))) == 1
    assert len(list(out.glob("주간자금계획_대표보고_*.pdf"))) == 1
    assert not list(out.glob("확인필요_*.xlsx"))
    # 확인된 확인필요 파일은 그대로 남는다
    assert reviews[0].exists()


def test_확인대기_주기검사가_완료를_감지한다(env):
    """REVIEW_WAIT 상태에서 '확인 완료' 저장을 10분 주기 검사가 잡는다."""
    cfg = env
    cfg.data.setdefault("options", {})["confirm_before_results"] = True
    now = now_local(cfg.timezone_name)
    make_full_inputs(cfg, now)
    state = StateManager(cfg.state_dir)
    from scheduler import AutomationService
    service = AutomationService(cfg, state, LOG)

    assert service.run_job(mode="auto").status == STATUS_REVIEW_WAIT

    # 미확인 상태의 주기 검사: 아무 일도 없다
    service._periodic_check()
    state.reload()
    assert state.state["last_run_status"] == STATUS_REVIEW_WAIT

    # '확인 완료' 저장 후 주기 검사 → 결과 생성
    review = next(cfg.folder("review").glob("확인필요_*.xlsx"))
    wb = load_workbook(review)
    wb["확인필요"]["B2"] = "예"
    wb.save(review)
    wb.close()
    service._periodic_check()
    state.reload()
    assert state.state["last_run_status"] == STATUS_SUCCESS
    assert list(cfg.folder("output").glob("주간자금계획_경영보고_*.xlsx"))


def test_실행창_확인대기가_저장을_즉시_감지(env):
    """'지금 실행' 창 대기: 확인 완료 저장 순간 바로 결과 생성."""
    cfg = env
    cfg.data.setdefault("options", {})["confirm_before_results"] = True
    make_full_inputs(cfg, NOW)
    state = StateManager(cfg.state_dir)
    assert _run(cfg, state, mode="manual").status == STATUS_REVIEW_WAIT

    # 대기 0분이면 기다리지 않고 그대로 종료
    cfg.data["options"]["confirm_wait_minutes"] = 0
    res = app_module._wait_for_confirmation(cfg, state, LOG, now=NOW)
    assert res.status == STATUS_REVIEW_WAIT

    # '확인 완료' 저장 → 즉시 감지해 결과 3종 생성
    cfg.data["options"]["confirm_wait_minutes"] = 5
    review = next(cfg.folder("review").glob("확인필요_*.xlsx"))
    wb = load_workbook(review)
    wb["확인필요"]["B2"] = "예"
    wb.save(review)
    wb.close()
    res = app_module._wait_for_confirmation(cfg, state, LOG,
                                            poll_seconds=0.01, now=NOW)
    assert res.status == STATUS_SUCCESS, res.message
    assert list(cfg.folder("output").glob("주간자금계획_경영보고_*.xlsx"))


def test_확인감시_틱이_저장을_감지한다(env):
    """상주 감시(30초 간격) 틱: 확인 완료 저장을 잡아 결과를 만든다."""
    cfg = env
    cfg.data.setdefault("options", {})["confirm_before_results"] = True
    now = now_local(cfg.timezone_name)
    make_full_inputs(cfg, now)
    state = StateManager(cfg.state_dir)
    from scheduler import AutomationService
    service = AutomationService(cfg, state, LOG)

    assert service.run_job(mode="auto").status == STATUS_REVIEW_WAIT
    service._review_watch_tick()          # 미확인 → 변화 없음
    state.reload()
    assert state.state["last_run_status"] == STATUS_REVIEW_WAIT

    review = next(cfg.folder("review").glob("확인필요_*.xlsx"))
    wb = load_workbook(review)
    wb["확인필요"]["B2"] = "예"
    wb.save(review)
    wb.close()
    service._review_watch_tick()
    state.reload()
    assert state.state["last_run_status"] == STATUS_SUCCESS


def test_대외비_가림_해제시_상세_표시(env):
    """mask_confidential=False: 대외비도 거래처·신청자 상세로 노출."""
    cfg = env
    cfg.data.setdefault("options", {})["mask_confidential"] = False
    make_full_inputs(cfg, NOW)
    state = StateManager(cfg.state_dir)
    assert _run(cfg, state).status == STATUS_SUCCESS

    out = next(cfg.folder("output").glob("주간자금계획_경영보고_*.xlsx"))
    wb = load_workbook(out, read_only=True)
    try:
        texts = " ".join(str(v)
                         for row in wb["주간보고"].iter_rows(values_only=True)
                         for v in row if v is not None)
    finally:
        wb.close()
    assert "가상급여처리" in texts        # 대외비 거래처가 그대로 보인다

    # 통합 행에도 개별 대외비 행이 유지된다 (집계행 아님)
    import team_loader
    rows = [{"confidential": True, "거래처": "엠에스바이오텍",
             "지출내용": "외상매입금", "신청자": "김정원",
             "예상금액": 20_000_000, "반영상태": "정상반영",
             "자금계획 반영일": date(2026, 9, 23), "요청ID": "경영-1"}]
    shown = team_loader.mask_confidential_rows(rows, enabled=False)
    assert shown[0]["거래처"] == "엠에스바이오텍"
    assert shown[0]["신청자"] == "김정원"


def test_확인필요_처리지시_라운드트립(tmp_path):
    """일자·내용·금액 분리 표기 + '처리'(계획에 반영) 지시 읽기."""
    import excel_report as er
    issues = [
        {"구분": "계획 없는 실제출금", "은행": "농협",
         "일자": date(2026, 9, 17), "내용": "결제대행사",
         "금액": 123456, "원본파일": "농협.csv"},
        {"구분": "정기지출 누락 의심", "일자": date(2026, 9, 23),
         "내용": "METLIFE — 팀 지출예정 파일에서 찾지 못함",
         "금액": 7641812, "원본파일": "자동추정_지출목록.xlsx",
         "지시항목": "METLIFE"},
    ]
    path = tmp_path / "확인필요_test.xlsx"
    er.create_issue_workbook(issues, path, week_key="2026-W39",
                             signature="sig")
    assert er.load_review_directives(path) == []   # 아직 지시 없음

    wb = load_workbook(path)
    ws = wb["확인필요"]
    # 계획 없는 실제출금: 일자(E)·내용(F)·금액(G) 분리 표기
    assert ws["E5"].value.date() == date(2026, 9, 17)
    assert ws["F5"].value == "결제대행사" and ws["G5"].value == 123456
    assert ws["I5"].value is None            # 누락 아닌 행엔 처리 칸 없음
    # 누락 의심 행: 처리 기본 '반영 안 함'
    assert ws["I6"].value == "반영 안 함"
    ws["I6"] = "계획에 반영"
    ws["G6"] = 8_000_000                     # 금액을 고치면 고친 값으로 반영
    wb.save(path)
    wb.close()

    assert er.load_review_directives(path) == [
        {"일자": date(2026, 9, 23), "금액": 8_000_000.0, "항목": "METLIFE"}]


def test_일별_비고_지출내역_요약():
    """4주일별계획 비고란: 반영일 기준 지출 내역이 요약된다."""
    from datetime import date
    daily = [{"일자": date(2026, 9, 23), "비고": ""},
             {"일자": date(2026, 9, 24), "비고": ""}]
    masked = [
        {"반영상태": "정상반영", "자금계획 반영일": date(2026, 9, 23),
         "거래처": "와우프레스", "팀명": "디자인팀", "예상금액": 25000},
        {"반영상태": "정상반영", "자금계획 반영일": date(2026, 9, 23),
         "confidential": True, "지출내용": "기타 대외비 지출",
         "예상금액": 632250},
        {"반영상태": "취소", "자금계획 반영일": date(2026, 9, 24),
         "거래처": "취소된거래처", "예상금액": 99999},
    ]
    app_module._annotate_daily_notes(daily, masked)
    assert "와우프레스 25,000" in daily[0]["비고"]
    assert "기타 대외비 지출 632,250" in daily[0]["비고"]
    assert "취소된거래처" not in (daily[1]["비고"] or "")


def test_지난자료_보관기간_지나면_삭제(tmp_path):
    """99_지난자료: 보관 일수가 지난 파일은 지우고 빈 폴더도 정리한다."""
    import os
    import time as _time
    import backup_manager

    class _Cfg:
        def folder(self, name):
            return tmp_path / {"archive": "99"}[name]

    cfg = _Cfg()
    old_dir = cfg.folder("archive") / "지난결과"
    old_dir.mkdir(parents=True)
    old_file = old_dir / "주간자금계획_옛날.xlsx"
    old_file.write_text("x")
    past = _time.time() - 20 * 86400
    os.utime(old_file, (past, past))
    recent = cfg.folder("archive") / "지난결과_최근.xlsx"
    recent.write_text("y")

    removed = backup_manager.purge_old_archive(cfg, keep_days=14)
    assert removed == 1
    assert not old_file.exists() and not old_dir.exists()  # 빈 폴더 정리
    assert recent.exists()                                  # 최근 파일 유지
    # keep_days=0이면 정리하지 않는다
    assert backup_manager.purge_old_archive(cfg, keep_days=0) == 0


def test_은행폴더_계좌별_최신파일만_유지(tmp_path):
    """모든 계좌가 더 최신 파일로 대체된 은행 파일만 지난자료로 이동한다."""
    from datetime import date
    import backup_manager

    class _Cfg:
        def __init__(self, base):
            self.base = base

        def bank_dir(self, bank):
            return self.base / "03" / bank

        def folder(self, name):
            return self.base / {"archive": "99"}[name]

    cfg = _Cfg(tmp_path)
    for bank, names in (("농협", ["구파일.csv", "신파일.csv"]),
                        ("국민은행", ["주계좌_최신.xls", "네이버_구.xls"])):
        d = cfg.bank_dir(bank)
        d.mkdir(parents=True)
        for n in names:
            (d / n).write_text("x")

    rows = [
        # 농협 한 계좌: 구파일(9/13) < 신파일(9/18) → 구파일 이동
        {"은행": "농협", "계좌": "A", "거래일": date(2026, 9, 13),
         "원본파일": "구파일.csv"},
        {"은행": "농협", "계좌": "A", "거래일": date(2026, 9, 18),
         "원본파일": "신파일.csv"},
        # 국민 두 계좌: 서로 다른 계좌라 둘 다 최신 → 유지
        {"은행": "국민은행", "계좌": "B", "거래일": date(2026, 9, 18),
         "원본파일": "주계좌_최신.xls"},
        {"은행": "국민은행", "계좌": "C", "거래일": date(2026, 9, 15),
         "원본파일": "네이버_구.xls"},
    ]
    moved = backup_manager.archive_superseded_bank_files(cfg, rows)
    assert moved == 1
    assert not (cfg.bank_dir("농협") / "구파일.csv").exists()
    assert (cfg.bank_dir("농협") / "신파일.csv").exists()
    assert (cfg.bank_dir("국민은행") / "네이버_구.xls").exists()
    assert (cfg.folder("archive") / "지난입력파일" / "농협" / "구파일.csv").exists()
