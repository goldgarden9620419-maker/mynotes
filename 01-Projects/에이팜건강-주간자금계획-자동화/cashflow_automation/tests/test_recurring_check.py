# -*- coding: utf-8 -*-
"""일괄 반영/제외 스위치와 금주 정기지출 체크 테스트."""
from datetime import date, timedelta

from openpyxl import load_workbook

import forecast_engine as fe

BASE = date(2026, 9, 21)  # 월요일
RECUR = [{"정기지출명": "SKB", "대표 지급일": 25,
          "평균 월지출": 220_000.0, "신뢰도": "상"},
         {"정기지출명": "코웨이", "대표 지급일": 28,
          "평균 월지출": 113_398.0, "신뢰도": "상"}]


def test_일괄_전체제외_적용과_개별수정_보존(tmp_path):
    draft = tmp_path / "자동추정_지출목록.xlsx"
    fe.refresh_auto_draft_file(draft, RECUR, BASE)
    # 기본(개별 관리): 새 행은 '반영'
    drafts, excluded = fe.load_auto_drafts(draft, BASE)
    assert len(drafts) == 2 and excluded == 0

    # 사용자가 안내 시트에서 '전체 제외' 선택 → 다음 갱신 때 전 행 적용
    wb = load_workbook(draft)
    wb[fe.GUIDE_SHEET].cell(row=6, column=2, value="전체 제외")
    wb.save(draft)
    wb.close()
    fe.refresh_auto_draft_file(draft, RECUR, BASE)
    drafts, excluded = fe.load_auto_drafts(draft, BASE)
    assert not drafts and excluded == 2

    # 이후 개별 행을 '반영'으로 되돌리면 다음 갱신에도 보존된다
    wb = load_workbook(draft)
    wb[fe.AUTO_DRAFT_SHEET].cell(row=2, column=5, value="반영")
    wb.save(draft)
    wb.close()
    fe.refresh_auto_draft_file(draft, RECUR, BASE)
    drafts, excluded = fe.load_auto_drafts(draft, BASE)
    assert len(drafts) == 1 and excluded == 1


def test_전체제외_기본값이면_생성부터_제외(tmp_path):
    draft = tmp_path / "자동추정_지출목록.xlsx"
    fe.refresh_auto_draft_file(draft, RECUR, BASE, default_mode="전체 제외")
    drafts, excluded = fe.load_auto_drafts(draft, BASE)
    assert not drafts and excluded == 2
    wb = load_workbook(draft)
    assert wb[fe.GUIDE_SHEET].cell(row=6, column=2).value == "전체 제외"
    wb.close()


def test_공휴일_시트_생성과_로드(tmp_path):
    """기준파일 '공휴일' 시트: 최초 생성 시 기본 공휴일을 채우고,
    이미 있으면 사용자 관리분을 그대로 읽는다."""
    from openpyxl import Workbook
    base = tmp_path / "기준.xlsx"
    wb = Workbook()
    wb.active.title = "카드결제기준"
    wb.save(base)
    wb.close()

    assert fe.ensure_holiday_sheet(base) > 0
    assert fe.ensure_holiday_sheet(base) == 0      # 재실행 시 안 건드림
    holidays = fe.load_holidays(base)
    assert holidays[date(2026, 9, 25)] == "추석"
    assert date(2026, 10, 9) in holidays

    # 사용자가 임시공휴일을 추가하면 그대로 읽힌다
    wb = load_workbook(base)
    wb[fe.HOLIDAY_SHEET].append([date(2026, 11, 3), "임시공휴일"])
    wb.save(base)
    wb.close()
    assert fe.load_holidays(base)[date(2026, 11, 3)] == "임시공휴일"


def test_주말실행이면_대표보고_일별전망은_차주(tmp_path):
    """display_week_start를 차주 월요일로 주면 금주일별이 차주 월~금."""
    fc = fe.build_forecast([], BASE, 1_000_000, [], [], [], [0.8], 0.8,
                           display_week_start=BASE + timedelta(days=7))
    days = fc["rate_scenarios"][0.8]["금주일별"]
    assert [d for d, _b, _s in days] == [BASE + timedelta(days=7 + i)
                                         for i in range(5)]


def test_금주체크_판정과_구간(tmp_path):
    start, end = date(2026, 9, 21), date(2026, 10, 2)
    draft_rows = [
        {"일자": date(2026, 9, 23), "정기지출명": "NH기업카드",
         "예상금액": 6_066_502.0, "반영": "제외"},   # 별칭으로 취합과 연결
        {"일자": date(2026, 9, 23), "정기지출명": "우리카드결제대금",
         "예상금액": 26_000_000.0, "반영": "제외"},  # 같은 취합 건을 재사용 못함
        {"일자": date(2026, 9, 25), "정기지출명": "SKB",
         "예상금액": 220_000.0, "반영": "제외"},     # 취합에 없음 → 누락
        {"일자": date(2026, 9, 28), "정기지출명": "코웨이",
         "예상금액": 113_398.0, "반영": "반영"},     # 평균 금액 자동 반영
        {"일자": date(2026, 10, 15), "정기지출명": "한화생명",
         "예상금액": 1_000_000.0, "반영": "제외"},   # 구간 밖
    ]
    plans = [{"자금계획 반영일": date(2026, 9, 23),
              "예상금액": 10_000_000.0, "거래처": "농협카드",
              "지출내용": "8월 카드대금"}]
    aliases = {"NH기업카드": ["농협카드"]}
    variable = [{"정기지출명": "㈜에이팜", "대표 지급일": 30,
                 "평균 월지출": 1_369_025.0, "성격": "변동"}]
    rows = fe.weekly_recurring_check(draft_rows, plans, aliases,
                                     start, end, variable_items=variable)
    by = {r["항목"]: r for r in rows}
    assert "한화생명" not in by
    assert by["NH기업카드"]["판정"] == fe.CHECK_TEAM_OK
    assert by["NH기업카드"]["팀제출금액"] == 10_000_000.0
    # 취합 한 건은 한 항목만 커버 — 유사도 높은 NH(별칭 일치)가 가져가고
    # 우리카드는 누락 의심으로 남는다
    assert by["우리카드결제대금"]["누락"]
    assert by["SKB"]["누락"] and by["SKB"]["판정"] == fe.CHECK_MISSING
    assert by["코웨이"]["판정"] == fe.CHECK_AUTO and not by["코웨이"]["누락"]
    # 성격 '변동' 항목은 목록에 없어도 매월 만들어 검사한다 (9/30 도래)
    assert by["㈜에이팜"]["예정일"] == date(2026, 9, 30)
    assert by["㈜에이팜"]["누락"]
    assert [r["예정일"] for r in rows] == sorted(r["예정일"] for r in rows)
