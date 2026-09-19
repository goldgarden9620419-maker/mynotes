# -*- coding: utf-8 -*-
"""라이브 양식 채우기 테스트: 수식 보존 + 주차 날짜 재고정."""
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook

from live_report import fill_live_workbook, verify_live_workbook

OLD_MONDAY = date(2026, 9, 14)
NEW_MONDAY = date(2026, 9, 21)

_DAILY_C = ("=INDEX('설정및분류'!$C$6:$C$12,WEEKDAY(A{r},2))*'요약'!$B$13")
_OLD_SUMIFS = ("=SUMIFS('4주일별계획'!$C$6:$C$33,'4주일별계획'!$A$6:$A$33,"
               "\">=\"&DATE(2026,9,14),'4주일별계획'!$A$6:$A$33,"
               "\"<=\"&DATE(2026,9,20))")


def _make_stub_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws["B6"], ws["B13"] = 100, 0.8
    ws["B14"] = "='4주일별계획'!J33"
    ws["B15"] = "=MIN('4주일별계획'!J6:J33)"

    ws = wb.create_sheet("4주일별계획")
    for i in range(28):
        r = 6 + i
        ws.cell(row=r, column=1,
                value=datetime(2026, 9, 14) + timedelta(days=i))
        ws.cell(row=r, column=3, value=_DAILY_C.format(r=r))
        ws.cell(row=r, column=8, value=f"=C{r}+D{r}-E{r}-F{r}-G{r}")
        ws.cell(row=r, column=10, value=f"=I{r}+H{r}")

    ws = wb.create_sheet("13주주별계획")
    for w in range(13):
        r = 6 + w
        ws.cell(row=r, column=1, value=f"{w + 1}주차")
        if w < 4:
            for col in ("C", "D", "E", "F", "G"):
                ws[f"{col}{r}"] = _OLD_SUMIFS
        else:
            ws.cell(row=r, column=3,
                    value="=SUM('설정및분류'!$C$6:$C$12)*'요약'!$B$13")

    wb.create_sheet("정기지출분석")
    ws = wb.create_sheet("계좌내역통합_RAW")
    # 옛 수동 붙여넣기용 잔여 수식 (자동화에서는 정리 대상)
    ws.cell(row=50, column=1, value="=IF('주간계좌_붙여넣기'!B8=\"\",\"\",1)")
    wb.create_sheet("주간계좌_붙여넣기")
    ws = wb.create_sheet("설정및분류")
    for i in range(7):
        ws.cell(row=6 + i, column=3, value=0)
    wb.save(path)
    wb.close()


def _fake_report(base: date) -> dict:
    daily = []
    for i in range(28):
        d = base + timedelta(days=i)
        daily.append({"일자": d, "확정·기타입금": 0, "팀별 송금예정": 0,
                      "카드결제": 500000 if i == 3 else 0,
                      "자동이체": 0, "기타지출": 100000 if i == 0 else 0,
                      "비고": "서울보증보험 632,250" if i == 2 else ""})
    weekly = [{"확정기타입금": 0, "송금예정": 0, "카드결제": 0,
               "자동이체": 0, "기타지출": 0} for _ in range(13)]
    weekly[5]["기타지출"] = 700000
    return {
        "meta": {"base_date": base},
        "forecast": {"weekday_avg": {i: 1000000 for i in range(5)},
                     "daily": daily, "weekly": weekly},
        "total_balance": 12345678,
        "history_stats": {"외부입금": 1, "외부출금": 2, "온라인입금": 3,
                          "주평균온라인": 4},
        "balances": {("국민은행", "596001-04-154577"): 12345678},
        "account_last_dates": {("국민은행", "596001-04-154577"):
                               base - timedelta(days=3)},
        "account_labels": {},
        "recurring": [{"은행": "우리은행", "정기지출명": "사회보험",
                       "분류": "4대보험", "발생개월수": 6, "거래건수": 6,
                       "평균 월지출": 100.0, "최소 월지출": 90.0,
                       "최대 월지출": 110.0, "대표 지급일": 10,
                       "신뢰도": "상"}],
        "bank_rows": [{"거래일": base - timedelta(days=3),
                       "거래일시": None, "은행": "국민은행",
                       "계좌": "596001-04-154577", "출금액": 0.0,
                       "입금액": 50000.0, "거래후잔액": 12345678.0,
                       "적요": "전자금융", "기재내용·상대방": "테스트",
                       "취급점": "", "자동분류": "", "내부이체": False,
                       "정기지출후보": False, "현금유출입": 50000.0,
                       "반영상태": "정상반영"}],
        "integrated_masked": [
            {"요청ID": "디자인-20260918-001", "팀명": "디자인팀",
             "신청자": "김담당", "품의승인": "승인",
             "지급예정일": date(2026, 9, 25),
             "자금계획 반영일": date(2026, 10, 23),
             "거래처": "OO인쇄", "지출내용": "리플렛 인쇄",
             "예상금액": 25000, "지급방법": "법인카드",
             "카드구분": "국민카드", "확정여부": "확정",
             "진행상태": "신규", "최종수정일": date(2026, 9, 18),
             "원본파일": "디자인팀_지출계획.xlsx",
             "반영상태": "정상반영", "확인사항": ""},
            {"요청ID": "경영-대외비-집계", "팀명": "경영지원팀",
             "신청자": "", "품의승인": "승인대기",
             "지급예정일": date(2026, 9, 23),
             "자금계획 반영일": date(2026, 9, 23),
             "거래처": "기타 대외비 지출", "지출내용": "(대외비 집계)",
             "예상금액": 632250, "지급방법": "계좌송금",
             "카드구분": "", "확정여부": "미확정",
             "진행상태": "신규", "최종수정일": datetime(2026, 9, 18, 10, 0),
             "원본파일": "경영지원팀_대외비_지출계획.xlsx",
             "반영상태": "정상반영",
             "확인사항": "승인대기 상태(정책상 반영)"},
        ],
        "apalm_expenses": [
            {"일자": date(2026, 9, 23), "출처": "팀 지출계획",
             "반영": "미반영(별도 관리)", "팀명": "경영지원팀",
             "거래처": "서원회계법인", "지출내용": "기장료",
             "예상금액": 330000, "지급방법": "계좌송금", "비고": "에이팜"},
            {"일자": date(2026, 9, 28), "출처": "정기지출 추정",
             "반영": "반영", "팀명": "", "거래처": "㈜에이팜",
             "지출내용": "정기지출 추정(자동 초안): ㈜에이팜 신뢰도 상",
             "예상금액": 5000000, "지급방법": "", "비고": ""},
        ],
    }


def test_라이브양식_다음주로_재고정(tmp_path):
    template = tmp_path / "템플릿.xlsx"
    _make_stub_template(template)
    out = tmp_path / "결과.xlsx"
    fill_live_workbook(template, _fake_report(NEW_MONDAY), out)

    assert verify_live_workbook(out, NEW_MONDAY)
    wb = load_workbook(out)
    daily = wb["4주일별계획"]
    # 날짜가 새 주차로 이동
    assert daily["A6"].value.date() == NEW_MONDAY
    assert daily["A33"].value.date() == NEW_MONDAY + timedelta(days=27)
    # 반영률 수식은 그대로
    assert "INDEX" in daily["C6"].value and "$B$13" in daily["C6"].value
    # 지출 값 기록 (1일차 기타지출, 4일차 카드)
    assert daily["G6"].value == 100000
    assert daily["F9"].value == 500000
    # 비고란에 그날 지출 내역 요약 (3일차), 없는 날은 비움
    assert daily["K8"].value == "서울보증보험 632,250"
    assert daily["K6"].value is None
    weekly = wb["13주주별계획"]
    # 1주차 SUMIFS가 새 날짜로 재작성
    assert "DATE(2026,9,21)" in weekly["C6"].value
    assert "DATE(2026,9,27)" in weekly["C6"].value
    # 5주 이후 반영률 수식 유지 + 6주차 기타지출 값
    assert "$B$13" in weekly["C10"].value
    assert weekly["G11"].value == 700000
    summary = wb["요약"]
    assert summary["B6"].value == 12345678
    assert summary["B14"].value.startswith("=")  # 요약 수식 보존
    raw = wb["계좌내역통합_RAW"]
    assert raw["C2"].value == "국민은행"
    assert raw["F2"].value == 50000
    # 수동 붙여넣기 시트 제거 + 잔여 참조 수식 정리
    assert "주간계좌_붙여넣기" not in wb.sheetnames
    assert raw["A50"].value is None
    # 정기지출분석은 사용자 편집용이라 맨 끝 시트에 둔다
    assert wb.sheetnames[-1] == "정기지출분석"
    # 머리글 필터 + 성격 드롭다운 + K4 일괄 변경 셀이 붙는다
    rec = wb["정기지출분석"]
    assert rec.auto_filter.ref == "A5:K6"          # 자료 1행
    assert rec["K4"].value == "변경 안 함"
    rec_dvs = [str(dv.formula1)
               for dv in rec.data_validations.dataValidation]
    assert any("전체 변동" in f for f in rec_dvs)
    assert any("정기,변동,제외" in f for f in rec_dvs)
    # 지출계획 취합 시트가 새로 생성되어 자동 반영된다
    assert "지출계획_취합" in wb.sheetnames
    exp = wb["지출계획_취합"]
    assert exp["A5"].value == "요청ID"
    assert exp["A6"].value == "디자인-20260918-001"
    assert exp["I6"].value == 25000            # 예상금액
    assert exp["I6"].number_format == "#,##0"
    def _d(v):  # openpyxl은 날짜를 datetime으로 읽는다
        return v.date() if isinstance(v, datetime) else v
    assert _d(exp["E6"].value) == date(2026, 9, 25)  # 지급예정일
    assert exp["A7"].value == "경영-대외비-집계"
    assert _d(exp["N7"].value) == date(2026, 9, 18)  # datetime→date 변환
    assert exp["Q7"].value == "승인대기 상태(정책상 반영)"
    assert exp["A8"].value is None
    assert exp.freeze_panes == "A6"
    assert exp.auto_filter.ref == "A5:Q7"   # 머리글 필터 (데이터 2행)
    # 에이팜 지출계획 시트: 취합 시트 바로 다음, 미반영/반영 구분 + 합계
    names = wb.sheetnames
    assert names.index("에이팜 지출계획") == names.index("지출계획_취합") + 1
    ap = wb["에이팜 지출계획"]
    assert ap["A5"].value == "일자"
    assert _d(ap["A6"].value) == date(2026, 9, 23)
    assert ap["D6"].value == "미반영(별도 관리)"
    assert ap["F6"].value == "서원회계법인"
    assert ap["H6"].value == 330000
    assert ap["J6"].value == "에이팜"
    assert ap["D7"].value == "반영" and ap["F7"].value == "㈜에이팜"
    assert ap["G8"].value == "합계(자금계획 미반영)"
    assert ap["H8"].value == 330000
    assert ap["G9"].value == "합계(자금계획 반영)"
    assert ap["H9"].value == 5000000
    wb.close()


def test_지출시트_재실행시_잔여행_정리(tmp_path):
    """행이 줄어든 다음 주 실행에서 이전 잔여 데이터가 남지 않는다."""
    template = tmp_path / "템플릿.xlsx"
    _make_stub_template(template)
    out1 = tmp_path / "결과1.xlsx"
    fill_live_workbook(template, _fake_report(NEW_MONDAY), out1)
    report2 = _fake_report(NEW_MONDAY)
    report2["integrated_masked"] = report2["integrated_masked"][:1]
    out2 = tmp_path / "결과2.xlsx"
    fill_live_workbook(out1, report2, out2)
    wb = load_workbook(out2)
    exp = wb["지출계획_취합"]
    assert exp["A6"].value == "디자인-20260918-001"
    assert exp["A7"].value is None
    assert exp["I7"].value is None
    wb.close()


def _side(cell, name):
    s = getattr(cell.border, name, None)
    return s.style if s is not None else None


def _assert_exec_box(ws, first_row, last_row):
    """실행일~차주 금요일 구간이 하나의 붉은 외곽 상자인지 확인."""
    assert _side(ws.cell(row=first_row, column=1), "top") == "medium"
    assert _side(ws.cell(row=first_row, column=1), "left") == "medium"
    assert _side(ws.cell(row=first_row, column=11), "top") == "medium"
    assert _side(ws.cell(row=last_row, column=11), "right") == "medium"
    assert _side(ws.cell(row=last_row, column=6), "bottom") == "medium"
    # 상자 내부는 격자가 아니다 (셀별 테두리 아님)
    mid = (first_row + last_row) // 2
    assert _side(ws.cell(row=mid, column=3), "top") != "medium"
    assert _side(ws.cell(row=mid, column=3), "left") != "medium"
    # 구간 밖은 표시가 없다
    assert _side(ws.cell(row=last_row + 1, column=1), "left") != "medium"
    if first_row > 6:
        assert _side(ws.cell(row=first_row - 1, column=1), "left") != "medium"
    # 예전 TODAY() 조건부서식은 남아있지 않다
    assert not [rule for rules in ws.conditional_formatting
                for rule in rules.rules
                if rule.formula and "WEEKDAY(TODAY()" in rule.formula[0]]


def test_실행일_차주금요일_붉은_상자(tmp_path):
    """붉은 상자는 실행일부터 차주 금요일까지를 하나로 묶는다."""
    template = tmp_path / "템플릿.xlsx"
    _make_stub_template(template)
    # 월요일(9/21) 실행 → 9/21 ~ 차주 금요일 10/2 (6~17행)
    out1 = tmp_path / "결과1.xlsx"
    fill_live_workbook(template, _fake_report(NEW_MONDAY), out1)
    wb = load_workbook(out1)
    _assert_exec_box(wb["4주일별계획"], 6, 17)
    # 비고란은 자동 줄바꿈
    assert wb["4주일별계획"]["K8"].alignment.wrap_text
    wb.close()
    # 목요일(9/24) 실행 → 9/24 ~ 10/2 (9~17행), 재실행에도 상자는 하나
    report2 = _fake_report(NEW_MONDAY)
    report2["meta"]["run_date"] = NEW_MONDAY + timedelta(days=3)
    out2 = tmp_path / "결과2.xlsx"
    fill_live_workbook(out1, report2, out2)
    wb = load_workbook(out2)
    _assert_exec_box(wb["4주일별계획"], 9, 17)
    wb.close()


def test_라이브양식_검증은_이전주면_실패(tmp_path):
    template = tmp_path / "템플릿.xlsx"
    _make_stub_template(template)
    out = tmp_path / "결과.xlsx"
    fill_live_workbook(template, _fake_report(NEW_MONDAY), out)
    assert not verify_live_workbook(out, OLD_MONDAY)
