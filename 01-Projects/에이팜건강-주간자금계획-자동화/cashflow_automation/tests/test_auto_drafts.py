# -*- coding: utf-8 -*-
"""자동추정 지출 목록 파일 관리(이관·갱신·반영) 테스트."""
from datetime import date

from openpyxl import Workbook, load_workbook

import forecast_engine as fe

BASE = date(2026, 9, 21)  # 월요일


def _base_workbook(path):
    wb = Workbook()
    wb.active.title = "카드결제기준"
    ws = wb.create_sheet("주간조정")
    ws.append(["일자", "조정입금", "조정지출", "내용"])
    ws.append([date(2026, 9, 28), None, 1_369_025,
               "정기지출 추정(자동 초안): ㈜에이팜 신뢰도 하"])
    ws.append([date(2026, 10, 5), None, 220_000,
               "정기지출 추정(자동 초안): SKB 신뢰도 상"])
    ws.append([date(2026, 9, 25), None, 500_000, "수동 입력 지출"])
    ws.append([date(2026, 9, 30), 10_021_867, None,
               "관계사 외상대 입금 추정(주식회사에이팜, 말일)"])
    wb.save(path)
    wb.close()


def test_주간조정_자동초안_이관(tmp_path):
    base = tmp_path / "기준.xlsx"
    _base_workbook(base)
    draft = tmp_path / "자동추정_지출목록.xlsx"
    overrides = {"㈜에이팜": {"분류": "외상매입금", "성격": "변동"}}

    assert fe.migrate_auto_drafts(base, draft, overrides) == 2
    # 목록 파일: 에이팜은 '제외', SKB는 '반영'
    wb = load_workbook(draft)
    rows = {r[1]: r for r in
            wb[fe.AUTO_DRAFT_SHEET].iter_rows(min_row=2, values_only=True)}
    assert rows["㈜에이팜"][4] == "제외"
    assert rows["SKB"][4] == "반영" and rows["SKB"][2] == 220_000
    wb.close()
    # 주간조정에는 수동 조정·입금 추정만 남는다
    remain = fe.load_adjustments(base)
    assert {a["내용"] for a in remain} == {
        "수동 입력 지출", "관계사 외상대 입금 추정(주식회사에이팜, 말일)"}
    # 파일이 이미 있으면 재이관하지 않는다
    assert fe.migrate_auto_drafts(base, draft, overrides) == 0


def test_목록_갱신은_사용자_수정을_보존(tmp_path):
    draft = tmp_path / "자동추정_지출목록.xlsx"
    recurring = [{"정기지출명": "SKB", "대표 지급일": 15,
                  "평균 월지출": 220_000.0, "신뢰도": "상"}]
    added, pruned = fe.refresh_auto_draft_file(draft, recurring, BASE)
    # 4주(기준일 9/21 + 27일 = 10/18) 안: 10/15 한 건만
    # (5주차 이후는 13주 계획이 정기지출을 직접 배분 — 이중 반영 방지)
    assert (added, pruned) == (1, 0)

    # 사용자가 10/15 금액을 수정하고 지난 행을 하나 흉내낸다
    wb = load_workbook(draft)
    ws = wb[fe.AUTO_DRAFT_SHEET]
    for row in ws.iter_rows(min_row=2):
        if row[0].value and row[0].value.month == 10:
            row[2].value = 999_999
    ws.append([date(2026, 9, 1), "SKB", 220_000, "상", "반영", ""])
    wb.save(draft)
    wb.close()

    added, pruned = fe.refresh_auto_draft_file(draft, recurring, BASE)
    assert added == 0 and pruned == 1     # 중복 추가 없음, 지난 행 정리
    drafts, excluded = fe.load_auto_drafts(draft, BASE)
    amounts = {d["일자"]: d["조정지출"] for d in drafts}
    assert amounts[date(2026, 10, 15)] == 999_999   # 수정 보존
    assert excluded == 0
    assert all(a["내용"].startswith("정기지출 추정(자동 초안): SKB")
               for a in drafts)

    # 4주 밖 미래 행은 아직 반영하지 않는다 (그 주가 오면 반영)
    from openpyxl import load_workbook as _lw
    wb = _lw(draft)
    wb[fe.AUTO_DRAFT_SHEET].append(
        [date(2026, 11, 15), "SKB", 220_000, "상", "반영", ""])
    wb.save(draft)
    wb.close()
    drafts2, _ = fe.load_auto_drafts(draft, BASE)
    assert all(d["일자"] <= date(2026, 10, 18) for d in drafts2)


def test_제외_표시는_반영하지_않는다(tmp_path):
    draft = tmp_path / "자동추정_지출목록.xlsx"
    fe.refresh_auto_draft_file(
        draft, [{"정기지출명": "코웨이", "대표 지급일": 5,
                 "평균 월지출": 113_398.0, "신뢰도": "상"}], BASE)
    wb = load_workbook(draft)
    ws = wb[fe.AUTO_DRAFT_SHEET]
    first = next(ws.iter_rows(min_row=2))
    first[4].value = "제외"
    wb.save(draft)
    wb.close()
    drafts, excluded = fe.load_auto_drafts(draft)
    assert excluded == 1
    assert all(d["일자"] != first[0].value for d in drafts)


def test_별칭_등록시_같은달_취합입력이면_금액무관_제외(tmp_path):
    """카드대금처럼 금액이 변하는 정기지출: 별칭으로 이중 반영을 막는다."""
    adjustments = [
        # 이름이 달라 기존 유사도 규칙으로는 못 잡고, 금액도 25% 넘게 다름
        {"일자": date(2026, 9, 23), "조정입금": 0.0, "조정지출": 6_066_502.0,
         "내용": "정기지출 추정(자동 초안): NH기업카드 신뢰도 중"},
        # 별칭 없음 → 기존 규칙 그대로 유지(이름 불일치라 반영 유지)
        {"일자": date(2026, 9, 23), "조정입금": 0.0, "조정지출": 1_000_000.0,
         "내용": "정기지출 추정(자동 초안): 한화생 신뢰도 중"},
    ]
    plans = [{"자금계획 반영일": date(2026, 9, 23),
              "예상금액": 10_000_000.0, "거래처": "농협카드",
              "지출내용": "8월 카드대금"}]
    aliases = {"NH기업카드": ["농협카드"]}
    kept, skipped = fe.filter_duplicate_adjustments(
        adjustments, plans, aliases=aliases)
    assert len(skipped) == 1 and "NH기업카드" in skipped[0]["내용"]
    assert len(kept) == 1 and "한화생" in kept[0]["내용"]
    # 다른 달의 취합 입력은 제외 근거가 아니다
    plans_oct = [dict(plans[0], **{"자금계획 반영일": date(2026, 10, 23)})]
    kept2, skipped2 = fe.filter_duplicate_adjustments(
        adjustments, plans_oct, aliases=aliases)
    assert not skipped2 and len(kept2) == 2


def test_별칭_시트_생성과_로드(tmp_path):
    draft = tmp_path / "자동추정_지출목록.xlsx"
    fe.refresh_auto_draft_file(draft, [], BASE)
    aliases = fe.load_draft_aliases(draft)
    assert aliases.get("NH기업카드") == ["농협카드"]   # 기본 예시 행


def test_표_서식과_일자순_정렬(tmp_path):
    """갱신 때마다 일자순 정렬 + 표 서식 + 제외 회색 조건부서식."""
    draft = tmp_path / "자동추정_지출목록.xlsx"
    fe.refresh_auto_draft_file(
        draft, [{"정기지출명": "SKB", "대표 지급일": 15,
                 "평균 월지출": 220_000.0, "신뢰도": "상"}], BASE)
    # 순서 뒤죽박죽 행 추가 후 재갱신 → 정렬 확인
    wb = load_workbook(draft)
    ws = wb[fe.AUTO_DRAFT_SHEET]
    ws.append([date(2026, 9, 22), "코웨이", 113_398, "상", "반영", ""])
    wb.save(draft)
    wb.close()
    fe.refresh_auto_draft_file(draft, [], BASE)
    wb = load_workbook(draft)
    ws = wb[fe.AUTO_DRAFT_SHEET]
    dates = [ws.cell(row=r, column=1).value for r in (2, 3)]
    assert dates[0].date() == date(2026, 9, 22)      # 9/22이 10/15보다 먼저
    assert dates[1].date() == date(2026, 10, 15)
    # 머리글 남색 + 필터 + 틀 고정 + 제외 회색 규칙
    assert ws.cell(row=1, column=1).fill.start_color.rgb.endswith("1F4E79")
    assert ws.auto_filter.ref.startswith("A1:F")
    assert ws.freeze_panes == "A2"
    assert any("제외" in str(r.formula[0])
               for rules in ws.conditional_formatting for r in rules.rules)
    wb.close()


def test_반영열_드롭다운(tmp_path):
    """반영 열(E)에 '반영/제외' 드롭다운이 생성·유지된다."""
    draft = tmp_path / "자동추정_지출목록.xlsx"
    fe.refresh_auto_draft_file(
        draft, [{"정기지출명": "SKB", "대표 지급일": 15,
                 "평균 월지출": 220_000.0, "신뢰도": "상"}], BASE)
    wb = load_workbook(draft)
    dvs = wb[fe.AUTO_DRAFT_SHEET].data_validations.dataValidation
    assert any("반영" in str(dv.formula1) for dv in dvs)
    wb.close()
    # 재갱신해도 드롭다운이 중복 생성되지 않는다
    fe.refresh_auto_draft_file(draft, [], BASE)
    wb = load_workbook(draft)
    dvs = wb[fe.AUTO_DRAFT_SHEET].data_validations.dataValidation
    assert sum(1 for dv in dvs if "반영" in str(dv.formula1)) == 1
    wb.close()
