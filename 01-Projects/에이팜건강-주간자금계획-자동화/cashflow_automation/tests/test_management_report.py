# -*- coding: utf-8 -*-
"""주간 자금 경영보고(대화형) 워크북 테스트."""
from datetime import date

from openpyxl import load_workbook

import management_report as mr


def _report():
    base = date(2026, 9, 21)  # 월요일
    daily = []
    for i in range(28):
        d = date.fromordinal(base.toordinal() + i)
        daily.append({"일자": d, "실적": False, "온라인 예상입금": 0.0,
                      "확정·기타입금": 100000.0 if i == 1 else 0.0,
                      "팀별 송금예정": 0.0, "카드결제": 0.0,
                      "자동이체": 0.0, "기타지출": 0.0})
    return {
        "meta": {"company": "(주)에이팜건강", "base_date": base,
                 "run_at": "2026-09-21 09:10:00"},
        "total_balance": 10_000_000,
        "balances": {("우리은행", "220351"): 4_000_000,
                     ("국민은행", "169124"): 6_000_000},
        "account_labels": {}, "account_last_dates": {},
        "forecast": {
            "daily": daily, "rate": 0.8, "actual_until": None,
            "start_balance": 10_000_000,
            "weekday_avg": {i: 1_000_000 if i < 5 else 0 for i in range(7)},
            "rate_scenarios": {
                0.8: {"4주 기말잔액": -5_000_000, "4주 최저잔액": -5_000_000,
                      "자금부족 예상일": date(2026, 10, 10)},
                0.9: {"4주 기말잔액": -2_000_000, "4주 최저잔액": -2_000_000,
                      "자금부족 예상일": date(2026, 10, 12)},
                1.0: {"4주 기말잔액": 1_000_000, "4주 최저잔액": 500_000,
                      "자금부족 예상일": None},
            },
        },
        "holidays": {date(2026, 9, 24): "추석 연휴",
                     date(2026, 9, 25): "추석"},
        "week_expenses": [
            {"일자": date(2026, 9, 22), "구분": "물류팀",
             "내용": "한진택배 택배비", "금액": 500_000,
             "지급방법": "계좌송금"},
            {"일자": date(2026, 10, 14), "구분": "경영지원팀",
             "내용": "대외비 급여·인건비(대외비)", "금액": 3_000_000,
             "지급방법": "계좌송금"},
        ],
        "stability_target": 0,
    }


def test_경영보고_생성과_수식(tmp_path):
    out = tmp_path / "주간자금계획_경영보고_test.xlsx"
    mr.create_management_workbook(_report(), out)
    assert mr.verify_management_workbook(out)

    wb = load_workbook(out)
    ws = wb["경영보고"]
    # ① 총잔액 합계 수식 — 총잔액은 12행 고정 (2026-09-23 재배치)
    assert str(ws["C12"].value).startswith("=SUM(")
    # ④ 4주(28일) 일별 표(140~167행): 반영률 참조 + 지출 SUMIFS + 시작잔액
    assert "$F$12" in str(ws["C140"].value)
    # 반영률(F12)은 총잔액 행 옆 드롭다운 (60~100%)
    rate_dvs = [dv for dv in ws.data_validations.dataValidation
                if "80%" in str(dv.formula1)]
    assert rate_dvs and "F12" in str(rate_dvs[0].sqref)
    assert "SUMIFS" in str(ws["D140"].value)
    assert "$E$16" in str(ws["D140"].value)     # ② 지출 표 구간 참조
    assert str(ws["E140"].value).startswith("=$K$139")
    assert ws["A167"].value.date() == date(2026, 10, 18)  # 28일째
    assert str(ws["C168"].value).startswith("=MIN(E140:E167")
    # 실행일(월 9/21)~차주 금요일(10/2) 구간이 하나의 붉은 상자 (격자 아님)
    def _side(cell, name):
        s = getattr(cell.border, name, None)
        return s.style if s is not None else None
    assert _side(ws["A140"], "top") == "medium"
    assert _side(ws["A140"], "left") == "medium"
    assert _side(ws["F151"], "bottom") == "medium"
    assert _side(ws["I151"], "right") == "medium"   # 상자는 내역 열(I)까지
    assert _side(ws["C143"], "left") != "medium"
    assert _side(ws["A152"], "left") != "medium"   # 구간 밖은 없음
    # ② 4주 지출예정 표(16~109행 — 현황 바로 다음): 항목·합계
    assert ws["A16"].value is not None and ws["E16"].value == 500000
    assert ws["D17"].value == "대외비 급여·인건비(대외비)"
    assert str(ws["E110"].value).startswith("=SUMIFS")  # 보류 제외 합계
    assert ws.auto_filter.ref == "A15:I109"   # 지급일별 필터
    # ⑤ 필요 추가 입금: 금액은 좁은 B열을 피해 C~F열 (##### 방지)
    assert ws["A174"].value == 0.8 and ws["A176"].value == 1.0
    assert ws["B174"].value is None                # B열엔 금액 없음
    assert ws["C174"].value == -5_000_000          # 4주 기말잔액
    assert "MAX(0,$C$172" in str(ws["F174"].value)
    # ⑥ 전달 메모 수식
    memo = str(ws["A180"].value)
    assert memo.startswith("=IF(") and "건강사업팀" in memo
    # 공휴일 예상입금 0: 숨김 요일평균(J)이 0 → 입금 수식도 0
    assert ws["J143"].value == 0 and ws["J144"].value == 0  # 추석 9/24~25
    assert ws["J141"].value == 1_000_000                    # 평일은 평균
    # 주말·공휴일 일자는 붉은 글자 (9/24 추석=143행, 9/26 토=145행)
    assert str(ws["A143"].font.color.rgb).endswith("C00000")
    assert str(ws["B143"].font.color.rgb).endswith("C00000")
    assert str(ws["A145"].font.color.rgb).endswith("C00000")
    assert not str(ws["A141"].font.color.rgb                 # 화요일은 검정
                   if ws["A141"].font.color else "").endswith("C00000")
    # 인쇄: A4 세로 폭 맞춤 + 빈 지출행 숨김 (자료 2행 + 예비 3행만 표시)
    assert ws.page_setup.orientation == "portrait"
    assert int(ws.page_setup.fitToWidth or 0) == 1
    assert "$A$1" in str(ws.print_area) and "$I$" in str(ws.print_area)
    assert not ws.row_dimensions[20].hidden      # 예비행 (16+2건 뒤 3행)
    assert ws.row_dimensions[21].hidden          # 그 밖의 빈 행은 숨김
    assert ws.row_dimensions[109].hidden
    # ④ 실행일(9/21)~차주 금요일(10/2)만 표시 — 이후 일자 행은 숨김
    assert not ws.row_dimensions[140].hidden     # 9/21 (실행일)
    assert not ws.row_dimensions[151].hidden     # 10/2 (차주 금요일)
    assert ws.row_dimensions[152].hidden         # 10/3부터 숨김
    assert ws.row_dimensions[167].hidden
    # ② 창 밖 지급일 행 숨김: 9/22는 표시, 10/14는 숨김 (합계에는 포함)
    assert not ws.row_dimensions[16].hidden
    assert ws.row_dimensions[17].hidden
    wb.close()


def test_경영보고_실행일_기준_지난_일자_숨김(tmp_path):
    """주중(수 9/23) 실행이면 월·화 행이 숨고 창은 9/23~10/2."""
    rep = _report()
    rep["meta"]["run_date"] = date(2026, 9, 23)
    out = tmp_path / "주간자금계획_경영보고_midweek.xlsx"
    mr.create_management_workbook(rep, out)
    assert mr.verify_management_workbook(out)

    wb = load_workbook(out)
    ws = wb["경영보고"]
    assert ws.row_dimensions[140].hidden         # 9/21 (지난 일자)
    assert ws.row_dimensions[141].hidden         # 9/22
    assert not ws.row_dimensions[142].hidden     # 9/23 (실행일)
    assert not ws.row_dimensions[151].hidden     # 10/2 (차주 금요일)
    assert ws.row_dimensions[152].hidden         # 10/3부터 숨김
    # 붉은 상자도 실행일 행(142)에서 시작
    top = ws["A142"].border.top
    assert top is not None and top.style == "medium"
    # ② 지난 지급일(9/22) 행도 숨김
    assert ws.row_dimensions[16].hidden
    assert ws.row_dimensions[17].hidden          # 10/14 (창 밖)
    wb.close()


def test_입금예정_표와_실지출_대조_확인란(tmp_path):
    """③ 입금예정: 온라인 예상(반영률 연동 수식)·확정입금·실입금 표시.
    ④ 지출예정: 당일 대조 결과를 실지출(G)·차이(H)로 잇고, 차이 나는
    행은 주황 강조 + 확인(I) 드롭다운 (2026-09-22 사용자 요청)."""
    rep = _report()
    rep["meta"]["run_date"] = date(2026, 9, 22)
    rep["week_expenses"][0]["요청ID"] = "물류-20260918-001"
    rep["forecast"]["today_actual"] = {
        "일자": date(2026, 9, 22), "온라인": 123_456.0,
        "기타입금": 50_000.0, "출금": 200_000.0, "순증감": -26_544.0}
    rep["adjustments"] = [{"일자": date(2026, 9, 23),
                           "조정입금": 1_000_000.0, "조정지출": 0.0,
                           "내용": "거래처 정산 입금"}]
    rep["intraday_check"] = {"대조내역": [
        {"요청ID": "물류-20260918-001", "팀명": "물류팀",
         "거래처": "한진택배", "예상금액": 500_000.0, "집행액": 200_000.0,
         "잔여": 300_000.0, "상태": "일부지급", "사유": "잔여 있음"}]}
    out = tmp_path / "주간자금계획_경영보고_intraday.xlsx"
    mr.create_management_workbook(rep, out)
    assert mr.verify_management_workbook(out)

    wb = load_workbook(out)
    ws = wb["경영보고"]
    # ③ 첫 행(114) = 실행일(9/22) 온라인 예상: ④의 요일평균(J141)×반영률
    assert ws["C114"].value == "온라인 예상"
    assert "J141" in str(ws["E114"].value) \
        and "$F$12" in str(ws["E114"].value)
    assert ws["F114"].value == 123456                # 실입금(은행 확인)
    assert str(ws["G114"].value).startswith("=IF(")  # 차이 수식
    # 기타 실입금(계획 외) 행
    assert ws["C115"].value == "기타 실입금" and ws["F115"].value == 50000
    # 9/23 확정입금 항목 행 (온라인 행 다음) — 일자·금액은 노란 칸
    conf = [r for r in range(114, 136) if ws[f"C{r}"].value == "확정입금"]
    assert conf and ws[f"E{conf[0]}"].value == 1000000
    assert ws[f"D{conf[0]}"].value == "거래처 정산 입금"
    assert str(ws[f"E{conf[0]}"].fill.start_color.rgb).endswith("FFF2CC")
    # 합계 행: 예상·실입금 SUM 수식
    assert "SUM(E114" in str(ws["E136"].value)
    assert "SUM(F114" in str(ws["F136"].value)
    # ② 실지출 대조: 집행액·차이 수식·주황 강조·확인 드롭다운
    assert ws["G16"].value == 200000
    assert "E16-G16" in str(ws["H16"].value)
    assert str(ws["G16"].fill.start_color.rgb).endswith("FFE699")
    dv = [v for v in ws.data_validations.dataValidation
          if "보류" in str(v.formula1) and "적용" in str(v.formula1)]
    assert dv and any("I16" in str(v.sqref) for v in dv)
    # 대조 결과가 없는 행(미래 지급일)은 실지출 빈칸
    assert ws["G17"].value in ("", None)
    wb.close()


def test_정기지출_체크_시트를_만들지_않는다(tmp_path):
    # 2026-09-23 사용자 확정: 정기지출 관련은 결과파일에 넣지 않는다 —
    # 누락 의심은 확인필요 파일의 '정기지출누락' 시트에서만 확인받는다
    rep = _report()
    rep["meta"]["run_date"] = date(2026, 9, 21)
    rep["recurring_check"] = [
        {"예정일": date(2026, 9, 25), "항목": "SKB",
         "예상금액": 220_000, "관리상태": "제외",
         "팀제출일": None, "팀제출금액": None,
         "누락": True, "판정": "누락 의심 — 팀 재제출 요청"},
    ]
    out = tmp_path / "주간자금계획_경영보고_check.xlsx"
    mr.create_management_workbook(rep, out)
    assert mr.verify_management_workbook(out)

    wb = load_workbook(out)
    assert wb.sheetnames == [mr.SHEET_NAME, mr.CEO_SHEET]
    wb.close()


def test_실지출_수기입력과_보류_제외_수식(tmp_path):
    """2026-09-23 사용자 요청: ② 실지출(G)은 전 행 수기 입력 칸(차이 H
    자동 계산), 확인(I)은 적용/보류 드롭다운 — '보류' 행은 ④·합계
    SUMIFS 조건에서 빠진다. ④ '상태' 열은 그 날짜의 지출 내역 나열로."""
    rep = _report()
    out = tmp_path / "주간자금계획_경영보고_hold.xlsx"
    mr.create_management_workbook(rep, out)
    wb = load_workbook(out)
    ws = wb["경영보고"]
    # 대조 결과가 없는 행도 실지출(G)은 노란 입력칸 + 차이(H) 수식
    assert str(ws["G16"].fill.start_color.rgb).endswith("FFF2CC")
    assert "E16-G16" in str(ws["H16"].value)
    # 예비 행(20)에도 요일·차이 수식과 적용/보류 드롭다운이 깔려 있다
    assert "WEEKDAY(A20" in str(ws["B20"].value)
    assert "E20-G20" in str(ws["H20"].value)
    dv = [v for v in ws.data_validations.dataValidation
          if "보류" in str(v.formula1)]
    assert dv and any("I20" in str(v.sqref) for v in dv)
    # ④ 지출 열은 보류 행을 뺀 SUMIFS, 합계(E110)도 보류 제외
    assert '"<>보류"' in str(ws["D140"].value)
    assert '"<>보류"' in str(ws["E110"].value)
    # ④ '지출 내역'(F, F:I 병합): 그 날짜의 내용·금액 나열 — 배열
    # 수식(TEXTJOIN)으로 저장해야 엑셀·LibreOffice에서 계산된다
    assert ws["F139"].value == "지출 내역"
    v140 = ws["F140"].value
    f140 = getattr(v140, "text", None) or str(v140)
    assert type(v140).__name__ == "ArrayFormula"
    assert f140.startswith("=IFERROR(")
    assert "TEXTJOIN" in f140 and "$D$16" in f140 and "보류" in f140
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "F140:I140" in merged
    wb.close()
