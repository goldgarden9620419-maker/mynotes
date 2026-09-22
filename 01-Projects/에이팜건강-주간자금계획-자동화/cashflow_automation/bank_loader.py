# -*- coding: utf-8 -*-
"""은행 거래내역 읽기·표준화 (16번 항목).

농협·우리은행·국민은행의 XLS/XLSX/CSV를 읽어 표준 항목으로 변환한다.
은행마다 열 이름이 달라 후보 키워드로 헤더를 탐지한다.
과거 거래 누적(bank_history.csv)도 여기서 관리한다 — 정기지출 분석과
온라인 입금 예측은 14일치보다 긴 이력이 필요하기 때문이다.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Optional

from common import (
    BANK_REFLECT_OK, BANKS, normalize_text, parse_amount, parse_date,
    parse_datetime,
)

# 표준 필드별 헤더 후보 (공백 제거 후 부분일치)
_FIELD_CANDIDATES: dict[str, list[str]] = {
    "거래일시": ["거래일시"],
    "거래일": ["거래일자", "거래날짜", "거래일", "일자", "날짜",
             "월/일", "월일"],
    "거래시간": ["거래시간", "시간"],
    "계좌": ["계좌번호", "계좌"],
    "출금액": ["출금금액", "출금액", "출금(원)", "지급금액", "지급(원)",
              "출금", "지급"],
    "입금액": ["입금금액", "입금액", "입금(원)", "입금"],
    "거래후잔액": ["거래후잔액", "거래후 잔액", "잔액(원)", "잔액"],
    "적요": ["적요", "거래내용", "거래구분"],
    "기재내용·상대방": ["기재내용", "보낸분/받는분", "의뢰인/수취인",
                    "거래기록사항", "기록사항", "상대방", "받는분", "내용"],
    "취급점": ["취급점", "거래점", "처리점", "송금점", "점포"],
}
_MIN_HEADER_MATCHES = 3

SUPPORTED_SUFFIXES = {".xls", ".xlsx", ".csv", ".txt", ".tsv",
                      ".htm", ".html"}

HISTORY_COLUMNS = ["거래일시", "거래일", "은행", "계좌", "출금액", "입금액",
                   "거래후잔액", "적요", "기재내용·상대방", "취급점", "자동분류"]


def _clean_header(value: Any) -> str:
    return re.sub(r"[\s ]+", "", normalize_text(value))


def _match_field(header: str) -> Optional[str]:
    key = _clean_header(header)
    if not key:
        return None
    for field, candidates in _FIELD_CANDIDATES.items():
        for cand in candidates:
            if cand.replace(" ", "") in key:
                return field
    return None


def _detect_header(rows: list[list[Any]]) -> Optional[tuple[int, dict[int, str]]]:
    """상위 30행에서 헤더 행을 찾는다."""
    best = None
    for i, row in enumerate(rows[:30]):
        mapping: dict[int, str] = {}
        for col, value in enumerate(row):
            field = _match_field(value)
            if field and field not in mapping.values():
                mapping[col] = field
        score = len(mapping)
        has_date = any(f in ("거래일", "거래일시") for f in mapping.values())
        has_amount = any(f in ("출금액", "입금액") for f in mapping.values())
        if score >= _MIN_HEADER_MATCHES and has_date and has_amount:
            if best is None or score > len(best[1]):
                best = (i, mapping)
    return best


_ACCOUNT_PATTERN = re.compile(r"(\d{4,8}(?:[- ]\d{2,6}){1,3}|\d{10,14})")


def _find_account_hint(rows: list[list[Any]], header_idx: int) -> str:
    """헤더 위쪽에서 '계좌번호: xxx' 형태의 안내를 찾는다.

    은행에 따라 '계좌번호' 라벨과 번호가 다른 셀에 있으므로
    행 전체를 이어붙여 검색한다.
    """
    for row in rows[:header_idx]:
        row_text = " ".join(normalize_text(v) for v in row if v not in
                            (None, ""))
        if "계좌" in row_text:
            m = _ACCOUNT_PATTERN.search(row_text)
            if m:
                return m.group(1).strip()
    return ""


def _looks_like_date_token(token: str) -> bool:
    """YYYYMMDD·YYMMDD로 읽히는 숫자 토큰은 계좌가 아니라 날짜다."""
    if len(token) == 8 and token[:2] in ("19", "20"):
        return 1 <= int(token[4:6]) <= 12 and 1 <= int(token[6:8]) <= 31
    if len(token) == 6:
        return 1 <= int(token[2:4]) <= 12 and 1 <= int(token[4:6]) <= 31
    return False


def _account_hint_from_filename(stem: str) -> str:
    """파일 안에 계좌번호가 없을 때(더존 내보내기 등) 파일명에서 찾는다.

    '국민 4577 20260922.xlsx'처럼 계좌 뒷자리를 파일명에 적어두면
    그 숫자를 계좌 구분자로 쓴다. 날짜로 보이는 토큰은 제외하고,
    공백은 계좌 구분자로 보지 않는다('4577 20260922' 오결합 방지).
    """
    text = normalize_text(stem)
    m = re.search(r"(\d{4,8}(?:-\d{2,6}){1,3}|\d{10,14})", text)
    if m and not _looks_like_date_token(m.group(1)):
        return m.group(1)
    best = ""
    for token in re.findall(r"\d{4,}", text):
        if _looks_like_date_token(token):
            continue
        if len(token) >= len(best):
            best = token
    return best


_MONTH_DAY_RE = re.compile(r"^(\d{1,2})[-/.월]\s*(\d{1,2})일?$")


def _parse_month_day(value: Any, ref: Optional[date] = None) -> Optional[date]:
    """'09-21'·'9/21'처럼 연도 없는 월-일(더존 내보내기)을 날짜로 만든다.

    거래내역에 미래 날짜는 없으므로, 기준일(ref)보다 일주일 넘게
    미래가 되는 해석은 지난해 날짜로 본다 (연말·연초 파일 대비).
    """
    text = normalize_text(value)
    m = _MONTH_DAY_RE.match(text)
    if not m:
        return None
    month, day = int(m.group(1)), int(m.group(2))
    if not (1 <= month <= 12 and 1 <= day <= 31):
        return None
    ref = ref or date.today()
    for year in (ref.year, ref.year - 1):
        try:
            cand = date(year, month, day)
        except ValueError:
            continue
        if cand <= ref + timedelta(days=7):
            return cand
    return None


# ---------------------------------------------------------------------------
# 파일 읽기
# ---------------------------------------------------------------------------

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
# 스타일 XML이 깨진 xlsx 복구용 최소 스타일 (셀 서식 참조가 남아 있어도
# 안전하게 열리도록 빈 서식을 넉넉히 채운다)
_XF = '<xf numFmtId="0" fontId="0" fillId="0" borderId="0"/>'
_MIN_STYLES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<styleSheet xmlns="http://schemas.openxmlformats.org/'
    'spreadsheetml/2006/main">'
    '<fonts count="1"><font><sz val="11"/><name val="Calibri"/>'
    '</font></fonts>'
    '<fills count="2"><fill><patternFill patternType="none"/></fill>'
    '<fill><patternFill patternType="gray125"/></fill></fills>'
    '<borders count="1"><border><left/><right/><top/><bottom/>'
    '<diagonal/></border></borders>'
    f'<cellStyleXfs count="1">{_XF}</cellStyleXfs>'
    f'<cellXfs count="512">{_XF * 512}</cellXfs>'
    '</styleSheet>')


def _read_xlsx_plain(path: Path) -> list[list[Any]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_xlsx_rows(path: Path) -> list[list[Any]]:
    """xlsx 읽기. 스타일 XML이 깨진 은행 파일은 복구해서 재시도한다."""
    try:
        return _read_xlsx_plain(path)
    except Exception:
        import tempfile
        import zipfile
        try:
            with zipfile.ZipFile(path) as zin:
                names = zin.namelist()
                with tempfile.NamedTemporaryFile(
                        suffix=".xlsx", delete=False) as tmp:
                    tmp_path = Path(tmp.name)
                with zipfile.ZipFile(tmp_path, "w",
                                     zipfile.ZIP_DEFLATED) as zout:
                    for name in names:
                        if name == "xl/styles.xml":
                            zout.writestr(name, _MIN_STYLES)
                        else:
                            zout.writestr(name, zin.read(name))
        except Exception:
            raise ValueError("xlsx 파일을 열 수 없습니다 (손상)")
        try:
            return _read_xlsx_plain(tmp_path)
        finally:
            try:
                tmp_path.unlink()
            except OSError:
                pass


def _read_xls_rows(path: Path) -> list[list[Any]]:
    import xlrd
    book = xlrd.open_workbook(str(path))
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(datetime(*xlrd.xldate_as_tuple(
                    cell.value, book.datemode)))
            else:
                row.append(cell.value)
        rows.append(row)
    return rows


def _read_text(path: Path) -> Optional[str]:
    raw = path.read_bytes()
    encodings = ("utf-8-sig", "cp949", "euc-kr", "utf-8")
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        encodings = ("utf-16",) + encodings
    for encoding in encodings:
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return None


class _HtmlTableParser:
    """은행이 .xls 확장자로 주는 HTML 표를 행 목록으로 바꾼다."""

    def __init__(self):
        from html.parser import HTMLParser

        outer = self

        class _P(HTMLParser):
            def __init__(self):
                super().__init__(convert_charrefs=True)
                self.rows: list[list[str]] = []
                self._row: Optional[list[str]] = None
                self._cell: Optional[list[str]] = None

            def handle_starttag(self, tag, attrs):
                if tag == "tr":
                    self._row = []
                elif tag in ("td", "th"):
                    self._cell = []
                elif tag == "br" and self._cell is not None:
                    self._cell.append(" ")

            def handle_endtag(self, tag):
                if tag in ("td", "th") and self._cell is not None:
                    self._row = self._row if self._row is not None else []
                    self._row.append(" ".join(
                        "".join(self._cell).split()))
                    self._cell = None
                elif tag == "tr" and self._row is not None:
                    self.rows.append(self._row)
                    self._row = None

            def handle_data(self, data):
                if self._cell is not None:
                    self._cell.append(data)

        outer.parser = _P()

    def parse(self, text: str) -> list[list[str]]:
        self.parser.feed(text)
        self.parser.close()
        return self.parser.rows


def _read_raw_rows(path: Path) -> list[list[Any]]:
    """확장자와 무관하게 파일 내용을 판별해 표를 읽는다.

    은행 다운로드 파일은 확장자와 실제 형식이 다른 경우가 많다:
    .xls인데 HTML 표, .xlsx인데 구형 xls, 스타일이 깨진 xlsx, CSV 등.
    zip(xlsx) → OLE(구형 xls) → HTML 표 → CSV/TSV 순서로 시도한다
    (2026-09-21 사용자 요청: 형식 구분 없이 읽기).
    """
    with open(path, "rb") as f:
        head = f.read(8)
    if head[:2] == b"PK":
        return _read_xlsx_rows(path)
    if head == _OLE_MAGIC:
        return _read_xls_rows(path)
    text = _read_text(path)
    if text is not None:
        stripped = text.lstrip().lower()
        if stripped.startswith("<") and (
                "<table" in stripped or "<tr" in stripped
                or "<html" in stripped):
            rows = _HtmlTableParser().parse(text)
            if rows:
                return rows
        delimiter = "\t" if text.count("\t") > text.count(",") else ","
        return [row for row in
                csv.reader(io.StringIO(text), delimiter=delimiter)]
    raise ValueError("알 수 없는 파일 형식입니다 "
                     "(xlsx·xls·HTML 표·CSV 어느 것으로도 읽지 못함)")


def load_bank_file(path: Path, bank: str) -> tuple[list[dict], list[dict]]:
    """은행 파일 하나를 표준 행으로 변환한다. (행, 확인필요) 반환."""
    path = Path(path)
    issues: list[dict] = []
    try:
        raw_rows = _read_raw_rows(path)
    except Exception as exc:
        return [], [{"구분": "손상된 파일", "은행": bank,
                     "내용": f"파일 열기 실패: {exc}", "원본파일": path.name}]
    header = _detect_header(raw_rows)
    if header is None:
        return [], [{"구분": "필수 열 누락", "은행": bank,
                     "내용": "거래내역 헤더를 찾지 못했습니다",
                     "원본파일": path.name}]
    header_idx, mapping = header
    account_hint = (_find_account_hint(raw_rows, header_idx)
                    or _account_hint_from_filename(path.stem))

    rows: list[dict] = []
    for order, raw in enumerate(raw_rows[header_idx + 1:], start=1):
        values = {field: (raw[col] if col < len(raw) else None)
                  for col, field in mapping.items()}
        row = _normalize_bank_row(values, bank, path.name, account_hint, order)
        if row is None:
            continue
        if row.get("_잘못된날짜"):
            issues.append({"구분": "잘못된 날짜", "은행": bank,
                           "내용": f"{order}행 거래일 해석 불가",
                           "원본파일": path.name})
            continue
        rows.append(row)
    if not rows:
        issues.append({"구분": "은행 자료 누락", "은행": bank,
                       "내용": "읽을 수 있는 거래가 없습니다",
                       "원본파일": path.name})
    return rows, issues


def _excel_serial(value: Any) -> Optional[datetime]:
    """엑셀 날짜 일련값 → datetime. 스타일이 깨진 xlsx를 복구해 읽으면
    날짜 서식이 사라져 숫자로 오므로 그럴듯한 범위만 날짜로 되살린다."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if not 20000 <= float(value) <= 80000:      # 1954년~2119년
        return None
    return datetime(1899, 12, 30) + timedelta(days=float(value))


def _normalize_bank_row(values: dict, bank: str, filename: str,
                        account_hint: str, order: int) -> Optional[dict]:
    tx_datetime = (parse_datetime(values.get("거래일시"))
                   or _excel_serial(values.get("거래일시")))
    tx_date = parse_date(values.get("거래일")) or (
        tx_datetime.date() if tx_datetime else None)
    if tx_date is None:
        serial = _excel_serial(values.get("거래일"))
        if serial is not None:
            tx_date = serial.date()
    if tx_date is None:
        # 더존 내보내기는 '09-21'처럼 연도 없는 월-일을 쓴다
        tx_date = _parse_month_day(values.get("거래일"))
    if tx_datetime is None and tx_date is None:
        raw_date = normalize_text(values.get("거래일")
                                  or values.get("거래일시"))
        if not raw_date:
            return None  # 빈 행
        # '총 862건', '합계'('합  계' 포함) 같은 요약 행은 거래가 아니다
        compact = raw_date.replace(" ", "")
        if compact.startswith("총") or "합계" in compact \
                or compact.endswith("건"):
            return None
        return {"_잘못된날짜": True}
    if tx_datetime is None and values.get("거래시간") is not None:
        t = normalize_text(values.get("거래시간"))
        m = re.match(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", t)
        if m and tx_date is not None:
            tx_datetime = datetime.combine(tx_date, time(
                int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)))
    if tx_date is None and tx_datetime is not None:
        tx_date = tx_datetime.date()

    out_amount = parse_amount(values.get("출금액")) or 0.0
    in_amount = parse_amount(values.get("입금액")) or 0.0
    balance = parse_amount(values.get("거래후잔액"))
    if out_amount == 0 and in_amount == 0 and balance is None:
        return None  # 합계·안내 행

    return {
        "거래일시": tx_datetime,
        "거래일": tx_date,
        "은행": bank,
        "계좌": normalize_text(values.get("계좌")) or account_hint,
        "출금액": out_amount,
        "입금액": in_amount,
        "거래후잔액": balance,
        "적요": normalize_text(values.get("적요")),
        "기재내용·상대방": normalize_text(values.get("기재내용·상대방")),
        "취급점": normalize_text(values.get("취급점")),
        "자동분류": "",
        "내부이체": False,
        "정기지출후보": False,
        "현금유출입": in_amount - out_amount,
        "원본파일": filename,
        "반영상태": BANK_REFLECT_OK,
        "row_order": order,
    }


def load_all_banks(cfg) -> dict:
    """세 은행 폴더의 모든 지원 파일을 읽는다."""
    all_rows: list[dict] = []
    issues: list[dict] = []
    missing_banks: list[str] = []
    for bank in BANKS:
        bank_dir = cfg.bank_dir(bank)
        files = sorted(p for p in bank_dir.glob("*")
                       if p.suffix.lower() in SUPPORTED_SUFFIXES
                       and not p.name.startswith("~$"))
        if not files:
            missing_banks.append(bank)
            issues.append({"구분": "은행 자료 누락", "은행": bank,
                           "내용": f"{bank} 폴더에 거래내역 파일이 없습니다",
                           "원본파일": ""})
            continue
        for path in files:
            rows, file_issues = load_bank_file(path, bank)
            all_rows.extend(rows)
            issues.extend(file_issues)
    unify_account_labels(all_rows)
    return {"rows": all_rows, "issues": issues, "missing_banks": missing_banks}


def unify_account_labels(rows: list[dict]) -> list[dict]:
    """짧은 계좌 라벨을 같은 은행의 전체 계좌번호와 뒷자리로 맞춰 통일한다.

    더존 내보내기처럼 계좌번호가 없는 파일은 파일명 뒷자리(예: '국민 4577')로
    구분하는데, 같은 계좌의 은행 원본('…154577')과 라벨이 다르면 잔액이
    이중 합산되고 내부이체 짝 판정도 어긋난다. 뒷자리가 일치하는 전체
    번호가 같은 은행에 정확히 하나면 그 라벨로 바꿔 같은 계좌로 묶는다.
    """
    by_bank: dict[str, set[str]] = {}
    for row in rows:
        account = row.get("계좌") or ""
        if account:
            by_bank.setdefault(row.get("은행") or "", set()).add(account)
    mapping: dict[tuple, str] = {}
    for bank, accounts in by_bank.items():
        digits = {a: re.sub(r"\D", "", a) for a in accounts}
        for short in accounts:
            if len(digits[short]) < 4:
                continue
            longer = [a for a in accounts
                      if a != short and len(digits[a]) > len(digits[short])
                      and digits[a].endswith(digits[short])]
            if len(longer) == 1:
                mapping[(bank, short)] = longer[0]
    for row in rows:
        key = (row.get("은행") or "", row.get("계좌") or "")
        if key in mapping:
            row["계좌"] = mapping[key]
    return rows


# ---------------------------------------------------------------------------
# 잔액 요약
# ---------------------------------------------------------------------------

def summarize_balances(rows: list[dict]) -> tuple[dict[tuple, float], float]:
    """(은행, 계좌)별 최신 거래후잔액과 전체 합계."""
    latest: dict[tuple, tuple] = {}
    for row in rows:
        if row.get("반영상태") != BANK_REFLECT_OK:
            continue
        balance = row.get("거래후잔액")
        if balance is None:
            continue
        key = (row["은행"], row.get("계좌") or "")
        sort_key = (row.get("거래일") or date.min,
                    row.get("거래일시") or datetime.min,
                    row.get("row_order", 0))
        if key not in latest or sort_key >= latest[key][0]:
            latest[key] = (sort_key, balance)
    balances = {key: value[1] for key, value in latest.items()}
    return balances, sum(balances.values())


# ---------------------------------------------------------------------------
# 거래 이력 누적 (정기지출·입금 예측용)
# ---------------------------------------------------------------------------

def load_history(history_path: Path) -> list[dict]:
    history_path = Path(history_path)
    if not history_path.exists():
        return []
    rows: list[dict] = []
    with open(history_path, encoding="utf-8-sig", newline="") as f:
        for record in csv.DictReader(f):
            row = {
                "거래일시": parse_datetime(record.get("거래일시")),
                "거래일": parse_date(record.get("거래일")),
                "은행": record.get("은행", ""),
                "계좌": record.get("계좌", ""),
                "출금액": parse_amount(record.get("출금액")) or 0.0,
                "입금액": parse_amount(record.get("입금액")) or 0.0,
                "거래후잔액": parse_amount(record.get("거래후잔액")),
                "적요": record.get("적요", ""),
                "기재내용·상대방": record.get("기재내용·상대방", ""),
                "취급점": record.get("취급점", ""),
                "자동분류": record.get("자동분류", ""),
                "내부이체": record.get("내부이체", "") == "True",
                "정기지출후보": False,
                "현금유출입": (parse_amount(record.get("입금액")) or 0.0)
                            - (parse_amount(record.get("출금액")) or 0.0),
                "원본파일": record.get("원본파일", "history"),
                "반영상태": BANK_REFLECT_OK,
                "row_order": 0,
            }
            if row["거래일"] is not None:
                rows.append(row)
    return rows


def save_history(history_path: Path, rows: list[dict]) -> None:
    history_path = Path(history_path)
    history_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = history_path.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(HISTORY_COLUMNS + ["내부이체", "원본파일"])
        for row in sorted(rows, key=lambda r: (r.get("거래일") or date.min,
                                               r.get("거래일시") or datetime.min)):
            writer.writerow([
                row["거래일시"].strftime("%Y-%m-%d %H:%M:%S")
                if row.get("거래일시") else "",
                row["거래일"].strftime("%Y-%m-%d") if row.get("거래일") else "",
                row.get("은행", ""), row.get("계좌", ""),
                row.get("출금액", 0), row.get("입금액", 0),
                "" if row.get("거래후잔액") is None else row["거래후잔액"],
                row.get("적요", ""), row.get("기재내용·상대방", ""),
                row.get("취급점", ""), row.get("자동분류", ""),
                bool(row.get("내부이체")), row.get("원본파일", ""),
            ])
    tmp.replace(history_path)
